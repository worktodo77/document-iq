"""Spreadsheet fidelity: a cell's value, or the meaning of its value, is never
lost or changed without a word said.

The 2026-09-10 fidelity sweep found ``.xlsx``/``.xls`` extraction
(``src/dociq/ingest/extract.py``: ``_extract_xlsx``, ``_extract_xls`` and
their helpers) losing cell values and meanings silently -- the confirmed
defect (E1) together with its class. These tests pin the FIXED behavior,
against the corpus fixtures (``workbook_constructs`` ->
``18_workbook_constructs.xlsx``, ``legacy_workbook`` ->
``19_legacy_workbook.xls``) and against small workbooks each test builds for
the siblings a fixture does not hold (review-fix round 1, 2026-09-14): time
and duration cells, stored empty formula results, numbers that are no date,
header/footer codes, threaded comments, and the rest named per test.

Output is pinned EXACTLY where the wording is decided: a substring check lets
a wrong line, a wrong order or a stray code letter through.

Every sentinel word is invented and is not a substring of any other sentinel
in the same workbook (client-data rule D-12).
"""

from __future__ import annotations

import datetime
import html
import io
import re
import zipfile
import xml.etree.ElementTree as ET

import make_fixtures
import openpyxl
import openpyxl.chart
import openpyxl.comments
import pytest

from dociq.contracts import ProcessingStatus, RunConfig
from dociq.identify import bates
from dociq.ingest import extract as ex
from dociq.ingest import walker
from dociq.ingest.pagemodel import normalize

from .conftest import FIXTURES

UNREAD = ex.M_XLSX_SHEET_UNREAD
PART_UNREAD = "spreadsheet part could not be read"
"""The transient marker's phrase, spelled out: a test that read it from the
module would pass against a module that never defined it."""
PAGE_NOTE_XLSX = "XLSX has no page boundaries; one synthetic page per worksheet"
PAGE_NOTE_XLS = "XLS has no page boundaries; one synthetic page per worksheet"
FORMULA_NOTE_ONE = ("1 cell(s): a formula cell had no stored value; the formula is "
                    "shown in the cell's place instead of the missing value")
CHART1_NOTE = (f"chartsheet 'Chart1': its chart, and any print header or footer on it, "
               f"were not read ({UNREAD})")
CAP_LINE = "[not read: the row cap was reached before this sheet]"

REGISTER_PAGE = (
    "[sheet: Register]\n"
    "OSPREY\n"
    "Ref\tAmount\tNote\n"
    "\t10\t15%\n"
    "\t20\t15.00%\n"
    "HOLLY ¶ IVY\t=SUM(B2:B3)\t\tTAB STOP\n"
    "[blank rows 5-6]\n"
    "JUNIPER\n"
    "KESTREL <https://example.invalid/LARK>\n"
    "[comment on B2 by NIGHTJAR] MAGPIE\n"
    "PUFFIN")
"""Fixture 18's Register page, whole: every E-finding's output in place."""

LEGACY_PAGE = ("[sheet: Legacy]\n"
               "RAVENXLS\t2024-07-16\tTRUE\t#DIV/0!\t5\n"
               "[blank row 2]\n"
               "SWIFT")
"""Fixture 19's only page, whole."""


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _pages(name: str, opt: ex.ExtractOptions | None = None):
    path = FIXTURES / name
    return ex.extract(path.name, path.read_bytes(), opt)


def _page_with(pages, prefix: str):
    """The first page whose text starts with ``prefix``, or ``None``.

    Never indexes into ``pages`` directly -- a wrong page count must fail an
    assertion with a clear message, not raise ``IndexError`` first.
    """
    return next((p for p in pages if p.text.startswith(prefix)), None)


def _assert_found(text: str, sentinel: str, note: str = "") -> None:
    """``sentinel in text``, via ``str.find`` so a miss reports its position."""
    pos = text.find(sentinel)
    assert pos != -1, (
        f"{sentinel!r} not found in extracted text (str.find -> {pos}). "
        f"{note}Text was:\n{text!r}")


def _texts(doc) -> list[str]:
    return [p.text for p in doc.pages]


def _assert_doc(doc, pages: list[str], notes: list[str]) -> None:
    """Status FULL, and the pages and notes exactly as given."""
    assert doc.status is ProcessingStatus.FULL, (doc.status, doc.error)
    assert _texts(doc) == pages, f"pages were {_texts(doc)!r}"
    assert list(doc.notes) == notes, f"notes were {list(doc.notes)!r}"


def _xlsx(build) -> bytes:
    """An ``.xlsx`` built by openpyxl: ``build(workbook)`` fills it in."""
    wb = openpyxl.Workbook()
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _patch(raw: bytes, edits: dict, renames: dict | None = None) -> bytes:
    """Rewrite package parts as raw XML -- for constructs openpyxl cannot
    write, or writes differently from Excel. ``edits`` maps a part name to
    ``fn(old text or None) -> new text or None`` (``None`` deletes the part;
    a part not present is added when ``fn(None)`` returns text); ``renames``
    maps an old part name to a new one."""
    renames = renames or {}
    zin = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        seen = set()
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename in edits:
                seen.add(info.filename)
                new = edits[info.filename](data.decode("utf-8"))
                if new is None:
                    continue
                data = new.encode("utf-8")
            zout.writestr(renames.get(info.filename, info.filename), data)
        for name, fn in edits.items():
            if name not in seen:
                new = fn(None)
                if new is not None:
                    zout.writestr(name, new.encode("utf-8"))
    return out.getvalue()


def _part_names(raw: bytes) -> list[str]:
    return zipfile.ZipFile(io.BytesIO(raw)).namelist()


def _extract(name: str, raw: bytes):
    return ex.extract(name, raw)


def _detected_dates(tmp_path, files: dict[str, bytes]) -> dict[str, tuple]:
    """``{filename: detected_dates}`` from a real ``walker.run`` over a folder
    holding exactly ``files`` -- the product path that fills the index and the
    screen's document date."""
    src = tmp_path / "src"
    src.mkdir()
    for name, data in files.items():
        (src / name).write_bytes(data)
    config = RunConfig(source_root=str(src), output_root=str(tmp_path / "out"),
                       ocr_engine_version=ex.ocr_engine_version())
    result = walker.run(config, walker.WalkOptions(ocr_enabled=False, resume=False))
    return {d.filename: d.detected_dates for d in result.documents}


def _set_header_footer(sheet_xml: str, header_footer_xml: str) -> str:
    return re.sub(r"<headerFooter>.*?</headerFooter>", header_footer_xml,
                  sheet_xml, flags=re.S)


def _is(ref: str | None, text: str) -> str:
    """An inline-string ``<c>``, with ``r=`` only when ``ref`` is given."""
    r = f' r="{ref}"' if ref else ""
    return f'<c{r} t="inlineStr"><is><t>{text}</t></is></c>'


def _rows_xlsx(sheet_data: str, dimension: str | None, *, name: str = "Order",
               head: str = "") -> bytes:
    """A one-sheet ``.xlsx`` whose ``sheetData`` is exactly ``sheet_data``
    (raw XML, rows and cells in whatever order and form it says), declaring
    ``dimension`` -- or none at all when it is ``None``. ``head`` is raw XML
    put before ``sheetData`` (``sheetFormatPr``, ``cols``)."""
    def build(wb):
        ws = wb.active
        ws.title = name
        ws["A1"] = "PLACEHOLDER"

    def sheet(xml: str) -> str:
        xml = re.sub(r"<sheetFormatPr[^>]*/>", "", xml)
        xml = re.sub(r"<sheetData>.*?</sheetData>|<sheetData/>",
                     lambda _m: head + f"<sheetData>{sheet_data}</sheetData>", xml, flags=re.S)
        dim = "" if dimension is None else f'<dimension ref="{dimension}"/>'
        return re.sub(r'<dimension ref="[^"]*"/>', lambda _m: dim, xml)

    return _patch(_xlsx(build), {"xl/worksheets/sheet1.xml": sheet})


def _once_in_pool(fn, nth: int = 1, when=lambda *a, **k: True):
    """``fn``, raising MemoryError on its ``nth`` qualifying call made on an
    extraction-POOL thread -- never on the walker's serial retry: a failure
    that load caused and a second read clears."""
    import threading

    state = {"calls": 0}

    def wrapper(*args, **kwargs):
        if (threading.current_thread().name.startswith("dociq-doc")
                and when(*args, **kwargs)):
            state["calls"] += 1
            if state["calls"] == nth:
                raise MemoryError("simulated memory pressure inside the pool")
        return fn(*args, **kwargs)

    return wrapper


# ---------------------------------------------------------------------------
# E1: formulas with no stored value
# ---------------------------------------------------------------------------


def test_e1_formula_with_no_stored_value_shows_the_formula():
    """E1 (confirmed): ``data_only=True`` never computes a formula, so the
    cell showed nothing. The formula is shown and ONE marked note, pinned
    exactly, counts it; a workbook with no such cell carries no such note.

    Sibling (review A2): Excel stores the result of ``=IF(x="","",x)`` as
    ``<c t="str"><f>..</f><v/></c>`` -- a stored value that is the empty
    string. It is blank, not missing: no formula text, no note. Only a
    ``t="str"`` formula with NO ``<v>`` at all is missing."""
    doc = _pages("18_workbook_constructs.xlsx")
    reg = _page_with(doc.pages, "[sheet: Register]")
    assert reg is not None, "no page starts with '[sheet: Register]'"
    _assert_found(reg.text, "HOLLY ¶ IVY\t=SUM(B2:B3)\t\tTAB STOP",
                  "the formula cell's own formula, in its row. ")
    formula_notes = [n for n in doc.notes if ex.M_XLSX_FORMULA_NO_VALUE in n]
    assert formula_notes == [FORMULA_NOTE_ONE], f"notes were {list(doc.notes)!r}"
    assert ex.has_evidence_marker(FORMULA_NOTE_ONE)
    plain = _pages("06_register.xlsx")
    assert not [n for n in plain.notes if ex.M_XLSX_FORMULA_NO_VALUE in n], plain.notes

    def build(wb):
        ws = wb.active
        ws.title = "Idiom"
        ws["A1"], ws["B1"], ws["C1"] = "ALDER", "PLACEHOLDER1", "BIRCH"
        ws["A2"], ws["B2"] = "CEDAR", "PLACEHOLDER2"
        ws["A3"], ws["B3"] = "DOGWOOD", "PLACEHOLDER3"

    def cells(xml: str) -> str:
        xml = re.sub(r'<c r="B1"[^>]*>.*?</c>',
                     '<c r="B1" t="str"><f>IF(A1="ALDER","",A1)</f><v/></c>', xml)
        xml = re.sub(r'<c r="B2"[^>]*>.*?</c>',
                     '<c r="B2" t="str"><f>IF(A2="","",A2)</f><v></v></c>', xml)
        return re.sub(r'<c r="B3"[^>]*>.*?</c>',
                      '<c r="B3" t="str"><f>A3</f></c>', xml)

    raw = _patch(_xlsx(build), {"xl/worksheets/sheet1.xml": cells})
    _assert_doc(_extract("idiom.xlsx", raw),
                ["[sheet: Idiom]\nALDER\t\tBIRCH\nCEDAR\nDOGWOOD\t=A3"],
                [FORMULA_NOTE_ONE, PAGE_NOTE_XLSX])

    # Review r2 (unverified B): the raw-cell lookup behind the stored empty
    # string must line up with openpyxl's rows and columns on a SPARSE sheet:
    # a column after a gap (E4), a row with no r= after a numbered one, a cell
    # with no r= after a numbered one, a row after a gap in row numbers. A
    # reader that ignored r= shows the formula and a false marked note.
    empty = '<c{ref} t="str"><f>IF(1,"","")</f><v/></c>'
    sparse = ('<row r="1">' + _is("A1", "LBLONE") + "</row>"
              '<row r="4">' + _is("A4", "LBLFOUR") + empty.format(ref=' r="E4"') + "</row>"
              "<row>" + empty.format(ref="") + _is(None, "RLESSWORD") + "</row>"
              '<row r="7">' + _is("B7", "AFTERGAPWORD") + empty.format(ref="")
              + _is("E7", "EASTWORD") + "</row>")
    _assert_doc(_extract("sparse.xlsx", _rows_xlsx(sparse, "A1:E7", name="Sparse")),
                ["[sheet: Sparse]\nLBLONE\n[blank rows 2-3]\nLBLFOUR\n\tRLESSWORD\n"
                 "[blank row 6]\n\tAFTERGAPWORD\t\t\tEASTWORD"],
                [PAGE_NOTE_XLSX])


def test_e1_note_is_true_for_what_each_cell_shows():
    """E1's note says "the formula is shown" -- true only where a formula is.
    An ARRAY formula with no stored value carries its text on the object and
    shows it; a DATA-TABLE formula has no text and shows the placeholder, and
    the note says so for each count, exactly."""
    from openpyxl.worksheet.formula import ArrayFormula

    def build(wb):
        ws = wb.active
        ws.title = "Arrays"
        ws["A1"], ws["A2"] = 2, 3
        ws["B1"] = ArrayFormula("B1", "=SUM(A1:A2*2)")
        ws["C1"] = "PLACEHOLDERTABLE"

    raw = _patch(_xlsx(build), {"xl/worksheets/sheet1.xml": lambda x: re.sub(
        r'<c r="C1"[^>]*>.*?</c>',
        '<c r="C1"><f t="dataTable" ref="C1" dt2D="0" dtr="0" r1="A1"/></c>', x)})
    both = ("2 cell(s): a formula cell had no stored value; the formula is shown "
            "in the cell's place for 1; for the other 1 the formula's own text "
            "was not available, so '[formula with no stored value]' is shown")
    _assert_doc(_extract("arrays.xlsx", raw),
                ["[sheet: Arrays]\n2\t=SUM(A1:A2*2)\t[formula with no stored value]\n3"],
                [both, PAGE_NOTE_XLSX])

    only_table = _patch(_xlsx(build), {"xl/worksheets/sheet1.xml": lambda x: re.sub(
        r'<c r="B1"[^>]*>.*?</c>', "", re.sub(
            r'<c r="C1"[^>]*>.*?</c>',
            '<c r="C1"><f t="dataTable" ref="C1" dt2D="0" dtr="0" r1="A1"/></c>', x),
        flags=re.S)})
    placeholder_only = (
        "1 cell(s): a formula cell had no stored value; the formula's own text "
        "was not available either, so '[formula with no stored value]' is shown "
        "in the cell's place")
    _assert_doc(_extract("table.xlsx", only_table),
                ["[sheet: Arrays]\n2\t\t[formula with no stored value]\n3"],
                [placeholder_only, PAGE_NOTE_XLSX])


# ---------------------------------------------------------------------------
# E2: dates, times, durations, and numbers that are no date
# ---------------------------------------------------------------------------


def _clock_xls() -> bytes:
    return make_fixtures.xls_bytes([{
        "name": "Clocks",
        "cells": [(0, 0, "GEARWHEEL", 0),
                  (0, 1, 0.35, 20),        # h:mm
                  (0, 2, 1.5, 46),         # [h]:mm:ss, over a day
                  (0, 3, 0.0, 20),         # midnight
                  (0, 4, 0.75, 164),       # custom hh:mm
                  (0, 5, 61.25, 46),       # a duration past the 1900 leap-day quirk
                  (0, 6, 45489.5, 22)],    # a real date and time
    }], formats={164: "hh:mm"})


def _clock_xls_1904() -> bytes:
    return make_fixtures.xls_bytes([{
        "name": "Clocks1904",
        "cells": [(0, 0, "SPROCKET", 0), (0, 1, 0.35, 20), (0, 2, 1.5, 46),
                  (0, 3, 44027.0, 14)],    # 2024-07-16 in the 1904 system
    }], datemode=1)


def _clock_xlsx(epoch=None) -> bytes:
    def build(wb):
        if epoch is not None:
            wb.epoch = epoch
        ws = wb.active
        ws.title = "Clocks"
        ws["A1"] = "GEARWHEEL"
        ws["B1"] = datetime.time(8, 24)
        ws["C1"] = datetime.timedelta(days=1.5)
        ws["C1"].number_format = "[h]:mm:ss"
        ws["D1"] = datetime.time(0, 0)
        ws["E1"] = datetime.datetime(2024, 7, 16)
        ws["E1"].number_format = "yyyy-mm-dd"
    return _xlsx(build)


def test_e2_xls_date_formatted_cell_reads_as_iso_date():
    """E2 (confirmed): a date-formatted ``.xls`` cell reads as an ISO date;
    fixture 19's whole page is pinned (E3, E4, E15 with it).

    Sibling (review A1): xlrd types a TIME or DURATION cell as a date too, and
    ``xldate_as_datetime`` put it on the 1899-12-31 epoch -- a date the file
    does not hold. A time reads as a time, a duration as hours past 24, and
    the 1904 date system (``datemode`` 1) reads its own dates."""
    doc = _pages("19_legacy_workbook.xls")
    _assert_doc(doc, [LEGACY_PAGE], [PAGE_NOTE_XLS])
    _assert_doc(_extract("clocks.xls", _clock_xls()),
                ["[sheet: Clocks]\nGEARWHEEL\t08:24:00\t36:00:00\t00:00:00\t18:00:00"
                 "\t1470:00:00\t2024-07-16 12:00:00"],
                [PAGE_NOTE_XLS])
    _assert_doc(_extract("clocks1904.xls", _clock_xls_1904()),
                ["[sheet: Clocks1904]\nSPROCKET\t08:24:00\t36:00:00\t2024-07-16"],
                [PAGE_NOTE_XLS])


def test_e2_xlsx_duration_reads_as_hours_like_xls():
    """The ``.xlsx`` half of A1: openpyxl hands a ``[h]:mm:ss`` cell over as a
    ``timedelta``, which printed ``1 day, 12:00:00``. Both formats read a
    duration the same way, as Excel shows it -- in either date system."""
    from openpyxl.utils.datetime import MAC_EPOCH

    for epoch in (None, MAC_EPOCH):
        _assert_doc(_extract("clocks.xlsx", _clock_xlsx(epoch)),
                    ["[sheet: Clocks]\nGEARWHEEL\t08:24:00\t36:00:00\t00:00:00\t2024-07-16"],
                    [PAGE_NOTE_XLSX])


def _shift_cells():
    """Time-only formats holding a serial of 1 or more (review r2, A8): a
    weekly total past 24 hours under ``h:mm``, and a real timestamp whose
    format hides its date. ``(value, xlsx format, xls format index)``."""
    return [(1.25, "h:mm", 20), (2.5, "hh:mm:ss", 164), (1.75, "h:mm AM/PM", 18),
            (3.5, "mm:ss", 45), (45489.25, "h:mm", 20), (0.75, "h:mm", 20)]


def _shift_xls(datemode: int) -> bytes:
    return make_fixtures.xls_bytes([{
        "name": "Shifts",
        "cells": [(0, 0, "SHIFTLOG", 0)] + [
            (0, c, value, idx) for c, (value, _f, idx) in enumerate(_shift_cells(), start=1)],
    }], formats={164: "hh:mm:ss"}, datemode=datemode)


def _shift_xlsx(epoch) -> bytes:
    def build(wb):
        if epoch is not None:
            wb.epoch = epoch
        ws = wb.active
        ws.title = "Shifts"
        ws["A1"] = "SHIFTLOG"
        for c, (value, fmt, _idx) in enumerate(_shift_cells(), start=2):
            ws.cell(1, c, value).number_format = fmt
    return _xlsx(build)


def test_e2_time_and_duration_cells_add_no_detected_date(tmp_path):
    """Review A1's harm: the invented 1899/1900 dates reached
    ``DocumentRecord.detected_dates`` -- the index and the screen's document
    date. Through the real walk, each workbook's only date is its real one.

    Sibling (review r2, A8): a time-only format (``h:mm``, ``hh:mm:ss``,
    ``h:mm AM/PM``, ``mm:ss``) holding a serial of 1 or more showed an
    invented 1900/1904 date in both formats. The format shows no date, so
    none is rendered: the number reads as stored, with an unmarked note --
    in both formats and both date systems -- and a time under 1 stays a
    time."""
    from openpyxl.utils.datetime import MAC_EPOCH

    page = "[sheet: Shifts]\nSHIFTLOG\t1.25\t2.5\t1.75\t3.5\t45489.25\t18:00:00"
    note = ("5 cell(s) formatted as a time with no date held a value of a day or "
            "more; the number is shown as stored")
    for name, raw in (("shift.xls", _shift_xls(0)), ("shift1904.xls", _shift_xls(1))):
        _assert_doc(_extract(name, raw), [page], [note, PAGE_NOTE_XLS])
    for epoch in (None, MAC_EPOCH):
        _assert_doc(_extract("shift.xlsx", _shift_xlsx(epoch)), [page], [note, PAGE_NOTE_XLSX])
    assert not ex.has_evidence_marker(note)

    dates = _detected_dates(tmp_path, {
        "clocks.xls": _clock_xls(), "clocks1904.xls": _clock_xls_1904(),
        "clocks.xlsx": _clock_xlsx(), "shift.xls": _shift_xls(0),
        "shift1904.xls": _shift_xls(1), "shift.xlsx": _shift_xlsx(None),
        "shift1904.xlsx": _shift_xlsx(MAC_EPOCH)})
    assert dates == {"clocks.xls": ("2024-07-16",),
                     "clocks1904.xls": ("2024-07-16",),
                     "clocks.xlsx": ("2024-07-16",),
                     "shift.xls": (), "shift1904.xls": (),
                     "shift.xlsx": (), "shift1904.xlsx": ()}, dates


def test_e2_date_formatted_number_that_is_no_date_shows_the_number():
    """Review B2: one out-of-calendar number under a date format failed the
    WHOLE ``.xls`` workbook. A number that is no date in Excel's calendar --
    past 9999-12-31, negative, NaN -- reads as stored, beside the rest of the
    workbook, with an unmarked note counting them. ``.xlsx`` did not degrade
    to the number either: openpyxl substitutes ``#VALUE!`` for the first and
    counts a negative serial back into an invented 1899 date."""
    raw = make_fixtures.xls_bytes([
        {"name": "Costs",
         "cells": [(0, 0, "COSTROW", 0), (0, 1, 3000000.0, 14), (0, 2, -5.0, 14),
                   (0, 3, float("nan"), 14), (0, 4, 2958465.0, 14),
                   (0, 5, -0.5, 14)]},   # a negative FRACTION: never a 1899 date
        {"name": "Other", "cells": [(0, 0, "THIRDSHEET", 0)]},
    ])
    xls_note = ("4 cell(s) formatted as a date or time held a number that is no date "
                "in Excel's calendar (negative, past 9999-12-31, or not a number); the "
                "number is shown as stored")
    _assert_doc(_extract("costs.xls", raw),
                ["[sheet: Costs]\nCOSTROW\t3000000\t-5\tnan\t9999-12-31\t-0.5",
                 "[sheet: Other]\nTHIRDSHEET"],
                [xls_note, PAGE_NOTE_XLS])
    assert not ex.has_evidence_marker(xls_note)
    note = xls_note.replace("4 cell(s)", "3 cell(s)")

    def build(wb):
        ws = wb.active
        ws.title = "Costs"
        ws["A1"], ws["B1"], ws["C1"], ws["D1"] = "COSTROW", 3000000, -5, -0.5
        ws["B1"].number_format = ws["C1"].number_format = "yyyy-mm-dd"
        ws["D1"].number_format = "h:mm"

    with pytest.warns(UserWarning, match="outside the limits for dates"):
        doc = _extract("costs.xlsx", _xlsx(build))
    _assert_doc(doc, ["[sheet: Costs]\nCOSTROW\t3000000\t-5\t-0.5"],
                [note, PAGE_NOTE_XLSX])


# ---------------------------------------------------------------------------
# E3 / E15: .xls booleans, errors, numbers
# ---------------------------------------------------------------------------


def _legacy_row_xls() -> bytes:
    return make_fixtures.xls_bytes([{
        "name": "Row",
        "cells": [(0, 0, "EGRET\r\nFINCH", 0), (0, 1, False, 0),
                  (0, 2, make_fixtures.XlsError(0x2A), 0), (0, 3, 5.5, 0),
                  (0, 4, 0.15, 9), (0, 5, 0.15, 10), (0, 6, "TAB\tSTOP", 0)],
    }])


def test_e3_xls_boolean_and_error_render_their_own_text():
    """E3 (confirmed): TRUE and ``#DIV/0!`` on fixture 19, and their siblings
    on one ``.xls`` row: FALSE, a second error code (``#N/A``), a fraction, a
    CRLF inside a cell (ONE pilcrow), a tab inside a cell, and -- review
    unverified item -- the ``.xls`` percentages, which read as bare decimals
    while ``.xlsx`` read them as percentages."""
    doc = _pages("19_legacy_workbook.xls")
    _assert_doc(doc, [LEGACY_PAGE], [PAGE_NOTE_XLS])
    _assert_doc(_extract("row.xls", _legacy_row_xls()),
                ["[sheet: Row]\nEGRET ¶ FINCH\tFALSE\t#N/A\t5.5\t15%\t15.00%\tTAB STOP"],
                [PAGE_NOTE_XLS])


def test_e15_xls_whole_number_reads_as_an_integer():
    """E15: a whole number reads ``5``, not ``5.0``; a fraction keeps its
    fraction; and infinity or NaN under the General format reads as the
    number -- ``int(value)`` on either raised and failed the whole workbook."""
    doc = _pages("19_legacy_workbook.xls")
    sheet = _page_with(doc.pages, "[sheet: Legacy]")
    assert sheet is not None, "no page starts with '[sheet: Legacy]'"
    row0 = next((ln for ln in sheet.text.split("\n") if "RAVENXLS" in ln), None)
    assert row0 is not None, f"no line contains 'RAVENXLS'; text was {sheet.text!r}"
    assert row0.split("\t")[-1] == "5", row0
    raw = make_fixtures.xls_bytes([{"name": "Numbers", "cells": [
        (0, 0, "PLOVER", 0), (0, 1, 5.5, 0), (0, 2, float("inf"), 0),
        (0, 3, float("-inf"), 0), (0, 4, float("nan"), 0), (0, 5, -7.0, 0)]}])
    _assert_doc(_extract("numbers.xls", raw),
                ["[sheet: Numbers]\nPLOVER\t5.5\tinf\t-inf\tnan\t-7"],
                [PAGE_NOTE_XLS])


# ---------------------------------------------------------------------------
# E4 / E5: row positions and in-cell whitespace
# ---------------------------------------------------------------------------


def test_e4_blank_rows_collapse_to_one_marked_line():
    """E4 (confirmed): rows 5 and 6 of Register are one ``[blank rows 5-6]``.

    Sibling (review unverified): a row whose cells hold only whitespace is as
    blank as an empty row. Printed, it normalized to an empty line, and every
    row number after it was lost -- in both formats. A blank row AFTER the
    last data row, whitespace-only or not, still produces nothing."""
    doc = _pages("18_workbook_constructs.xlsx")
    reg = _page_with(doc.pages, "[sheet: Register]")
    assert reg is not None, "no page starts with '[sheet: Register]'"
    _assert_found(reg.text, "\n[blank rows 5-6]\nJUNIPER\n", "the collapsed blank run. ")

    def build(wb):
        ws = wb.active
        ws.title = "Space"
        ws["A1"], ws["A2"], ws["B3"], ws["A4"], ws["A5"], ws["C7"] = (
            "SISKIN", " ", "   ", "  ", "TWITE", " ")

    expected = ["[sheet: Space]\nSISKIN\n[blank rows 2-4]\nTWITE"]
    _assert_doc(_extract("space.xlsx", _xlsx(build)), expected, [PAGE_NOTE_XLSX])
    raw = make_fixtures.xls_bytes([{"name": "Space", "cells": [
        (0, 0, "SISKIN", 0), (1, 0, " ", 0), (2, 1, "   ", 0), (4, 0, "TWITE", 0),
        (6, 2, " ", 0)]}])
    _assert_doc(_extract("space.xls", raw), expected, [PAGE_NOTE_XLS])


def test_e5_newline_and_tab_inside_a_cell_stay_on_one_line():
    """E5: a cell's own newline becomes `` ¶ `` and its tab one space, so a row
    stays one line.

    Sibling (review B6): comment text and author go through the same
    cleaning, so one comment is one line -- an ordinary Excel note starts
    with ``Author:`` and a line break, and its second line was
    indistinguishable from a print footer."""
    doc = _pages("18_workbook_constructs.xlsx")
    reg = _page_with(doc.pages, "[sheet: Register]")
    assert reg is not None, "no page starts with '[sheet: Register]'"
    lines = reg.text.split("\n")
    assert "HOLLY ¶ IVY\t=SUM(B2:B3)\t\tTAB STOP" in lines, lines

    def build(wb):
        ws = wb.active
        ws.title = "Notes"
        ws["A1"], ws["A2"] = "GROUSEROW", "JACANAROW"
        ws["A1"].comment = openpyxl.comments.Comment("GROUSE:\nHERON", "GROUSE")
        ws["A2"].comment = openpyxl.comments.Comment("IBIS\t30", "JACANA")

    _assert_doc(_extract("notes.xlsx", _xlsx(build)),
                ["[sheet: Notes]\nGROUSEROW\nJACANAROW\n"
                 "[comment on A1 by GROUSE] GROUSE: ¶ HERON\n"
                 "[comment on A2 by JACANA] IBIS 30"],
                [PAGE_NOTE_XLSX])


def test_e5_a_date_a_cell_wraps_is_still_detected(tmp_path):
    """Review B4: E5's pilcrow sat between a wrapped date's parts, so
    ``16 July`` / ``2024`` stopped being a date for detection, and the index
    and screen date changed. Detection now reads the cell's break as a break
    (``extract.date_detection_text``); the page text keeps the pilcrow; and a
    pilcrow in a document that is not a spreadsheet stays a character."""
    def build(wb):
        ws = wb.active
        ws.title = "Letters"
        ws["A1"] = "Letter dated 16 July\n2024"
        ws["A2"] = "Issued March 1,\n2025"
        ws["A3"] = "Plain 2019-02-11"

    xls = make_fixtures.xls_bytes([{"name": "Letters", "cells": [
        (0, 0, "Letter dated 16 July\n2024", 0), (1, 0, "Issued March 1,\r\n2025", 0),
        (2, 0, "Plain 2019-02-11", 0)]}])
    page = ("[sheet: Letters]\nLetter dated 16 July ¶ 2024\nIssued March 1, ¶ 2025\n"
            "Plain 2019-02-11")
    assert _texts(_extract("letters.xlsx", _xlsx(build))) == [page]
    assert _texts(_extract("letters.xls", xls)) == [page]

    # Review r2 (A3): the decision follows the READER that wrote the page,
    # never the file's name. A workbook delivered under another name and
    # recovered by content sniffing keeps its wrapped dates; a Word file
    # named .xlsx gains no date from a pilcrow typed in its text.
    import docx

    word = docx.Document()
    word.add_paragraph("Agreed 9 May ¶ 2022 GANNETWORD")
    buf = io.BytesIO()
    word.save(buf)
    wrapped = ("2024-07-16", "2025-03-01", "2019-02-11")
    dates = _detected_dates(tmp_path, {
        "letters.xlsx": _xlsx(build), "letters.xls": xls, "letters.xlsm": _xlsx(build),
        "sheet_named.pdf": _xlsx(build), "sheet_named.docx": _xlsx(build),
        "legacy_named.msg": xls, "word.docx": buf.getvalue(),
        "word_named.xlsx": buf.getvalue(),
        # Typed text that copies a spreadsheet reader's page note is still
        # typed text: the reader is read from the page's notes, not its text.
        "memo.txt": (f"{PAGE_NOTE_XLSX}\nletter of March ¶ 12, 2020\n").encode("utf-8")})
    assert dates == {"letters.xlsx": wrapped, "letters.xls": wrapped,
                     "letters.xlsm": wrapped, "sheet_named.pdf": wrapped,
                     "sheet_named.docx": wrapped, "legacy_named.msg": wrapped,
                     "word.docx": (), "word_named.xlsx": (),
                     "memo.txt": ()}, dates


# ---------------------------------------------------------------------------
# E6: every tab keeps its page, whatever its kind
# ---------------------------------------------------------------------------


def _ledger_and_plot(wb, rows: int = 2):
    ws = wb.active
    ws.title = "Ledger"
    for i in range(1, rows + 1):
        ws[f"A{i}"] = f"ROW{i}"
    ws.oddHeader.center.text = "HERONHEAD"
    plot = wb.create_chartsheet("Plot")
    chart = openpyxl.chart.BarChart()
    chart.add_data(openpyxl.chart.Reference(ws, min_col=1, min_row=1, max_row=2))
    plot.add_chart(chart)


PLOT_NOTE = (f"chartsheet 'Plot': its chart, and any print header or footer on it, "
             f"were not read ({UNREAD})")


def test_e6_chartsheet_gets_its_own_page_in_tab_order():
    """E6: a chartsheet gets its page, in tab order, reading exactly
    ``[chartsheet: <name>]``, with a marked note naming it.

    Sibling (review unverified): the KIND comes from the relationship Type,
    not from ``chartsheets/`` in the part's path. A chartsheet stored at
    ``xl/plots/plot1.xml`` failed the whole workbook."""
    doc = _pages("18_workbook_constructs.xlsx")
    first_lines = [p.text.splitlines()[0] if p.text else "" for p in doc.pages]
    assert first_lines == ["[sheet: Register]", "[chartsheet: Chart1]", "[sheet: Later]"], (
        first_lines)
    assert doc.pages[1].text == "[chartsheet: Chart1]"
    assert list(doc.notes) == [CHART1_NOTE, FORMULA_NOTE_ONE, PAGE_NOTE_XLSX], doc.notes

    raw = _xlsx(_ledger_and_plot)
    part = next(n for n in _part_names(raw) if n.startswith("xl/chartsheets/sheet"))
    rels = next(n for n in _part_names(raw) if n.startswith("xl/chartsheets/_rels/"))
    moved = _patch(raw, {
        "[Content_Types].xml": lambda t: t.replace("/" + part, "/xl/plots/plot1.xml"),
        "xl/_rels/workbook.xml.rels": lambda t: t.replace("/" + part, "/xl/plots/plot1.xml"),
    }, renames={part: "xl/plots/plot1.xml", rels: "xl/plots/_rels/plot1.xml.rels"})
    _assert_doc(_extract("moved.xlsx", moved),
                ["[sheet: Ledger]\nHERONHEAD\nROW1\nROW2", "[chartsheet: Plot]"],
                [PLOT_NOTE, PAGE_NOTE_XLSX])


def test_e6_tab_order_comes_from_the_packages_relationships():
    """Review unverified: the workbook part was read from the fixed name
    ``xl/workbook.xml``. A package whose ``_rels/.rels`` names another part
    fell back to openpyxl's worksheet list: the chartsheet's page vanished,
    and no sheet's header was read -- with no note."""
    raw = _xlsx(_ledger_and_plot)
    moved = _patch(raw, {
        "[Content_Types].xml": lambda t: t.replace("/xl/workbook.xml", "/xl/book.xml"),
        "_rels/.rels": lambda t: t.replace("xl/workbook.xml", "xl/book.xml"),
    }, renames={"xl/workbook.xml": "xl/book.xml",
                "xl/_rels/workbook.xml.rels": "xl/_rels/book.xml.rels"})
    _assert_doc(_extract("book.xlsx", moved),
                ["[sheet: Ledger]\nHERONHEAD\nROW1\nROW2", "[chartsheet: Plot]"],
                [PLOT_NOTE, PAGE_NOTE_XLSX])


def test_e6_xls_chart_and_macro_tabs_keep_their_pages():
    """Review unverified: xlrd reads worksheets only, so a chart tab and an
    Excel 4.0 macro-sheet tab vanished and every later page number shifted.
    Each keeps its page, named from the workbook's own tab records, with a
    marked note."""
    raw = make_fixtures.xls_bytes([
        {"name": "Legacy", "cells": [(0, 0, "WOODLARK", 0)]},
        {"name": "Plot", "kind": "chart"},
        {"name": "Macro1", "kind": "macro"},
        {"name": "Annex", "cells": [(0, 0, "YELLOWHAMMER", 0)]},
    ])
    _assert_doc(_extract("tabs.xls", raw),
                ["[sheet: Legacy]\nWOODLARK", "[chartsheet: Plot]",
                 "[sheet: Macro1]\n[not read: an Excel 4.0 macro sheet]",
                 "[sheet: Annex]\nYELLOWHAMMER"],
                [PLOT_NOTE,
                 f"sheet 'Macro1': an Excel 4.0 macro sheet; its content was not read ({UNREAD})",
                 PAGE_NOTE_XLS])


def _three_sheets(wb):
    ws = wb.active
    ws.title = "First"
    ws["A1"] = "LAPWING"
    wb.create_sheet("Gone")["A1"] = "GONEWORD"
    wb.create_sheet("Last")["A1"] = "LINNETWORD"


def test_e6_dialog_sheet_and_missing_sheet_part_are_disclosed_as_unread():
    """Review unverified: a sheet whose part is missing from the file had no
    page and no note, shifting every later page; a dialog sheet read as an
    empty page with no note. Each keeps its page and says why it was not
    read, with a marked note."""
    missing = _patch(_xlsx(_three_sheets), {"xl/worksheets/sheet2.xml": lambda t: None})
    _assert_doc(_extract("missing.xlsx", missing),
                ["[sheet: First]\nLAPWING",
                 "[sheet: Gone]\n[not read: its part is missing from the file]",
                 "[sheet: Last]\nLINNETWORD"],
                [f"sheet 'Gone': its part is missing from the file; its content "
                 f"was not read ({UNREAD})", PAGE_NOTE_XLSX])

    dialog_xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  '<dialogsheet xmlns="http://schemas.openxmlformats.org/'
                  'spreadsheetml/2006/main"><sheetPr/></dialogsheet>')
    dialog = _patch(_xlsx(_three_sheets), {
        "xl/worksheets/sheet2.xml": lambda t: None,
        "xl/dialogsheets/sheet1.xml": lambda t: dialog_xml,
        "xl/_rels/workbook.xml.rels": lambda t: t.replace(
            'relationships/worksheet" Target="/xl/worksheets/sheet2.xml"',
            'relationships/dialogsheet" Target="/xl/dialogsheets/sheet1.xml"'),
        "[Content_Types].xml": lambda t: t.replace(
            '"/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-'
            'officedocument.spreadsheetml.worksheet+xml"',
            '"/xl/dialogsheets/sheet1.xml" ContentType="application/vnd.openxmlformats-'
            'officedocument.spreadsheetml.dialogsheet+xml"'),
    })
    _assert_doc(_extract("dialog.xlsx", dialog),
                ["[sheet: First]\nLAPWING", "[sheet: Gone]\n[not read: a dialog sheet]",
                 "[sheet: Last]\nLINNETWORD"],
                [f"sheet 'Gone': a dialog sheet; its content was not read ({UNREAD})",
                 PAGE_NOTE_XLSX])


def test_e6_duplicate_sheet_names_each_read_their_own_sheet():
    """Review C: two tabs with one name (a non-Excel writer can produce it)
    were looked up by name, and openpyxl returns the first -- the second
    page silently showed the first sheet's content. Sheets are matched by
    their own part."""
    def build(wb):
        ws = wb.active
        ws.title = "Data"
        ws["A1"] = "FIRSTSHEETWORD"
        wb.create_sheet("Other")["A1"] = "SECONDSHEETWORD"

    raw = _patch(_xlsx(build), {"xl/workbook.xml": lambda t: t.replace(
        'name="Other"', 'name="Data"')})
    _assert_doc(_extract("dups.xlsx", raw),
                ["[sheet: Data]\nFIRSTSHEETWORD", "[sheet: Data]\nSECONDSHEETWORD"],
                [PAGE_NOTE_XLSX])


# ---------------------------------------------------------------------------
# E7: the row cap
# ---------------------------------------------------------------------------


def test_e7_row_cap_note_names_what_was_lost(monkeypatch):
    """E7, pinned exactly with the cap biting inside Register: the in-page
    truncation line; a page for EVERY tab never opened, a chartsheet keeping
    its exact ``[chartsheet: <name>]`` label (review unverified: it read
    ``[sheet: ...]`` behind the cap); ONE marked note naming the truncated
    sheet and every sheet never opened; and no formula note, because the
    formula row was discarded to the cap, never shown."""
    monkeypatch.setattr(ex, "_XLSX_MAX_ROWS", 3)
    doc = _pages("18_workbook_constructs.xlsx")
    note = ("workbook truncated at 3 rows while reading 'Register'; the rest of "
            "that sheet, and the 2 sheet(s) never opened ('Chart1', 'Later'), "
            f"were not read ({UNREAD})")
    _assert_doc(doc,
                ["[sheet: Register]\nOSPREY\nRef\tAmount\tNote\n\t10\t15%\n\t20\t15.00%\n"
                 "[... workbook truncated at 3 rows]\n[comment on B2 by NIGHTJAR] MAGPIE\n"
                 "PUFFIN",
                 f"[chartsheet: Chart1]\n{CAP_LINE}",
                 f"[sheet: Later]\n{CAP_LINE}"],
                [note, PAGE_NOTE_XLSX])
    assert ex.has_evidence_marker(note)


def test_e7_xls_row_cap_gives_every_unopened_tab_its_page(monkeypatch):
    """E7 for ``.xls``: the same pages, labels and ONE note -- a chart tab
    behind the cap included."""
    monkeypatch.setattr(ex, "_XLSX_MAX_ROWS", 3)
    raw = make_fixtures.xls_bytes([
        {"name": "Ledger", "cells": [(r, 0, f"ROW{r + 1}", 0) for r in range(5)]},
        {"name": "Plot", "kind": "chart"},
        {"name": "Annex", "cells": [(0, 0, "YELLOWHAMMER", 0)]},
    ])
    _assert_doc(_extract("cap.xls", raw),
                ["[sheet: Ledger]\nROW1\nROW2\nROW3\n[... workbook truncated at 3 rows]",
                 f"[chartsheet: Plot]\n{CAP_LINE}", f"[sheet: Annex]\n{CAP_LINE}"],
                ["workbook truncated at 3 rows while reading 'Ledger'; the rest of "
                 "that sheet, and the 2 sheet(s) never opened ('Plot', 'Annex'), "
                 f"were not read ({UNREAD})", PAGE_NOTE_XLS])


def test_e7_cap_in_the_last_sheet_says_only_its_rest_was_lost(monkeypatch):
    """Review C: with no sheet after the truncated one, the note read "no
    further sheets not read". It says what was lost: the rest of that sheet."""
    monkeypatch.setattr(ex, "_XLSX_MAX_ROWS", 2)

    def build(wb):
        ws = wb.active
        ws.title = "Log"
        for i in range(1, 5):
            ws[f"A{i}"] = f"ENTRY{i}"

    _assert_doc(_extract("log.xlsx", _xlsx(build)),
                ["[sheet: Log]\nENTRY1\nENTRY2\n[... workbook truncated at 2 rows]"],
                ["workbook truncated at 2 rows while reading 'Log'; the rest of that "
                 f"sheet was not read ({UNREAD})", PAGE_NOTE_XLSX])


# ---------------------------------------------------------------------------
# E10: percentages
# ---------------------------------------------------------------------------


def test_e10_percentage_cell_reads_as_a_percentage():
    """E10: ``0%`` and ``0.00%`` on Register.

    Siblings (review unverified): a ``%`` escaped with a backslash is literal
    text Excel does not scale (``0\\%`` holding 15 read ``1500%``); ``?`` and
    ``#`` are digit placeholders; a negative keeps its sign; and only the
    format's FIRST section decides."""
    doc = _pages("18_workbook_constructs.xlsx")
    reg = _page_with(doc.pages, "[sheet: Register]")
    assert reg is not None, "no page starts with '[sheet: Register]'"
    _assert_found(reg.text, "\n\t10\t15%\n\t20\t15.00%\n", "both percentage cells. ")

    def build(wb):
        ws = wb.active
        ws.title = "Pct"
        for row, (value, fmt) in enumerate([
                (15, "0\\%"), (0.1525, "0.0?%"), (0.1525, "0.0#%"),
                (-0.125, "0.0%;[Red]-0.0%"), (0.15, "0.00;-0.00%"),
                (15, '0"%"')], start=1):
            ws[f"A{row}"] = value
            ws[f"A{row}"].number_format = fmt

    _assert_doc(_extract("pct.xlsx", _xlsx(build)),
                ["[sheet: Pct]\n15\n15.25%\n15.25%\n-12.5%\n0.15\n15"],
                [PAGE_NOTE_XLSX])


# ---------------------------------------------------------------------------
# E11: comments
# ---------------------------------------------------------------------------


def test_e11_cell_comment_is_read_after_the_sheets_rows():
    """E11: Register's page is pinned whole -- the comment line after the rows,
    before the footer.

    Siblings: comments print in CELL order whatever order the part holds
    them, each with its own author; rich-text runs join; and (review C) one
    comment whose ref is a range no longer drops every comment on the sheet."""
    doc = _pages("18_workbook_constructs.xlsx")
    reg = _page_with(doc.pages, "[sheet: Register]")
    assert reg is not None, "no page starts with '[sheet: Register]'"
    assert reg.text == REGISTER_PAGE, reg.text

    def build(wb):
        ws = wb.active
        ws.title = "Notes"
        ws["A1"], ws["B2"], ws["D5"] = "KINGLETROW", "MARTINROW", "DUNLINROW"
        ws["A1"].comment = openpyxl.comments.Comment("PLACEHOLDERA", "AUK")
        ws["B2"].comment = openpyxl.comments.Comment("PLACEHOLDERB", "AUK")

    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    comments = (
        f'<comments xmlns="{ns}"><authors><author>AUK</author>'
        '<author>SWIFTLET</author></authors><commentList>'
        '<comment ref="D5" authorId="1"><text><r><rPr><b/></rPr><t>SWIFTLET:</t></r>'
        '<r><t xml:space="preserve">&#10;LATERNOTE</t></r></text></comment>'
        '<comment ref="B2" authorId="0"><text><t>MARTIN</t></text></comment>'
        '<comment ref="A1:B1" authorId="0"><text><t>KINGLET</t></text></comment>'
        '</commentList></comments>')
    raw = _xlsx(build)
    part = next(n for n in _part_names(raw) if re.match(r"xl/comments/comment\d+\.xml$", n))
    _assert_doc(_extract("notes.xlsx", _patch(raw, {part: lambda t: comments})),
                ["[sheet: Notes]\nKINGLETROW\n\tMARTINROW\n[blank rows 3-4]\n\t\t\tDUNLINROW\n"
                 "[comment on A1:B1 by AUK] KINGLET\n"
                 "[comment on B2 by AUK] MARTIN\n"
                 "[comment on D5 by SWIFTLET] SWIFTLET: ¶ LATERNOTE"],
                [PAGE_NOTE_XLSX])


def test_e11_threaded_comments_are_credited_to_named_people():
    """Review unverified: Excel's default comment is THREADED -- the
    conversation in ``xl/threadedComments``, its people in ``xl/persons``, and
    a placeholder note credited to ``tc={id}`` for older readers. Only the
    placeholder was read, so no named person reached the page. Each thread
    entry is its own line credited to its person; the placeholder is not
    repeated; an ordinary note beside it still prints."""
    guid = "{00000000-0000-4000-8000-000000000001}"

    def build(wb):
        ws = wb.active
        ws.title = "Ledger"
        ws["A1"], ws["A2"] = "ALDER", "REEDROW"
        ws["A1"].comment = openpyxl.comments.Comment(
            "[Threaded comment]\n\nComment:\n    MAGPIE\nReplies:\n    NUTHATCH",
            "tc=" + guid)
        ws["A2"].comment = openpyxl.comments.Comment("BITTERN", "CRAKE")

    tc_ns = "http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments"
    thread = (
        f'<ThreadedComments xmlns="{tc_ns}">'
        f'<threadedComment ref="A1" dT="2026-01-01T00:00:00.00" personId="{{P1}}" '
        f'id="{guid}"><text>MAGPIE</text></threadedComment>'
        f'<threadedComment ref="A1" dT="2026-01-02T00:00:00.00" personId="{{P2}}" '
        f'id="{{00000000-0000-4000-8000-000000000002}}" parentId="{guid}">'
        '<text>NUTHATCH</text></threadedComment>'
        f'<threadedComment ref="A1" dT="2026-01-03T00:00:00.00" personId="{{P9}}" '
        f'id="{{00000000-0000-4000-8000-000000000003}}" parentId="{guid}">'
        '<text>WAXWING</text></threadedComment></ThreadedComments>')
    persons = (f'<personList xmlns="{tc_ns}">'
               '<person displayName="NIGHTJAR" id="{P1}" userId="N" providerId="None"/>'
               '<person displayName="ORIOLE" id="{P2}" userId="O" providerId="None"/>'
               '</personList>')
    rel = ('<Relationship Id="rIdTC" Type="http://schemas.microsoft.com/office/'
           '2017/10/relationships/threadedComment" '
           'Target="../threadedComments/threadedComment1.xml"/></Relationships>')
    person_rel = ('<Relationship Id="rIdP" Type="http://schemas.microsoft.com/office/'
                  '2017/10/relationships/person" Target="persons/person.xml"/>'
                  '</Relationships>')
    parts = {
        "xl/worksheets/_rels/sheet1.xml.rels": lambda t: t.replace("</Relationships>", rel),
        "xl/_rels/workbook.xml.rels": lambda t: t.replace("</Relationships>", person_rel),
        "xl/threadedComments/threadedComment1.xml": lambda t: thread,
        "xl/persons/person.xml": lambda t: persons,
    }
    lines = ("[sheet: Ledger]\nALDER\nREEDROW\n[comment on A1 by NIGHTJAR] MAGPIE\n"
             "[comment on A1 by ORIOLE] NUTHATCH\n"
             # A person the file names nobody for is credited to the id itself
             # (review r2 mutant R08: an empty author passed every test).
             "[comment on A1 by {P9}] WAXWING\n"
             "[comment on A2 by CRAKE] BITTERN")
    _assert_doc(_extract("threads.xlsx", _patch(_xlsx(build), parts)), [lines],
                [PAGE_NOTE_XLSX])

    # Review r2 (C): the persons part missing is disclosed only where a
    # threaded comment is credited to an id because of it -- never on a
    # workbook that has no threaded comment at all.
    missing = dict(parts, **{"xl/persons/person.xml": lambda t: None})
    ids_only = lines.replace("by NIGHTJAR", "by {P1}").replace("by ORIOLE", "by {P2}")
    _assert_doc(_extract("threads.xlsx", _patch(_xlsx(build), missing)), [ids_only],
                ["the workbook's persons part is missing from the file; 3 threaded "
                 "comment(s) are credited to person ids, not names "
                 f"({UNREAD})", PAGE_NOTE_XLSX])
    no_threads = {"xl/_rels/workbook.xml.rels": parts["xl/_rels/workbook.xml.rels"]}
    _assert_doc(_extract("nothreads.xlsx", _patch(_xlsx(build), no_threads)),
                ["[sheet: Ledger]\nALDER\nREEDROW\n"
                 "[comment on A1 by tc={00000000-0000-4000-8000-000000000001}] "
                 "[Threaded comment] ¶  ¶ Comment: ¶     MAGPIE ¶ Replies: ¶     NUTHATCH\n"
                 "[comment on A2 by CRAKE] BITTERN"],
                [PAGE_NOTE_XLSX])


# ---------------------------------------------------------------------------
# E12: print header and footer
# ---------------------------------------------------------------------------


def test_e12_print_header_and_footer_bound_the_sheets_page():
    """E12: Register's header is the line right after ``[sheet: Register]``,
    exactly ``OSPREY`` (no code letter, no ``&C``), and its footer is the
    page's last line, exactly ``PUFFIN``, after the comment.

    Siblings (review B3): Excel's full code set leaves no stray letter --
    color ``&K`` (hex or theme), fonts, sizes, picture ``&G``, strikethrough
    ``&S``, super/subscript ``&X``/``&Y``, file path ``&Z`` and the rest;
    ``&&`` is ``&``; the left, center and right sections are separate lines;
    even-page and first-page variants are read; and a red Bates stamp in the
    footer's left section parses with its own prefix."""
    doc = _pages("18_workbook_constructs.xlsx")
    reg = _page_with(doc.pages, "[sheet: Register]")
    assert reg is not None, "no page starts with '[sheet: Register]'"
    lines = reg.text.split("\n")
    assert lines[1] == "OSPREY", lines
    assert lines[-1] == "PUFFIN", lines
    assert lines[-2] == "[comment on B2 by NIGHTJAR] MAGPIE", lines

    def build(wb):
        ws = wb.active
        ws.title = "Stamps"
        ws["A1"] = "STAMPROW"
        ws.oddHeader.center.text = "PLACEHOLDERHEAD"

    header_footer = (
        '<headerFooter differentOddEven="1" differentFirst="1">'
        '<oddHeader>&amp;L&amp;"Arial,Bold"&amp;KFF0000LEFTWORD&amp;CMID&amp;&amp;WORD'
        '&amp;R&amp;K04+000RIGHTWORD</oddHeader>'
        '<oddFooter>&amp;L&amp;KFF0000QXZ 000123&amp;RPage &amp;P</oddFooter>'
        '<evenHeader>&amp;C&amp;S&amp;X&amp;Y&amp;E&amp;O&amp;H&amp;G&amp;Z&amp;F'
        '&amp;A&amp;D&amp;T&amp;N&amp;B&amp;I&amp;U&amp;12EVENWORD</evenHeader>'
        '<firstFooter>&amp;CFIRSTFOOTWORD</firstFooter></headerFooter>')
    raw = _patch(_xlsx(build), {"xl/worksheets/sheet1.xml": lambda t: _set_header_footer(
        t, header_footer)})
    doc = _extract("stamps.xlsx", raw)
    # The even header's &G places a picture: its code leaves no letter, and
    # the picture is named in a FINAL note (review-fix round 3), not dropped.
    _assert_doc(doc,
                ["[sheet: Stamps]\nLEFTWORD\nMID&WORD\nRIGHTWORD\nEVENWORD\nSTAMPROW\n"
                 "QXZ 000123\nPage\nFIRSTFOOTWORD"],
                [f"sheet 'Stamps': 1 picture(s) in its print header or footer were not "
                 f"read ({UNREAD})", PAGE_NOTE_XLSX])
    parsed = bates._parse_line("QXZ 000123")
    assert (parsed.prefix, parsed.number) == ("QXZ", 123), parsed


# ---------------------------------------------------------------------------
# E13: hyperlinks
# ---------------------------------------------------------------------------


def test_e13_hyperlink_target_follows_the_cell_text():
    """E13: KESTREL's external target follows it as `` <URL>``.

    Siblings (review unverified): a link over a RANGE is one link, attached
    to the range's first cell only -- it was copied into every cell, making
    empty cells non-blank; a location fragment on an external link is kept;
    an internal link adds nothing; and a link on a cell the part holds no
    ``<c>`` for, on a row past the last one it holds, still prints in its
    column and row."""
    from openpyxl.worksheet.hyperlink import Hyperlink

    doc = _pages("18_workbook_constructs.xlsx")
    reg = _page_with(doc.pages, "[sheet: Register]")
    assert reg is not None, "no page starts with '[sheet: Register]'"
    assert "KESTREL <https://example.invalid/LARK>" in reg.text.split("\n"), reg.text

    def build(wb):
        ws = wb.active
        ws.title = "Links"
        ws["A1"] = "LINNET"
        ws.merge_cells("A1:C1")
        ws["A1"].hyperlink = "https://example.invalid/NODDY"
        ws["A3"] = "PIPIT"
        ws["A3"].hyperlink = "https://example.invalid/PIPIT.pdf"
        ws["A5"] = "WAGTAIL"
        ws["A5"].hyperlink = Hyperlink(ref="A5", location="'Links'!A1")

    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

    def links(xml: str) -> str:
        xml = re.sub(r'<hyperlink ([^>]*)ref="A1" ', r'<hyperlink \1ref="A1:C1" ', xml)
        xml = re.sub(r'<hyperlink ([^>]*)ref="A3" ', r'<hyperlink \1ref="A3" location="page=5" ',
                     xml)
        return xml.replace("</hyperlinks>", f'<hyperlink xmlns:r="{r_ns}" ref="C9" '
                           'r:id="rIdROWLESS"/></hyperlinks>')

    def rels(xml: str) -> str:
        return xml.replace("</Relationships>", (
            '<Relationship Id="rIdROWLESS" Type="http://schemas.openxmlformats.org/'
            'officeDocument/2006/relationships/hyperlink" '
            'Target="https://example.invalid/ROWLESS" TargetMode="External"/>'
            "</Relationships>"))

    raw = _patch(_xlsx(build), {"xl/worksheets/sheet1.xml": links,
                                "xl/worksheets/_rels/sheet1.xml.rels": rels})
    sheet_xml = zipfile.ZipFile(io.BytesIO(raw)).read("xl/worksheets/sheet1.xml").decode()
    assert 'ref="A1:C1"' in sheet_xml and 'ref="C9"' in sheet_xml, sheet_xml
    assert '<c r="B1"' not in sheet_xml and '<row r="9"' not in sheet_xml, sheet_xml
    _assert_doc(_extract("links.xlsx", raw),
                ["[sheet: Links]\nLINNET <https://example.invalid/NODDY>\n[blank row 2]\n"
                 "PIPIT <https://example.invalid/PIPIT.pdf#page=5>\n[blank row 4]\nWAGTAIL\n"
                 "[blank rows 6-8]\n\t\t<https://example.invalid/ROWLESS>"],
                [PAGE_NOTE_XLSX])


# ---------------------------------------------------------------------------
# Constructs outside E1-E15 that lost evidence silently
# ---------------------------------------------------------------------------


def _stale_dimension(xml: str) -> str:
    return re.sub(r'<dimension ref="[^"]*"/>', '<dimension ref="A1"/>', xml)


def test_a_stale_dimension_does_not_truncate_rows_or_columns():
    """Review unverified (predates the package): a writer can leave
    ``<dimension ref="A1"/>`` on a sheet holding far more; read-only openpyxl
    stops at it, and 11 of these 12 cells vanished with status FULL."""
    def build(wb):
        ws = wb.active
        ws.title = "Stale"
        for r in range(1, 5):
            for c in "ABC":
                ws[f"{c}{r}"] = f"VIREO{c}{r}"

    raw = _patch(_xlsx(build), {"xl/worksheets/sheet1.xml": _stale_dimension})
    _assert_doc(_extract("stale.xlsx", raw),
                ["[sheet: Stale]\n" + "\n".join(
                    "\t".join(f"VIREO{c}{r}" for c in "ABC") for r in range(1, 5))],
                [PAGE_NOTE_XLSX])


def test_hidden_rows_columns_and_sheets_are_read_and_disclosed():
    """Review unverified: hidden rows, hidden columns, hidden and very hidden
    sheets were read with nothing saying they were hidden. They are still
    read, and an unmarked note says so, in both formats."""
    def build(wb):
        ws = wb.active
        ws.title = "Shown"
        ws["A1"], ws["C1"], ws["A2"] = "GOLDCREST", "HIDDENCOLWORD", "HIDDENROWWORD"
        ws.row_dimensions[2].hidden = True
        ws.column_dimensions["C"].hidden = True
        tucked = wb.create_sheet("Tucked")
        tucked["A1"] = "TUCKEDWORD"
        tucked.sheet_state = "hidden"
        buried = wb.create_sheet("Buried")
        buried["A1"] = "BURIEDWORD"
        buried.sheet_state = "veryHidden"

    pages = ["[sheet: Shown]\nGOLDCREST\t\tHIDDENCOLWORD\nHIDDENROWWORD",
             "[sheet: Tucked]\nTUCKEDWORD", "[sheet: Buried]\nBURIEDWORD"]
    notes = ["sheet 'Shown' marks 1 row(s) and 1 column(s) hidden; hidden cells "
             "are not skipped for being hidden",
             "sheet 'Tucked' is hidden in the workbook; it was read like any other sheet",
             "sheet 'Buried' is very hidden in the workbook; it was read like any "
             "other sheet"]
    _assert_doc(_extract("hidden.xlsx", _xlsx(build)), pages, notes + [PAGE_NOTE_XLSX])
    raw = make_fixtures.xls_bytes([
        {"name": "Shown", "hidden_rows": [1], "hidden_cols": [2],
         "cells": [(0, 0, "GOLDCREST", 0), (0, 2, "HIDDENCOLWORD", 0),
                   (1, 0, "HIDDENROWWORD", 0)]},
        {"name": "Tucked", "visibility": 1, "cells": [(0, 0, "TUCKEDWORD", 0)]},
        {"name": "Buried", "visibility": 2, "cells": [(0, 0, "BURIEDWORD", 0)]},
    ])
    _assert_doc(_extract("hidden.xls", raw), pages, notes + [PAGE_NOTE_XLS])
    assert not any(ex.has_evidence_marker(n) for n in notes)

    # Review r2 (C, and mutant R04): the other ways a sheet hides cells -- a
    # hidden column RANGE (B:D is three columns), a row of height 0, a column
    # of width 0, rows hidden by default -- are counted in both formats.
    head = ('<sheetFormatPr defaultRowHeight="15" zeroHeight="1"/>'
            '<cols><col min="2" max="4" width="9" hidden="1"/>'
            '<col min="6" max="6" width="0" customWidth="1"/></cols>')
    rows = ('<row r="1">' + _is("A1", "FIRSTSEEN") + _is("F1", "NARROWWORD") + "</row>"
            '<row r="2" ht="0" customHeight="1">' + _is("A2", "FLATROWWORD") + "</row>")
    folded = ["[sheet: Folded]\nFIRSTSEEN\t\t\t\t\tNARROWWORD\nFLATROWWORD"]
    folded_note = ("sheet 'Folded' marks 1 row(s) and 4 column(s) hidden, and hides every "
                   "row not given a height of its own; hidden cells are not skipped for "
                   "being hidden")
    _assert_doc(_extract("folded.xlsx", _rows_xlsx(rows, "A1:F2", name="Folded", head=head)),
                folded, [folded_note, PAGE_NOTE_XLSX])
    raw = make_fixtures.xls_bytes([
        {"name": "Folded", "hidden_cols": [1, 2, 3, 5], "zero_height_rows": [1],
         "rows_hidden_by_default": True,
         "cells": [(0, 0, "FIRSTSEEN", 0), (0, 5, "NARROWWORD", 0),
                   (1, 0, "FLATROWWORD", 0)]}])
    _assert_doc(_extract("folded.xls", raw), folded, [folded_note, PAGE_NOTE_XLS])


# ---------------------------------------------------------------------------
# How the per-sheet extras are read
# ---------------------------------------------------------------------------


def test_sheet_extras_never_parse_the_whole_worksheet_part(monkeypatch):
    """Review B1: header, footer and hyperlinks were found by parsing the
    WHOLE worksheet part into a tree, so memory grew with the sheet however
    low the row cap. With whole-part parsing made to fail, every extra still
    reads: the part is streamed."""
    real_fromstring = ET.fromstring

    def no_whole_sheet(data, *args, **kwargs):
        if b"<sheetData" in (data if isinstance(data, bytes) else data.encode()):
            raise MemoryError("a worksheet part was parsed whole")
        return real_fromstring(data, *args, **kwargs)

    monkeypatch.setattr(ET, "fromstring", no_whole_sheet)
    doc = _pages("18_workbook_constructs.xlsx")
    _assert_doc(doc, [REGISTER_PAGE, "[chartsheet: Chart1]", "[sheet: Later]\nQUETZAL"],
                [CHART1_NOTE, FORMULA_NOTE_ONE, PAGE_NOTE_XLSX])


def test_sheet_extras_read_failure_is_a_marked_note(monkeypatch):
    """Review B1: a failure reading a sheet's extras (a real MemoryError, on a
    loaded machine) was swallowed by a bare ``except: pass`` -- header,
    footer, links and comments gone, status FULL, no note. The rows and the
    separately stored comments still read, and each failed read is a MARKED
    note naming the sheet and what was not read."""
    real_fromstring, real_iterparse = ET.fromstring, ET.iterparse

    def failing_fromstring(data, *args, **kwargs):
        if b"<sheetData" in (data if isinstance(data, bytes) else data.encode()):
            raise MemoryError()
        return real_fromstring(data, *args, **kwargs)

    def failing_iterparse(source, *args, **kwargs):
        if "worksheets/" in getattr(source, "name", ""):
            raise MemoryError()
        return real_iterparse(source, *args, **kwargs)

    monkeypatch.setattr(ET, "fromstring", failing_fromstring)
    monkeypatch.setattr(ET, "iterparse", failing_iterparse)
    doc = _pages("18_workbook_constructs.xlsx")
    extras_note = ("its print header/footer, hyperlinks, hidden rows and columns, drawings "
                   "and true extent could not be read (MemoryError), so a cell or row "
                   "written out of order, or a text box, chart or picture, may be missing "
                   f"with no note of its own ({PART_UNREAD})")
    raw_note = ("sheet 'Register': its own cell XML could not be read (MemoryError), so a "
                "formula whose stored result is empty may be shown as its formula and "
                "counted as having no stored value, a date-formatted number that is no "
                "date may show as #VALUE!, a time-only cell may show the number openpyxl "
                "converted rather than the one stored, and a date, time or duration may "
                f"be a millisecond off Excel's rounding ({PART_UNREAD})")
    _assert_doc(doc,
                ["[sheet: Register]\nRef\tAmount\tNote\n\t10\t15%\n\t20\t15.00%\n"
                 "HOLLY ¶ IVY\t=SUM(B2:B3)\t\tTAB STOP\n[blank rows 5-6]\nJUNIPER\nKESTREL\n"
                 "[comment on B2 by NIGHTJAR] MAGPIE",
                 "[chartsheet: Chart1]", "[sheet: Later]\nQUETZAL"],
                [f"sheet 'Register': {extras_note}", raw_note, CHART1_NOTE,
                 f"sheet 'Later': {extras_note}",
                 FORMULA_NOTE_ONE, PAGE_NOTE_XLSX])
    # A read that raised is TRANSIENT (review-fix round 2): the walker's serial
    # retry re-reads it, and only a failure the file itself makes permanent
    # is final.
    for read_failure in (doc.notes[0], doc.notes[1], doc.notes[3]):
        assert ex.has_transient_marker(read_failure), read_failure
        assert not ex.has_final_marker(read_failure), read_failure
    assert ex.has_final_marker(CHART1_NOTE) and not ex.has_transient_marker(CHART1_NOTE)


def test_formula_and_sheet_list_read_failures_are_marked_notes(monkeypatch):
    """The two other bare fallbacks. The second, formula-text reading of the
    workbook failing left a formula with no stored value blank, with no note.
    The package's own sheet list failing to parse fell back to openpyxl's
    worksheet list -- the chartsheet's page gone, no sheet's header read, no
    note."""
    real_load = openpyxl.load_workbook

    def no_formula_reading(*args, **kwargs):
        if kwargs.get("data_only") is False:
            raise MemoryError()
        return real_load(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", no_formula_reading)
    doc = _pages("18_workbook_constructs.xlsx")
    formula_note = ("the workbook's formulas could not be read (MemoryError); a formula "
                    "cell with no stored value shows blank and is not counted "
                    f"({PART_UNREAD})")
    assert list(doc.notes) == [formula_note, CHART1_NOTE, PAGE_NOTE_XLSX], doc.notes
    assert ex.has_transient_marker(formula_note) and not ex.has_final_marker(formula_note)
    assert "HOLLY ¶ IVY\t\t\tTAB STOP" in doc.pages[0].text.split("\n"), doc.pages[0].text
    monkeypatch.setattr(openpyxl, "load_workbook", real_load)

    real_fromstring = ET.fromstring

    def no_sheet_list(data, *args, **kwargs):
        if b"<sheets>" in (data if isinstance(data, bytes) else data.encode()):
            raise MemoryError()
        return real_fromstring(data, *args, **kwargs)

    monkeypatch.setattr(ET, "fromstring", no_sheet_list)
    doc = _pages("18_workbook_constructs.xlsx")
    list_note = ("the workbook's own sheet list could not be read (MemoryError); tabs "
                 "follow openpyxl's list, which leaves out a sheet that names no part "
                 f"or a part openpyxl does not find in the file ({PART_UNREAD})")
    # Review-fix round 2 (A5): the fallback still knows each worksheet's part,
    # so the header, links, comments and raw cells all read.
    _assert_doc(doc, [REGISTER_PAGE, "[chartsheet: Chart1]", "[sheet: Later]\nQUETZAL"],
                [list_note, CHART1_NOTE, FORMULA_NOTE_ONE, PAGE_NOTE_XLSX])
    assert ex.has_transient_marker(list_note) and not ex.has_final_marker(list_note)


# ---------------------------------------------------------------------------
# Review-fix round 2 (D-55): what the second review found, and its class
# ---------------------------------------------------------------------------


def _walk_one(tmp_path, files: dict[str, bytes], workers: int = 2):
    """A real ``walker.run`` over exactly ``files``: ``(documents by
    filename, the run's load-dependent notes)``."""
    src = tmp_path / "src"
    src.mkdir()
    for name, data in files.items():
        (src / name).write_bytes(data)
    notes = walker.RunNotes()
    config = RunConfig(source_root=str(src), output_root=str(tmp_path / "out"),
                       ocr_engine_version=ex.ocr_engine_version())
    result = walker.run(config, walker.WalkOptions(workers=workers, ocr_enabled=False,
                                                   resume=False), notes)
    return {d.filename: d for d in result.documents}, list(notes.load_dependent)


_FIXTURE18_PAGES = [REGISTER_PAGE, "[chartsheet: Chart1]", "[sheet: Later]\nQUETZAL"]


@pytest.mark.parametrize("site", [
    "sheet extras stream", "comments part", "formula reading", "sheet list",
    "raw cell stream", "xls number formats"])
def test_a_read_failure_a_second_read_clears_is_retried_by_the_walker(
        tmp_path, monkeypatch, site):
    """Review r2 (A1): a MemoryError that a second read clears was labeled
    FINAL, so the walker never retried it and the run recorded the loss --
    header, comment, formula or duration gone -- as a property of the file.
    Every read that raised is TRANSIENT now: raised ONCE on a pool thread
    (never on the serial retry), the retry re-reads the file alone and the
    record the run keeps is the whole one."""
    import xlrd

    raw18 = (FIXTURES / "18_workbook_constructs.xlsx").read_bytes()
    if site == "sheet extras stream":
        monkeypatch.setattr(ex, "_xlsx_iter_rows", _once_in_pool(ex._xlsx_iter_rows, 1))
    elif site == "raw cell stream":
        monkeypatch.setattr(ex, "_xlsx_iter_rows", _once_in_pool(ex._xlsx_iter_rows, 2))
    elif site == "comments part":
        monkeypatch.setattr(ex, "_xlsx_legacy_comments",
                            _once_in_pool(ex._xlsx_legacy_comments))
    elif site == "formula reading":
        monkeypatch.setattr(openpyxl, "load_workbook", _once_in_pool(
            openpyxl.load_workbook, when=lambda *a, **k: k.get("data_only") is False))
    elif site == "sheet list":
        monkeypatch.setattr(ex, "_xlsx_sheet_order", _once_in_pool(ex._xlsx_sheet_order))
    else:
        monkeypatch.setattr(xlrd, "open_workbook", _once_in_pool(
            xlrd.open_workbook, when=lambda *a, **k: k.get("formatting_info")))

    if site == "xls number formats":
        docs, load_notes = _walk_one(tmp_path, {"clocks.xls": _clock_xls()})
        doc = docs["clocks.xls"]
        _assert_doc(doc, ["[sheet: Clocks]\nGEARWHEEL\t08:24:00\t36:00:00\t00:00:00\t18:00:00"
                          "\t1470:00:00\t2024-07-16 12:00:00"], [PAGE_NOTE_XLS])
    else:
        docs, load_notes = _walk_one(tmp_path, {"18_workbook_constructs.xlsx": raw18})
        doc = docs["18_workbook_constructs.xlsx"]
        _assert_doc(doc, _FIXTURE18_PAGES, [CHART1_NOTE, FORMULA_NOTE_ONE, PAGE_NOTE_XLSX])
    assert len(load_notes) == 1 and "RESOLVED" in load_notes[0], load_notes


def _threads_workbook(parts_edit=None) -> bytes:
    """A workbook with a legacy comment, a threaded comment, a hyperlink and
    a print header: every per-sheet read the extras make."""
    guid = "{00000000-0000-4000-8000-000000000001}"

    def build(wb):
        ws = wb.active
        ws.title = "Ledger"
        ws["A1"], ws["A2"] = "ALDER", "REEDROW"
        ws["A1"].comment = openpyxl.comments.Comment("placeholder", "tc=" + guid)
        ws["A2"].comment = openpyxl.comments.Comment("BITTERN", "CRAKE")
        ws["A2"].hyperlink = "https://example.invalid/REEDLINK"
        ws.oddHeader.center.text = "HERONHEAD"

    tc_ns = "http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments"
    parts = {
        "xl/worksheets/_rels/sheet1.xml.rels": lambda t: t.replace(
            "</Relationships>",
            '<Relationship Id="rIdTC" Type="http://schemas.microsoft.com/office/'
            '2017/10/relationships/threadedComment" '
            'Target="../threadedComments/threadedComment1.xml"/></Relationships>'),
        "xl/_rels/workbook.xml.rels": lambda t: t.replace(
            "</Relationships>",
            '<Relationship Id="rIdP" Type="http://schemas.microsoft.com/office/'
            '2017/10/relationships/person" Target="persons/person.xml"/></Relationships>'),
        "xl/threadedComments/threadedComment1.xml": lambda t: (
            f'<ThreadedComments xmlns="{tc_ns}"><threadedComment ref="A1" '
            f'personId="{{P1}}" id="{guid}"><text>MAGPIE</text></threadedComment>'
            "</ThreadedComments>"),
        "xl/persons/person.xml": lambda t: (
            f'<personList xmlns="{tc_ns}"><person displayName="NIGHTJAR" id="{{P1}}" '
            'userId="N" providerId="None"/></personList>'),
    }
    parts.update(parts_edit or {})
    return _patch(_xlsx(build), parts)


THREADS_PAGE = ("[sheet: Ledger]\nHERONHEAD\nALDER\nREEDROW <https://example.invalid/REEDLINK>\n"
                "[comment on A1 by NIGHTJAR] MAGPIE\n[comment on A2 by CRAKE] BITTERN")


def test_every_read_that_raised_is_a_transient_note_saying_what_it_cost(monkeypatch):
    """Review r2 (A1), the class: EACH exception path inside the spreadsheet
    readers -- relationships, threaded comments, legacy comments, persons,
    the rows part-way through, the formula rows part-way through, a legacy
    workbook's tab records and a sheet's own records -- is one TRANSIENT note
    that says what the failure cost, never FINAL, and the rest still reads."""
    raw = _threads_workbook()
    _assert_doc(_extract("threads.xlsx", raw), [THREADS_PAGE], [PAGE_NOTE_XLSX])

    def failing(fn, when=lambda *a, **k: True):
        def wrapper(*args, **kwargs):
            if when(*args, **kwargs):
                raise MemoryError()
            return fn(*args, **kwargs)
        return wrapper

    def check(expected_page, expected_note, name="threads.xlsx", data=raw, page_note=PAGE_NOTE_XLSX):
        doc = _extract(name, data)
        _assert_doc(doc, [expected_page], [expected_note, page_note])
        assert ex.has_transient_marker(expected_note), expected_note
        assert not ex.has_final_marker(expected_note), expected_note

    with monkeypatch.context() as m:
        m.setattr(ex, "_xlsx_rels", failing(ex._xlsx_rels, lambda z, part: "worksheets" in part))
        check("[sheet: Ledger]\nHERONHEAD\nALDER\nREEDROW",
              "sheet 'Ledger': its relationships, so its hyperlinks, comments and drawings, "
              f"could not be read (MemoryError) ({PART_UNREAD})")
    with monkeypatch.context() as m:
        m.setattr(ex, "_xlsx_threaded_comments", failing(ex._xlsx_threaded_comments))
        check("[sheet: Ledger]\nHERONHEAD\nALDER\nREEDROW <https://example.invalid/REEDLINK>\n"
              "[comment on A1 by tc={00000000-0000-4000-8000-000000000001}] placeholder\n"
              "[comment on A2 by CRAKE] BITTERN",
              f"sheet 'Ledger': its threaded comments could not be read (MemoryError) "
              f"({PART_UNREAD})")
    with monkeypatch.context() as m:
        m.setattr(ex, "_xlsx_legacy_comments", failing(ex._xlsx_legacy_comments))
        check("[sheet: Ledger]\nHERONHEAD\nALDER\nREEDROW <https://example.invalid/REEDLINK>\n"
              "[comment on A1 by NIGHTJAR] MAGPIE",
              f"sheet 'Ledger': its comments could not be read (MemoryError) ({PART_UNREAD})")
    with monkeypatch.context() as m:
        m.setattr(ex, "_xlsx_persons", failing(ex._xlsx_persons))
        check(THREADS_PAGE.replace("by NIGHTJAR", "by {P1}"),
              "the workbook's persons part could not be read (MemoryError); 1 threaded "
              f"comment(s) are credited to person ids, not names ({PART_UNREAD})")

    # What the FILE makes permanent stays FINAL: a relationship naming a
    # comments or threaded-comments part the package does not hold.
    for missing, page, note in (
            ("xl/threadedComments/threadedComment1.xml",
             "[sheet: Ledger]\nHERONHEAD\nALDER\nREEDROW <https://example.invalid/REEDLINK>\n"
             "[comment on A1 by tc={00000000-0000-4000-8000-000000000001}] placeholder\n"
             "[comment on A2 by CRAKE] BITTERN",
             "sheet 'Ledger': its threaded comments part is missing from the file; its "
             f"threaded comments were not read ({UNREAD})"),
            (next(n for n in _part_names(raw) if re.match(r"xl/comments/comment\d+\.xml$", n)),
             "[sheet: Ledger]\nHERONHEAD\nALDER\nREEDROW <https://example.invalid/REEDLINK>\n"
             "[comment on A1 by NIGHTJAR] MAGPIE",
             "sheet 'Ledger': its comments part is missing from the file; its comments were "
             f"not read ({UNREAD})")):
        doc = _extract("threads.xlsx", _patch(raw, {missing: lambda t: None}))
        _assert_doc(doc, [page], [note, PAGE_NOTE_XLSX])
        assert ex.has_final_marker(note) and not ex.has_transient_marker(note), note

    # openpyxl's own row reading failing part-way: rows before it are kept.
    from openpyxl.worksheet import _reader

    real_parse_row = _reader.WorkSheetParser.parse_row

    def parse_row_failing_at_two(self, row):
        if row.get("r") == "2":
            raise MemoryError()
        return real_parse_row(self, row)

    with monkeypatch.context() as m:
        m.setattr(_reader.WorkSheetParser, "parse_row", parse_row_failing_at_two)
        check("[sheet: Ledger]\nHERONHEAD\nALDER\n[comment on A1 by NIGHTJAR] MAGPIE\n"
              "[comment on A2 by CRAKE] BITTERN",
              "sheet 'Ledger': its rows from row 2 on could not be read (MemoryError), so "
              f"they are not on its page ({PART_UNREAD})")

    def build_formulas(wb):
        ws = wb.active
        ws.title = "Sums"
        ws["A1"], ws["A2"], ws["A3"] = "=1+1", "LABELROW", "=2+2"

    real_formula_parse_row = _reader.WorkSheetParser.parse_row

    def formula_rows_failing_at_three(self, row):
        if not self.data_only and row.get("r") == "3":
            raise MemoryError()
        return real_formula_parse_row(self, row)

    with monkeypatch.context() as m:
        m.setattr(_reader.WorkSheetParser, "parse_row", formula_rows_failing_at_three)
        doc = _extract("sums.xlsx", _xlsx(build_formulas))
        formula_note = ("sheet 'Sums': its formulas from row 3 on could not be read "
                        "(MemoryError), so a formula cell there with no stored value "
                        f"shows blank and is not counted ({PART_UNREAD})")
        _assert_doc(doc, ["[sheet: Sums]\n=1+1\nLABELROW"],
                    [formula_note, FORMULA_NOTE_ONE, PAGE_NOTE_XLSX])
        assert ex.has_transient_marker(formula_note) and not ex.has_final_marker(formula_note)

    xls = make_fixtures.xls_bytes([
        {"name": "Legacy", "cells": [(0, 0, "WOODLARK", 0)], "header": "&CHEADWORD"},
        {"name": "Plot", "kind": "chart"}])
    with monkeypatch.context() as m:
        m.setattr(ex, "_xls_sheet_records", failing(ex._xls_sheet_records))
        doc = _extract("tabs.xls", xls)
        records_note = ("sheet 'Legacy': its own records could not be read (MemoryError): "
                        "its print header and footer and dialog flag, and so a cell written "
                        "twice, a comment missing its text, and a text box, shape, chart, "
                        "picture or control, may be lost with no note of its own "
                        f"({PART_UNREAD})")
        _assert_doc(doc, ["[sheet: Legacy]\nWOODLARK", "[chartsheet: Plot]"],
                    [records_note, PLOT_NOTE, PAGE_NOTE_XLS])
        assert ex.has_transient_marker(records_note)
    with monkeypatch.context() as m:
        m.setattr(ex, "_xls_stream", failing(ex._xls_stream))
        doc = _extract("tabs.xls", xls)
        tabs_note = ("the workbook's own tab records could not be read (MemoryError); a "
                     "tab that is not a worksheet is named by its position, and no sheet's "
                     "own records were read: its print header and footer and dialog flag, "
                     "and so a cell written twice, a comment missing its text, and a text "
                     "box, shape, chart, picture or control, may be lost with no note of "
                     f"its own ({PART_UNREAD})")
        _assert_doc(doc, ["[sheet: Legacy]\nWOODLARK",
                          "[sheet: tab 2]\n[not read: a sheet of an unrecognized kind]"],
                    [tabs_note, "sheet 'tab 2': a sheet of an unrecognized kind; its content "
                     f"was not read ({UNREAD})", PAGE_NOTE_XLS])
        assert ex.has_transient_marker(tabs_note) and not ex.has_final_marker(tabs_note)


def test_a_cell_or_row_written_out_of_order_is_read_or_named():
    """Review r2 (A2): round 1's ``reset_dimensions()`` made openpyxl size
    each row by its LAST-written cell, so a cell written before a later
    column (C1, then A1) vanished with status FULL. Both readings are now
    given the sheet's real width, measured by the extras stream -- with a
    correct, a stale or no ``<dimension>``.

    Siblings (review r2, C): a row numbered at or below one written before
    it, and a second row with the same number, are skipped by openpyxl; the
    blank-row marker then claimed such a row was blank. It is named in a
    marked note and no marker covers it. A cell written twice in one row
    keeps only its later value, also named."""
    order = ('<row r="1">' + _is("C1", "ZEBRAWORD") + _is("A1", "ANTWORD") + "</row>"
             '<row r="2">' + _is("A2", "BEEWORD") + _is("B2", "COWWORD") + "</row>"
             '<row r="3">' + _is("D3", "DEERWORD") + _is("A3", "ELKWORD") + "</row>")
    expected = ["[sheet: Order]\nANTWORD\t\tZEBRAWORD\nBEEWORD\tCOWWORD\nELKWORD\t\t\tDEERWORD"]
    for dimension in ("A1:D3", "A1", None):
        _assert_doc(_extract("order.xlsx", _rows_xlsx(order, dimension)), expected,
                    [PAGE_NOTE_XLSX])

    rows = ('<row r="3">' + _is("A3", "THIRDROWWORD") + "</row>"
            '<row r="1">' + _is("A1", "FIRSTROWWORD") + "</row>"
            '<row r="5">' + _is("A5", "FIFTHROWWORD") + "</row>"
            '<row r="5">' + _is("A5", "TWINROWWORD") + "</row>"
            '<row r="6">' + _is("A6", "OLDCELLWORD") + _is("A6", "NEWCELLWORD") + "</row>")
    skipped = (f"sheet 'Order': 2 row(s) numbered at or below a row written before them "
               f"(row(s) 1, 5) were skipped by the reader; their content was not read "
               f"({UNREAD})")
    rewritten = ("sheet 'Order': 1 cell(s) were written a second time in the same row "
                 "and column, and only the later value of each is read; the earlier "
                 f"value was not read ({UNREAD})")
    _assert_doc(_extract("rows.xlsx", _rows_xlsx(rows, "A1:A6")),
                ["[sheet: Order]\n[blank row 2]\nTHIRDROWWORD\n[blank row 4]\nFIFTHROWWORD\n"
                 "NEWCELLWORD"],
                [skipped, rewritten, PAGE_NOTE_XLSX])
    assert ex.has_final_marker(skipped) and ex.has_final_marker(rewritten)


def _twin_parts_xlsx() -> bytes:
    """The ``.xlsx`` half of the ``.xls`` parity pair: a legacy comment, an
    external link with a location, a file link, an internal link, a range
    link, a link on a row past the last cell, and odd, even and first-page
    headers and footers in all three sections."""
    from openpyxl.worksheet.hyperlink import Hyperlink

    def build(wb):
        ws = wb.active
        ws.title = "Ledger"
        ws["A1"], ws["B1"] = "GANNETROW", 7
        ws["B1"].comment = openpyxl.comments.Comment("PELICANNOTE", "CURLEW")
        ws["A2"] = "SHAGROW"
        ws["A2"].hyperlink = Hyperlink(ref="A2", target="https://example.invalid/CORMORANT",
                                       location="page=5")
        ws["A3"] = "TERNROW"
        ws["A3"].hyperlink = "annex\\BOOBYFILE.pdf"
        ws["A4"] = "SKUAROW"
        ws["A4"].hyperlink = Hyperlink(ref="A4", location="'Ledger'!A1")
        ws["A5"] = "PETRELROW"
        ws["A5"].hyperlink = "https://example.invalid/RANGELINK"
        ws["A6"] = "FULMARROW"
        ws["A6"].hyperlink = "https://example.invalid/PASTROW"
        ws.oddHeader.left.text = "LEFTHEADWORD"
        ws.oddHeader.center.text = "HERONHEAD"
        ws.oddHeader.right.text = "QXZ 000321"
        ws.oddFooter.center.text = "EGRETFOOT"
        ws.evenHeader.center.text = "EVENHEADWORD"
        ws.evenFooter.center.text = "EVENFOOTWORD"
        ws.firstHeader.center.text = "FIRSTHEADWORD"
        ws.firstFooter.left.text = "FIRSTFOOTWORD"

    def sheet(xml: str) -> str:
        xml = re.sub(r'<hyperlink ([^>]*)ref="A5"', r'<hyperlink \1ref="A5:B6"', xml)
        # The link on A6 moves to C8: a row the part holds no cell for.
        xml = re.sub(r'<hyperlink ([^>]*)ref="A6"', r'<hyperlink \1ref="C8"', xml)
        return re.sub(r'<c r="A6"[^>]*>.*?</c>', "", xml)

    return _patch(_xlsx(build), {"xl/worksheets/sheet1.xml": sheet})


TWIN_PAGE = ("[sheet: Ledger]\nLEFTHEADWORD\nHERONHEAD\nQXZ 000321\nEVENHEADWORD\n"
             "FIRSTHEADWORD\nGANNETROW\t7\nSHAGROW <https://example.invalid/CORMORANT#page=5>\n"
             "TERNROW <annex\\BOOBYFILE.pdf>\nSKUAROW\nPETRELROW <https://example.invalid/RANGELINK>\n"
             "[blank rows 6-7]\n\t\t<https://example.invalid/PASTROW>\n"
             "[comment on B1 by CURLEW] PELICANNOTE\nEGRETFOOT\nEVENFOOTWORD\nFIRSTFOOTWORD")


def test_xls_comments_links_and_print_header_footer_read_as_xlsx_renders_them():
    """Review r2 (A4, a round-1 item never closed): ``.xls`` cell comments,
    hyperlinks and the print header and footer were dropped with status FULL
    and no note, although xlrd hands over the comment and link maps and the
    HEADER/FOOTER records sit in each sheet's own substream. The ``.xls``
    reads exactly as its ``.xlsx`` twin -- the record layout checked against
    a workbook real Excel 2016 saved both ways (scratchpad
    ``excel_fix2/probe_real_parity.py``)."""
    _assert_doc(_extract("twin.xlsx", _twin_parts_xlsx()), [TWIN_PAGE], [PAGE_NOTE_XLSX])
    raw = make_fixtures.xls_bytes([{
        "name": "Ledger",
        "cells": [(0, 0, "GANNETROW", 0), (0, 1, 7.0, 0), (1, 0, "SHAGROW", 0),
                  (2, 0, "TERNROW", 0), (3, 0, "SKUAROW", 0), (4, 0, "PETRELROW", 0)],
        "header": "&LLEFTHEADWORD&CHERONHEAD&RQXZ 000321",
        "footer": "&CEGRETFOOT",
        "even_first": ("&CEVENHEADWORD", "&CEVENFOOTWORD", "&CFIRSTHEADWORD",
                       "&LFIRSTFOOTWORD"),
        "notes": [(0, 1, "CURLEW", "PELICANNOTE")],
        "links": [(1, 1, 0, 0, "url", "https://example.invalid/CORMORANT", "page=5"),
                  (2, 2, 0, 0, "file", "annex\\BOOBYFILE.pdf", None),
                  (3, 3, 0, 0, "internal", None, "'Ledger'!A1"),
                  (4, 5, 0, 1, "url", "https://example.invalid/RANGELINK", None),
                  (7, 7, 2, 2, "url", "https://example.invalid/PASTROW", None)],
    }])
    _assert_doc(_extract("twin.xls", raw), [TWIN_PAGE], [PAGE_NOTE_XLS])


def test_each_format_discloses_the_siblings_the_other_already_did():
    """The parity table's gaps, both ways. ``.xls``: a dialog sheet (a
    worksheet substream Excel flags in its WSBOOL record) read as an empty
    page with no note; a NOTE record whose text records are missing, a link
    of a kind xlrd does not recognize, and a cell record written twice for
    one cell all vanished without a word. Cell records out of order, a stale
    DIMENSIONS record and two tabs of one name read as they should.
    ``.xlsx``: an Excel 4.0 macro sheet reads as the worksheet it is, and a
    sheet of a kind no reader knows is disclosed."""
    rewritten = ("sheet 'Ledger': 1 cell(s) were written a second time in the same row "
                 "and column, and only the later value of each is read; the earlier value "
                 f"was not read ({UNREAD})")
    raw = make_fixtures.xls_bytes([
        {"name": "Ledger", "dimensions": (0, 1, 0, 1), "dialog": False,
         "cells": [(2, 0, "THIRDROWWORD", 0), (0, 2, "EASTCELLWORD", 0),
                   (0, 0, "OLDCELLWORD", 0), (0, 0, "WOODLARK", 0), (1, 0, "YELLOWROW", 0)],
         "notes_without_text": [(0, 0, "CURLEW", "LOSTNOTE")],
         "links": [(1, 1, 0, 0, "unknown", None, None)]},
        {"name": "Dialog1", "dialog": True, "cells": [(0, 0, "CAPTIONWORD", 0)]},
        {"name": "Ledger", "cells": [(0, 0, "TWINTABWORD", 0)]},
        {"name": "Module1", "kind": "vbmodule"},
    ])
    _assert_doc(_extract("dialog.xls", raw),
                ["[sheet: Ledger]\nWOODLARK\t\tEASTCELLWORD\nYELLOWROW\nTHIRDROWWORD",
                 "[sheet: Dialog1]\n[not read: a dialog sheet]",
                 "[sheet: Ledger]\nTWINTABWORD",
                 "[sheet: Module1]\n[not read: a Visual Basic module sheet]"],
                [f"sheet 'Ledger': 1 comment(s) whose text records are missing or "
                 f"malformed were not read ({UNREAD})",
                 f"sheet 'Ledger': 1 hyperlink(s) of a kind this reader does not "
                 f"recognize were not read ({UNREAD})",
                 rewritten,
                 f"sheet 'Dialog1': a dialog sheet; its content was not read ({UNREAD})",
                 f"sheet 'Module1': a Visual Basic module sheet; its content was not read "
                 f"({UNREAD})",
                 PAGE_NOTE_XLS])

    macro_type = "http://schemas.microsoft.com/office/2006/relationships/xlMacrosheet"
    odd_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oddsheet"
    raw = _patch(_xlsx(_three_sheets), {
        "xl/_rels/workbook.xml.rels": lambda t: re.sub(
            r'Type="[^"]*" Target="/xl/worksheets/sheet3.xml"',
            f'Type="{odd_type}" Target="/xl/worksheets/sheet3.xml"', re.sub(
                r'Type="[^"]*" Target="/xl/worksheets/sheet2.xml"',
                f'Type="{macro_type}" Target="/xl/worksheets/sheet2.xml"', t))})
    _assert_doc(_extract("kinds.xlsm", raw),
                ["[sheet: First]\nLAPWING", "[sheet: Gone]\nGONEWORD",
                 "[sheet: Last]\n[not read: a sheet of a kind this reader does not recognize]"],
                [f"sheet 'Last': a sheet of a kind this reader does not recognize; its "
                 f"content was not read ({UNREAD})", PAGE_NOTE_XLSX])


def _tally_xlsx(rels_edit=None) -> bytes:
    """A stored empty-string formula (B1), a number no date can hold under a
    date format (C1), and a formula with no stored value (D1)."""
    def build(wb):
        ws = wb.active
        ws.title = "Tally"
        ws["A1"], ws["B1"], ws["C1"], ws["D1"], ws["A2"] = (
            "OSPREYROW", "PLACEHOLDERB", 3000000, "PLACEHOLDERD", "PLOVERROW")
        ws["C1"].number_format = "yyyy-mm-dd"

    def cells(xml: str) -> str:
        xml = re.sub(r'<c r="B1"[^>]*>.*?</c>',
                     '<c r="B1" t="str"><f>IF(A1="OSPREYROW","",A1)</f><v></v></c>', xml)
        return re.sub(r'<c r="D1"[^>]*>.*?</c>', '<c r="D1"><f>1+1</f></c>', xml)

    edits = {"xl/worksheets/sheet1.xml": cells}
    edits.update(rels_edit or {})
    return _patch(_xlsx(build), edits)


TALLY_PAGE = "[sheet: Tally]\nOSPREYROW\t\t3000000\t=1+1\nPLOVERROW"
NOT_A_DATE_ONE = ("1 cell(s) formatted as a date or time held a number that is no date in "
                  "Excel's calendar (negative, past 9999-12-31, or not a number); the "
                  "number is shown as stored")


def test_the_sheet_list_fallback_still_reads_stored_values(monkeypatch):
    """Review r2 (A5): when the workbook's own sheet list could not be read,
    every sheet's part was forgotten, so the raw-cell lookup could not see a
    stored empty string (a false FINAL 'no stored value' note) or a stored
    number openpyxl turned into ``#VALUE!`` (shown with nothing said). The
    fallback now takes each part from openpyxl's own path; the page and the
    notes are the intact package's, plus one transient note.

    Siblings (review r2, C): a package with no ``_rels/.rels`` (openpyxl
    finds the workbook through ``[Content_Types].xml``, and so does this),
    and one whose package relationship names the workbook in another CASE
    (OPC part names compare case-insensitively) need no fallback at all. A
    SHEET relationship in another case is a part openpyxl will not open: the
    note says exactly that, where it said the part was missing."""
    intact = [FORMULA_NOTE_ONE, NOT_A_DATE_ONE, PAGE_NOTE_XLSX]
    with pytest.warns(UserWarning, match="outside the limits for dates"):
        _assert_doc(_extract("tally.xlsx", _tally_xlsx()), [TALLY_PAGE], intact)
    with pytest.warns(UserWarning, match="outside the limits for dates"):
        _assert_doc(_extract("tally.xlsx", _tally_xlsx({"_rels/.rels": lambda t: None})),
                    [TALLY_PAGE], intact)
    with pytest.warns(UserWarning, match="outside the limits for dates"):
        _assert_doc(_extract("tally.xlsx", _tally_xlsx({
            "_rels/.rels": lambda t: t.replace("xl/workbook.xml", "xl/Workbook.xml")})),
                    [TALLY_PAGE], intact)
    _assert_doc(_extract("tally.xlsx", _tally_xlsx({
        "xl/_rels/workbook.xml.rels": lambda t: t.replace(
            "/xl/worksheets/sheet1.xml", "/xl/worksheets/Sheet1.xml")})),
                ["[sheet: Tally]\n[not read: openpyxl could not open it]"],
                [f"sheet 'Tally': openpyxl could not open it; its content was not read "
                 f"({UNREAD})", PAGE_NOTE_XLSX])

    real_fromstring = ET.fromstring

    def no_sheet_list(data, *args, **kwargs):
        if b"<sheets>" in (data if isinstance(data, bytes) else data.encode()):
            raise MemoryError()
        return real_fromstring(data, *args, **kwargs)

    monkeypatch.setattr(ET, "fromstring", no_sheet_list)
    list_note = ("the workbook's own sheet list could not be read (MemoryError); tabs "
                 "follow openpyxl's list, which leaves out a sheet that names no part "
                 f"or a part openpyxl does not find in the file ({PART_UNREAD})")
    with pytest.warns(UserWarning, match="outside the limits for dates"):
        _assert_doc(_extract("tally.xlsx", _tally_xlsx()), [TALLY_PAGE],
                    [list_note] + intact)


def test_xls_formats_unread_never_render_a_date_the_formats_cannot_vouch_for(tmp_path):
    """Review r2 (A6): when xlrd's formatting-aware open fails (here a
    PALETTE record whose color count disagrees with its size) and the plain
    open succeeds, formats are unknown -- and a duration read as an invented
    1900 date that reached detected dates, under an unmarked note. With no
    format to vouch for it, every date/time cell reads as the number stored,
    and the note is marked transient and counts them."""
    import struct

    palette = [(0x0092, struct.pack("<H", 56) + b"\0" * (4 * 55))]
    raw = make_fixtures.xls_bytes([{"name": "Clocks", "cells": [
        (0, 0, "SHIFTXLS", 0), (0, 1, 1.5, 46), (0, 2, 0.15, 9), (0, 3, 45489.0, 14),
        (0, 4, 0.5, 20)]}], globals_extra=palette)
    note = ("the workbook's number formats could not be read (XLRDError); 3 cell(s) "
            "whose format marks them a date or time are shown as the number stored, no "
            "duration or percentage format is applied to any other number, and hidden "
            f"rows and columns are not counted ({PART_UNREAD})")
    doc = _extract("clocks.xls", raw)
    _assert_doc(doc, ["[sheet: Clocks]\nSHIFTXLS\t1.5\t0.15\t45489\t0.5"], [note, PAGE_NOTE_XLS])
    assert ex.has_transient_marker(note) and not ex.has_final_marker(note)
    assert walker._dated(doc.pages) == ((), ()), walker._dated(doc.pages)

    # Review r3 (C): the count said "formatted as a date, time or duration"
    # and left out a bare [h] cell, which xlrd types a plain number. The note
    # counts what xlrd types a date and names every other number's formats.
    raw = make_fixtures.xls_bytes([{"name": "Spans", "cells": [
        (0, 0, "GOLDCRESTXLS", 0), (0, 1, 1.5, 164), (0, 2, 1.5, 165), (0, 3, 0.75, 166)]}],
        formats={164: "[h]:mm:ss", 165: "[h]", 166: "[mm]:ss"}, globals_extra=palette)
    note = ("the workbook's number formats could not be read (XLRDError); 2 cell(s) "
            "whose format marks them a date or time are shown as the number stored, no "
            "duration or percentage format is applied to any other number, and hidden "
            f"rows and columns are not counted ({PART_UNREAD})")
    _assert_doc(_extract("spans.xls", raw), ["[sheet: Spans]\nGOLDCRESTXLS\t1.5\t1.5\t0.75"],
                [note, PAGE_NOTE_XLS])


def test_hidden_notes_say_read_only_for_what_was_read(monkeypatch):
    """Review r2 (A7): the unmarked hidden-sheet note said "it was read like
    any other sheet" for a hidden chartsheet, a hidden sheet whose part is
    missing, hidden tabs behind the row cap and a hidden ``.xls`` macro
    sheet -- beside the marked note saying each was NOT read. It says the
    sheet was read only when it was; the hidden-cells note says hidden cells
    are not skipped, which stays true when the cap cuts a hidden row."""
    def build(wb):
        _ledger_and_plot(wb)
        wb["Plot"].sheet_state = "hidden"
        wb.create_sheet("Gone")["A1"] = "GONEWORD"
        wb["Gone"].sheet_state = "hidden"

    raw = _patch(_xlsx(build), {"xl/worksheets/sheet2.xml": lambda t: None})
    _assert_doc(_extract("hidden.xlsx", raw),
                ["[sheet: Ledger]\nHERONHEAD\nROW1\nROW2", "[chartsheet: Plot]",
                 "[sheet: Gone]\n[not read: its part is missing from the file]"],
                ["sheet 'Plot' is hidden in the workbook", PLOT_NOTE,
                 "sheet 'Gone' is hidden in the workbook",
                 f"sheet 'Gone': its part is missing from the file; its content was not "
                 f"read ({UNREAD})", PAGE_NOTE_XLSX])

    monkeypatch.setattr(ex, "_XLSX_MAX_ROWS", 2)

    def capped(wb):
        ws = wb.active
        ws.title = "Rows"
        for i in range(1, 6):
            ws[f"A{i}"] = f"RWORD{i}"
        ws.row_dimensions[5].hidden = True
        wb.create_sheet("HidA")["A1"] = "HIDAWORD"
        wb["HidA"].sheet_state = "hidden"
        wb.create_sheet("VeryB")["A1"] = "VERYBWORD"
        wb["VeryB"].sheet_state = "veryHidden"

    truncation = ("workbook truncated at 2 rows while reading 'Rows'; the rest of that "
                  "sheet, and the 2 sheet(s) never opened ('HidA', 'VeryB'), were not read "
                  f"({UNREAD})")
    _assert_doc(_extract("capped.xlsx", _xlsx(capped)),
                ["[sheet: Rows]\nRWORD1\nRWORD2\n[... workbook truncated at 2 rows]",
                 f"[sheet: HidA]\n{CAP_LINE}", f"[sheet: VeryB]\n{CAP_LINE}"],
                ["sheet 'Rows' marks 1 row(s) and 0 column(s) hidden; hidden cells are not "
                 "skipped for being hidden",
                 "sheet 'HidA' is hidden in the workbook",
                 "sheet 'VeryB' is very hidden in the workbook", truncation, PAGE_NOTE_XLSX])

    xls = make_fixtures.xls_bytes([
        {"name": "Ledger", "cells": [(r, 0, f"ROW{r + 1}", 0) for r in range(3)]},
        {"name": "Macro1", "kind": "macro", "visibility": 1},
        {"name": "Tucked", "visibility": 1, "cells": [(0, 0, "TUCKEDWORD", 0)]},
    ])
    _assert_doc(_extract("capped.xls", xls),
                ["[sheet: Ledger]\nROW1\nROW2\n[... workbook truncated at 2 rows]",
                 f"[sheet: Macro1]\n{CAP_LINE}", f"[sheet: Tucked]\n{CAP_LINE}"],
                ["sheet 'Macro1' is hidden in the workbook",
                 "sheet 'Tucked' is hidden in the workbook",
                 "workbook truncated at 2 rows while reading 'Ledger'; the rest of that "
                 "sheet, and the 2 sheet(s) never opened ('Macro1', 'Tucked'), were not "
                 f"read ({UNREAD})", PAGE_NOTE_XLS])
    monkeypatch.setattr(ex, "_XLSX_MAX_ROWS", 50000)
    _assert_doc(_extract("macro.xls", xls),
                ["[sheet: Ledger]\nROW1\nROW2\nROW3",
                 "[sheet: Macro1]\n[not read: an Excel 4.0 macro sheet]",
                 "[sheet: Tucked]\nTUCKEDWORD"],
                ["sheet 'Macro1' is hidden in the workbook",
                 f"sheet 'Macro1': an Excel 4.0 macro sheet; its content was not read ({UNREAD})",
                 "sheet 'Tucked' is hidden in the workbook; it was read like any other sheet",
                 PAGE_NOTE_XLS])


def test_the_streamed_sheet_read_holds_memory_flat():
    """Review r2 (A9): review round 1's B1 fix streams the worksheet part and
    clears each element, so memory holds one row whatever the sheet's size
    (the whole-part parse peaked at 2.3 GB). No test held it: a reader that
    never cleared passed every gate while its peak grew about 300 times.
    Traced peak around the extras stream (which now also measures the sheet's
    extent) and the raw-cell stream read to its last row stays under a small
    absolute bound, and stays FLAT against the sheet's size (review r3 C): a
    reader that clears each row but not ``sheetData``'s hold on it still
    passes a single-size bound, while its peak grows about 80 bytes a row --
    so the peak at 100,000 rows may exceed the peak at 25,000 by under
    0.5 MB."""
    import tracemalloc

    def traced_peak(n_rows: int) -> int:
        rows = "".join(f'<row r="{r}"><c r="A{r}"><v>{r}</v></c>'
                       f'<c r="B{r}" t="inlineStr"><is><t>ROWTEXT{r}</t></is></c></row>'
                       for r in range(1, n_rows + 1))
        sheet = (f'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                 f"<sheetData>{rows}</sheetData><headerFooter><oddHeader>&amp;CBIGHEADWORD"
                 f"</oddHeader></headerFooter></worksheet>").encode()
        del rows
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("xl/worksheets/sheet1.xml", sheet)
        del sheet
        with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as z:
            ex._xlsx_sheet_extras(z, "xl/worksheets/sheet1.xml", {})  # warm-up
            raw_cells = ex._XlsxRawCells(z, "xl/worksheets/sheet1.xml")
            tracemalloc.start()
            try:
                extras = ex._xlsx_sheet_extras(z, "xl/worksheets/sheet1.xml", {})
                last = raw_cells.get(n_rows, 1)
                _current, peak = tracemalloc.get_traced_memory()
            finally:
                tracemalloc.stop()
                raw_cells.close()
        assert extras.header_lines == ["BIGHEADWORD"] and extras.max_col == 2, extras
        assert last is not None and last.v_text == str(n_rows), last
        return peak

    small, large = traced_peak(25_000), traced_peak(100_000)
    assert small < 5 * 1024 * 1024, f"peak traced memory {small / 1e6:.2f} MB on 25,000 rows"
    assert large - small < 500_000, (
        f"peak traced memory grew with the sheet: {small / 1e6:.2f} MB on 25,000 rows, "
        f"{large / 1e6:.2f} MB on 100,000")


def test_xlsx_and_xls_render_booleans_whole_numbers_and_elapsed_time_alike():
    """Review r2 (C) and the format-rendering table: ``.xlsx`` printed a
    boolean ``True``/``False`` where ``.xls`` prints Excel's ``TRUE``/``FALSE``
    (E3 is exact); a whole number a writer stored as ``10.0`` read ``10.0``
    in ``.xlsx`` and ``10`` in ``.xls`` (E15); and an elapsed-seconds format
    (``[s]``) read as a duration in ``.xlsx`` and a bare number in ``.xls``."""
    def build(wb):
        ws = wb.active
        ws.title = "Flags"
        ws["A1"], ws["B1"], ws["C1"], ws["D1"] = True, False, "PLACEHOLDERC", 1.5
        ws["D1"].number_format = "[s]"
        ws["E1"], ws["F1"], ws["G1"] = "=1=1", "PLACEHOLDERF", "PLACEHOLDERG"
        ws["H1"], ws["I1"] = 15, 0.1525
        ws["H1"].number_format = "0\\%"
        ws["I1"].number_format = "0.0?%"

    def cells(xml: str) -> str:
        xml = re.sub(r'<c r="C1"[^>]*>.*?</c>', '<c r="C1"><v>10.0</v></c>', xml)
        xml = re.sub(r'<c r="F1"[^>]*>.*?</c>', '<c r="F1" t="e"><v>#N/A</v></c>', xml)
        # A STORED #VALUE! error is the error, never a number recovered for it.
        xml = re.sub(r'<c r="G1"[^>]*>.*?</c>', '<c r="G1" t="e"><v>#VALUE!</v></c>', xml)
        return re.sub(r'<c r="E1"[^>]*>.*?</c>', '<c r="E1" t="b"><f>1=1</f><v>1</v></c>', xml)

    page = "[sheet: Flags]\nTRUE\tFALSE\t10\t36:00:00\tTRUE\t#N/A\t#VALUE!\t15\t15.25%"
    _assert_doc(_extract("flags.xlsx", _patch(_xlsx(build), {"xl/worksheets/sheet1.xml": cells})),
                [page], [PAGE_NOTE_XLSX])
    raw = make_fixtures.xls_bytes([{"name": "Flags", "cells": [
        (0, 0, True, 0), (0, 1, False, 0), (0, 2, 10.0, 0), (0, 3, 1.5, 164), (0, 4, True, 0),
        (0, 5, make_fixtures.XlsError(0x2A), 0), (0, 6, make_fixtures.XlsError(0x0F), 0),
        (0, 7, 15.0, 165), (0, 8, 0.1525, 166)]}],
        formats={164: "[s]", 165: "0\\%", 166: "0.0?%"})
    _assert_doc(_extract("flags.xls", raw), [page], [PAGE_NOTE_XLS])


def test_a_cell_openpyxl_cannot_parse_stops_only_its_own_sheet():
    """The format table's NaN row: a ``<v>NaN</v>`` (no Excel writes one)
    makes openpyxl's row reader raise part-way through a sheet, and the whole
    workbook FAILED. The rows before it, and every other sheet, still read;
    the rest of that sheet is a transient note."""
    def build(wb):
        ws = wb.active
        ws.title = "Odd"
        ws["A1"], ws["A2"], ws["A3"] = "BEFOREWORD", 0.5, "AFTERWORD"
        wb.create_sheet("Next")["A1"] = "NEXTWORD"

    raw = _patch(_xlsx(build), {"xl/worksheets/sheet1.xml": lambda t: t.replace(
        "<v>0.5</v>", "<v>NaN</v>")})
    note = ("sheet 'Odd': its rows from row 2 on could not be read (ValueError), so "
            f"they are not on its page ({PART_UNREAD})")
    _assert_doc(_extract("odd.xlsx", raw), ["[sheet: Odd]\nBEFOREWORD", "[sheet: Next]\nNEXTWORD"],
                [note, PAGE_NOTE_XLSX])


def test_extras_failure_part_way_keeps_nothing_the_note_calls_unread(monkeypatch):
    """Review r2 (C, mutant R06): a failure raised part-way through the
    extras stream -- after the hyperlinks were read -- must not keep the
    half-read links, header or hidden counts while the note says they could
    not be read."""
    real = ex._xlsx_iter_rows

    def failing_after_hyperlinks(z, part):
        for row_no, el, path in real(z, part):
            yield row_no, el, path
            if ex._local(el.tag) == "hyperlinks":
                raise MemoryError()

    raw = _threads_workbook()
    calls = {"n": 0}

    def first_call_fails(z, part):
        calls["n"] += 1
        return failing_after_hyperlinks(z, part) if calls["n"] == 1 else real(z, part)

    monkeypatch.setattr(ex, "_xlsx_iter_rows", first_call_fails)
    note = ("sheet 'Ledger': its print header/footer, hyperlinks, hidden rows and columns, "
            "drawings and true extent could not be read (MemoryError), so a cell or row "
            "written out of order, or a text box, chart or picture, may be missing with no "
            f"note of its own ({PART_UNREAD})")
    _assert_doc(_extract("threads.xlsx", raw),
                ["[sheet: Ledger]\nALDER\nREEDROW\n[comment on A1 by NIGHTJAR] MAGPIE\n"
                 "[comment on A2 by CRAKE] BITTERN"],
                [note, PAGE_NOTE_XLSX])


def test_the_accounting_line_does_not_say_the_bytes_are_missing():
    """Review r2 (C, round 1's C2 still open): a chartsheet's chart or rows past
    the row cap are FINAL gaps whose bytes ARE in the file; the accounting
    line said they were "NOT in the corpus"."""
    from dociq.contracts import DocumentRecord, PageKind, PageRecord, RunResult
    from dociq.verify import accounting

    lost = DocumentRecord(
        doc_id="", rel_path="book.xlsx", filename="book.xlsx", sha256="3" * 64,
        size_bytes=1, ext=".xlsx", status=ProcessingStatus.FULL,
        pages=(PageRecord(page_no=1, text="[chartsheet: Chart1]", kind=PageKind.SYNTHETIC),),
        notes=(CHART1_NOTE,))
    report = accounting.check(RunResult(config=RunConfig(source_root="s", output_root="o"),
                                        documents=(lost,)))
    assert report.evidence_line == (
        "EVIDENCE GAPS — 1 document(s) carry a disclosed evidence gap that re-reading "
        "will not change"), report.evidence_line

    # Review r3 (C): the line said the gap was "content that was NOT read",
    # false for a FINAL marker that loses nothing -- section recognition
    # failing keeps every page. The line is true for every FINAL marker.
    kept = DocumentRecord(
        doc_id="", rel_path="memo.pdf", filename="memo.pdf", sha256="4" * 64,
        size_bytes=1, ext=".pdf", status=ProcessingStatus.FULL,
        pages=(PageRecord(page_no=1, text="MEMOWORD", kind=PageKind.NATIVE),),
        notes=(f"{ex.M_SECTIONS}: IndexError; every page is kept",))
    report = accounting.check(RunResult(config=RunConfig(source_root="s", output_root="o"),
                                        documents=(kept,)))
    assert report.documents_evidence_lost == 1, report
    assert report.evidence_line == (
        "EVIDENCE GAPS — 1 document(s) carry a disclosed evidence gap that re-reading "
        "will not change"), report.evidence_line
    assert "NOT read" not in report.render(), report.render()


# ---------------------------------------------------------------------------
# Derived class test: no cell's own literal text may be lost, whatever else
# happens to it.
# ---------------------------------------------------------------------------


def _assert_every_string_survives(name: str, raw: bytes) -> int:
    """Every string in ``xl/sharedStrings.xml`` (each ``si``'s text, its runs
    joined) and every inline-string cell of every WORKSHEET part appears,
    E5-transformed and normalized with dociq's own ``normalize``, in the
    extracted text. Returns how many strings were checked."""
    doc = _extract(name, raw)
    full_text = normalize("\n".join(p.text for p in doc.pages))
    strings: list[str] = []
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        names = z.namelist()
        if "xl/sharedStrings.xml" in names:
            sst = z.read("xl/sharedStrings.xml").decode("utf-8")
            for si in re.findall(r"<si>(.*?)</si>", sst, re.S):
                runs = re.findall(r"<t[^>]*>(.*?)</t>", si, re.S)
                strings.append(html.unescape("".join(runs)))
        for part in names:
            if not re.match(r"xl/worksheets/sheet\d+\.xml$", part):
                continue
            xml = z.read(part).decode("utf-8")
            for is_block in re.findall(r"<is>(.*?)</is>", xml, re.S):
                runs = re.findall(r"<t[^>]*>(.*?)</t>", is_block, re.S)
                strings.append(html.unescape("".join(runs)))
    assert strings, f"expected at least one shared or inline string in {name}"
    for s in strings:
        # E5 is an exact, disclosed transformation of a cell's own line breaks
        # and tabs, not a loss, so each string is looked for through it.
        e5 = (s.replace("\r\n", "\n").replace("\r", "\n")
               .replace("\n", " ¶ ").replace("\t", " "))
        needle = normalize(e5)
        pos = full_text.find(needle)
        assert pos != -1, (
            f"{name}: cell text {s!r} (E5-transformed {e5!r}, normalized "
            f"{needle!r}) not found anywhere in the extracted text (str.find -> {pos})")
    return len(strings)


def test_every_shared_or_inline_string_cell_survives_extraction():
    """No shared or inline string cell's text is lost, in fixture 18 (whose
    strings openpyxl writes inline) AND in a workbook holding its strings in
    ``xl/sharedStrings.xml`` the way Excel writes them -- a branch fixture 18
    never reached (review C). That workbook also declares a stale
    ``<dimension>``, the case where strings vanished."""
    path = FIXTURES / "18_workbook_constructs.xlsx"
    assert _assert_every_string_survives(path.name, path.read_bytes()) >= 7

    def build(wb):
        ws = wb.active
        ws.title = "Shared"
        ws["A1"] = "PLACEHOLDER"

    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    sst = (f'<sst xmlns="{ns}" count="3" uniqueCount="3">'
           '<si><t>SHAREDWORD</t></si>'
           '<si><r><t>RUNPART</t></r><r><t>JOINED</t></r></si>'
           '<si><t xml:space="preserve">LINEONE&#10;LINETWO</t></si></sst>')
    sheet_data = ('<sheetData><row r="1"><c r="A1" t="s"><v>0</v></c>'
                  '<c r="B1" t="s"><v>1</v></c></row>'
                  '<row r="3"><c r="C3" t="s"><v>2</v></c></row></sheetData>')
    raw = _patch(_xlsx(build), {
        "xl/sharedStrings.xml": lambda t: sst,
        "xl/worksheets/sheet1.xml": lambda t: _stale_dimension(
            re.sub(r"<sheetData>.*?</sheetData>", sheet_data, t, flags=re.S)),
        "xl/_rels/workbook.xml.rels": lambda t: t.replace(
            "</Relationships>",
            '<Relationship Id="rIdSST" Type="http://schemas.openxmlformats.org/'
            'officeDocument/2006/relationships/sharedStrings" '
            'Target="sharedStrings.xml"/></Relationships>'),
        "[Content_Types].xml": lambda t: t.replace(
            "</Types>",
            '<Override PartName="/xl/sharedStrings.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>'),
    })
    assert _assert_every_string_survives("shared.xlsx", raw) == 3


# ---------------------------------------------------------------------------
# Review-fix round 3 (D-55): a number at every magnitude, Excel's rounding,
# Custom Views, and what a worksheet shows beside its cells.
# ---------------------------------------------------------------------------


def _one_sheet_both(name: str, rows: list[list[tuple[float, str]]],
                    customs: dict[str, int] | None = None) -> dict[str, bytes]:
    """The same cells as an ``.xlsx`` and an ``.xls``
    (:func:`make_fixtures.xls_bytes`): row ``n`` is ``R<n>`` then each
    ``(value, number format)``. ``customs`` maps a format string to the
    ``.xls`` format index it is declared under (built-ins need none).

    Each ``.xlsx`` ``<v>`` holds the value's shortest round-trip spelling, as
    Excel writes it: openpyxl writes 16 significant digits, which would make
    ``1.1270833333333334`` a different double."""
    builtin = {"General": 0, "0": 1, "0.00": 2, "0%": 9, "0.00%": 10, "mm-dd-yy": 14,
               "h:mm": 20, "[h]:mm:ss": 46}
    index = dict(builtin, **(customs or {}))
    stand_in = {}

    def build(wb):
        ws = wb.active
        ws.title = name
        for r, row in enumerate(rows, start=1):
            ws.cell(r, 1, f"R{r}")
            for c, (value, fmt) in enumerate(row, start=2):
                placeholder = 7_000_000 + len(stand_in)
                stand_in[f"<v>{placeholder}</v>"] = f"<v>{value!r}</v>"
                ws.cell(r, c, placeholder).number_format = fmt

    def exact(xml: str) -> str:
        return re.sub(r"<v>7\d{6}</v>", lambda m: stand_in[m.group(0)], xml)

    raw = _xlsx(build)
    cells = []
    for r, row in enumerate(rows):
        cells.append((r, 0, f"R{r + 1}", 0))
        cells.extend((r, c, value, index[fmt]) for c, (value, fmt) in enumerate(row, start=1))
    formats = {i: f for f, i in (customs or {}).items()}
    return {"xlsx": _patch(raw, {"xl/worksheets/sheet1.xml": exact}),
            "xls": make_fixtures.xls_bytes([{"name": name, "cells": cells}], formats=formats)}


def test_a_number_reads_as_stored_at_every_magnitude():
    """Review r3 (B1): a whole-valued double printed as ``str(int(value))``
    spelled out its binary expansion -- ``1.23456789012346E+18`` read
    ``1234567890123460096`` and ``1E+30`` a 31-digit number, digits the file
    does not hold and Excel never shows -- and a percentage of one
    (``123456789012346011648%``, ``1E+308`` as ``inf%``) likewise, in both
    formats, status FULL, no note.

    The class (the format table, excel_fix3/formats_fixed.md): at 1E+15,
    2**53, 1.23456789012346E+18, 6.02E+23, 1E+30, 1E+308, 5E-324, 1E-20,
    negative zero, the largest and smallest integers Excel stores (15 digits,
    the typed extremes, an RK record's 30 bits) and a 17-digit time, every
    format family shows the stored value exactly -- its shortest spelling --
    or, where a format rounds, Excel's 15 significant digits. Never an
    invented digit, never ``-0``."""
    big = "0" * 295
    table = [  # stored, General and 0, 0%, date, time only, duration
        (1e15, "1000000000000000", "100000000000000000%", None, None, None),
        (2.0 ** 53, "9007199254740992", "900719925474099000%", None, None, None),
        (1.23456789012346e18, "1.23456789012346e+18", "123456789012346000000%", None, None, None),
        (6.02e23, "6.02e+23", "602" + "0" * 23 + "%", None, None, None),
        (1e30, "1e+30", "1" + "0" * 32 + "%", None, None, None),
        (1e308, "1e+308", "1" + "0" * 310 + "%", None, None, None),
        (5e-324, "5e-324", "0%", "00:00:00", "00:00:00", "0:00:00"),
        (1e-20, "1e-20", "0%", "00:00:00", "00:00:00", "0:00:00"),
        (-0.0, "0", "0%", "00:00:00", "00:00:00", "0:00:00"),
        (999999999999999.0, "999999999999999", "99999999999999900%", None, None, None),
        (-999999999999999.0, "-999999999999999", "-99999999999999900%", None, None, None),
        (9.99999999999999e307, "9.99999999999999e+307", "999999999999999" + big + "%",
         None, None, None),
        (-9.99999999999999e307, "-9.99999999999999e+307", "-999999999999999" + big + "%",
         None, None, None),
        (536870911.0, "536870911", "53687091100%", None, None, "12884901864:00:00"),
        (-536870912.0, "-536870912", "-53687091200%", None, None, "-12884901888:00:00"),
        (1.1270833333333334, "1.1270833333333334", "113%", "1900-01-01 03:03:00", None,
         "27:03:00"),
    ]
    rows, lines = [], []
    for value, stored, pct, date, clock, span in table:
        shown = [stored, stored, pct, date or stored, clock or stored, span or stored]
        rows.append([(value, f) for f in ("General", "0", "0%", "mm-dd-yy", "h:mm",
                                          "[h]:mm:ss")])
        lines.append("\t".join([f"R{len(lines) + 1}"] + shown))
    page = "[sheet: Sizes]\n" + "\n".join(lines)
    notes = [_not_a_date_note_text(25),
             "10 cell(s) formatted as a time with no date held a value of a day or more; "
             "the number is shown as stored"]
    books = _one_sheet_both("Sizes", rows)
    with pytest.warns(UserWarning, match="outside the limits for dates"):
        _assert_doc(_extract("sizes.xlsx", books["xlsx"]), [page], notes + [PAGE_NOTE_XLSX])
    _assert_doc(_extract("sizes.xls", books["xls"]), [page], notes + [PAGE_NOTE_XLS])


def _not_a_date_note_text(count: int) -> str:
    return (f"{count} cell(s) formatted as a date or time held a number that is no date in "
            "Excel's calendar (negative, past 9999-12-31, or not a number); the number is "
            "shown as stored")


# What real Excel 16 displayed for each (value, format): its own Range.Text,
# recorded through COM (review_excel_r3/verify_2/v01_excel_display.tsv and
# excel_fix3/excel_real/x01_numbers.tsv). Invented numbers only.
EXCEL_SHOWN = [
    (0.125, "0%", "13%"), (0.145, "0%", "15%"), (0.005, "0%", "1%"), (0.285, "0%", "29%"),
    (-0.125, "0%", "-13%"), (-0.001, "0%", "0%"), (0.0125, "0.0%", "1.3%"),
    (0.375, "0%", "38%"), (0.625, "0%", "63%"), (0.875, "0%", "88%"), (0.025, "0%", "3%"),
    (0.135, "0%", "14%"), (0.115, "0%", "12%"), (0.155, "0%", "16%"), (0.165, "0%", "17%"),
    (0.335, "0.0%", "33.5%"), (0.00125, "0.00%", "0.13%"), (0.12, "0%", "12%"),
    (0.1249, "0%", "12%"), (0.1251, "0%", "13%"), (0.15, "0%", "15%"),
    (0.15, "0.00%", "15.00%"), (-0.004, "0%", "0%"), (-1e-05, "0.0%", "0.0%"),
    (0.5, "0%", "50%"), (1.005, "0%", "101%"),
    (1e15, "0%", "100000000000000000%"), (2.0 ** 53, "0%", "900719925474099000%"),
    (1.23456789012346e18, "0.00%", "123456789012346000000.00%"),
    (999999999999999.0, "0%", "99999999999999900%"),
    (100000000000000.5, "0%", "10000000000000000%"),
    (-2.5e20, "0%", "-25000000000000000000000%"),
    (1234567890123450.0, "0.00%", "123456789012345000.00%"),
    (123456789012345678.0, "0%", "12345678901234600000%"),
    (0.000123456789012345678, "0.00%", "0.01%"), (45489.354166666664, "0.00%", "4548935.42%"),
    (0.123456789012345678, "0.0000000000000000000%", "12.3456789012346000000%"),
    (-1e-20, "0%", "0%"), (-0.004, "0.0%", "-0.4%"), (0.30000000000000004, "0%", "30%"),
    (5e-324, "0.00%", "0.00%"),
    (0.00146484375, "hh:mm:ss.000", "00:02:06.563"),
    (0.00341796875, "hh:mm:ss.000", "00:04:55.313"),
    (0.00048828125, "hh:mm:ss.000", "00:00:42.188"),
    (0.00244140625, "hh:mm:ss.000", "00:03:30.938"),
    (4.6875e-07, "hh:mm:ss.000", "00:00:00.041"), (1.5625e-07, "hh:mm:ss.000", "00:00:00.014"),
    (0.123456789, "hh:mm:ss.000", "02:57:46.667"), (0.99999999999, "hh:mm:ss.000", "00:00:00.000"),
    (1.00146484375, "[h]:mm:ss.000", "24:02:06.563"),
    (0.00341796875, "[h]:mm:ss.000", "0:04:55.313"),
    (2.00048828125, "[h]:mm:ss.000", "48:00:42.188"),
    (45489.00146484375, "[h]:mm:ss.000", "1091736:02:06.562"),
    (1.1270833333333334, "[h]:mm:ss.000", "27:03:00.000"),
    (45489.00146484375, "yyyy-mm-dd hh:mm:ss.000", "2024-07-16 00:02:06.562"),
    (45489.00341796875, "yyyy-mm-dd hh:mm:ss.000", "2024-07-16 00:04:55.312"),
    (45489.00048828125, "yyyy-mm-dd hh:mm:ss.000", "2024-07-16 00:00:42.187"),
    (100.00146484375, "yyyy-mm-dd hh:mm:ss.000", "1900-04-09 00:02:06.563"),
    (45489.354166666664, "yyyy-mm-dd hh:mm:ss.000", "2024-07-16 08:30:00.000"),
]


def _as_the_page_spells(fmt: str, excel_text: str) -> str:
    """Real Excel's display of a cell, in the page's own spelling: a
    percentage as it is; a time or date-time as ISO with its millisecond as
    microseconds (none when it is zero, no time at all at midnight); a
    duration's millisecond only when there is one."""
    if fmt.endswith("%"):
        return excel_text
    head, ms = excel_text.rsplit(".", 1)
    if fmt.startswith("[h]"):
        return head if ms == "000" else excel_text
    if fmt.startswith("yyyy") and head.endswith(" 00:00:00") and ms == "000":
        return head[:-9]
    return head if ms == "000" else f"{excel_text}000"


def test_percentages_times_and_durations_round_as_real_excel_shows_them():
    """Review r3 (B3): a percentage was rounded half to even on the binary
    product -- 0.125 under ``0%`` read ``12%`` where Excel shows ``13%``,
    -0.001 read ``-0%`` -- 14 of 26 real-Excel cells disagreed, in both
    formats. The class: EVERY place the package rounds a number to fewer
    digits goes through Excel's one rule (15 significant digits, a binary tie
    at the 16th toward zero; then half away from zero; never a negative
    zero): a percentage's decimals, and a time's, date-time's or duration's
    millisecond, which rounded half to even too (0.00146484375 of a day read
    ``00:02:06.562``; Excel shows ``.563``). Every cell here is one real Excel
    displayed, and both formats must show exactly that."""
    customs = {}
    for _v, fmt, _t in EXCEL_SHOWN:
        if fmt not in ("0%", "0.00%") and fmt not in customs:
            customs[fmt] = 164 + len(customs)
    rows = [[(value, fmt)] for value, fmt, _t in EXCEL_SHOWN]
    page = "[sheet: Shown]\n" + "\n".join(
        f"R{n}\t{_as_the_page_spells(fmt, text)}"
        for n, (_v, fmt, text) in enumerate(EXCEL_SHOWN, start=1))
    books = _one_sheet_both("Shown", rows, customs)
    _assert_doc(_extract("shown.xlsx", books["xlsx"]), [page], [PAGE_NOTE_XLSX])
    _assert_doc(_extract("shown.xls", books["xls"]), [page], [PAGE_NOTE_XLS])


# The records real Excel 16 wrote for a Custom View saved with print settings
# and no header or footer (review_excel_r3/verify_1/real/vB_late.xls, the
# worksheet substream from USERSVIEWBEGIN to USERSVIEWEND, byte for byte).
_EXCEL_VIEW_BLOCK = bytes.fromhex(
    "aa0140000c6419fefb79414faeba12099ea8c464010000006400000040000000030000003c0000200000300000001900"
    "00000000000000000000000000000000ffffffff1d000f00030000000000000100000000000000140000001500000083"
    "000200000084000200000026000800666666666666e63f27000800666666666666e63f28000800000000000000e83f29"
    "000800000000000000e83fa100220000000000010001000100040000000064333333333333d33f333333333333d33f3c"
    "009c0826009c08000000000000000000000c6419fefb79414faeba12099ea8c4643c330000000000000000ab01020001"
    "00")


def _xls_record(opcode: int, data: bytes) -> bytes:
    import struct

    return struct.pack("<HH", opcode, len(data)) + data


def _xls_hf(opcode: int, text: str) -> bytes:
    """A HEADER or FOOTER record as Excel writes it: empty, or a 16-bit
    character count, a flags byte and the characters."""
    import struct

    return _xls_record(opcode, struct.pack("<HB", len(text), 0) + text.encode("latin-1")
                       if text else b"")


def _xls_view_block(header: str, footer: str, inside: bytes = b"") -> bytes:
    """Excel's Custom View records (:data:`_EXCEL_VIEW_BLOCK`) holding the
    view's own ``header`` and ``footer``, and ``inside`` after them."""
    empty = _xls_hf(0x0014, "") + _xls_hf(0x0015, "")
    assert _EXCEL_VIEW_BLOCK.count(empty) == 1
    return _EXCEL_VIEW_BLOCK.replace(
        empty, _xls_hf(0x0014, header) + _xls_hf(0x0015, footer) + inside)


def _xls_shape(ot: int, obj_id: int, text: str = "", extra: bytes = b"") -> bytes:
    """A drawing object as Excel writes one: an OBJ record (ftCmo with the
    object's type and id, ``extra`` subrecords, ftEnd) and, for ``text``, a
    TXO record with the text and its formatting runs in CONTINUE records."""
    import struct

    obj = _xls_record(0x005D, struct.pack("<HHHHH", 0x15, 0x12, ot, obj_id, 0x6011)
                      + b"\0" * 12 + extra + b"\0" * 4)
    if not text:
        return obj
    encoded = text.encode("latin-1")
    return (obj + _xls_record(0x00EC, bytes.fromhex("00000df000000000"))
            + _xls_record(0x01B6, struct.pack("<HH6sHHH", 0x0212, 0, b"\0" * 6,
                                              len(text), 16, 0) + b"\0\0")
            + _xls_record(0x003C, b"\x00" + encoded)
            + _xls_record(0x003C, struct.pack("<HH4xHH4x", 0, 0, len(text), 0)))


_PICTURE = bytes.fromhex("07000200ffff080002000000")          # ftCf, ftPioGrbit
_EMBEDDED = bytes.fromhex("07000200ffff0800020000000900040000000000")  # + ftPictFmla


def _viewed_xlsx(own: str, view: str, view_extra: str = "", parts=None) -> bytes:
    """The reviewer's real-Excel Custom View workbook as an ``.xlsx``: sheet
    ``Sheet9`` holding ``BODYTOKEN``, its own ``headerFooter`` markup
    ``own``, and a Custom View (the element real Excel wrote, less its
    printer-settings reference) holding ``view`` and ``view_extra``."""
    def build(wb):
        ws = wb.active
        ws.title = "Sheet9"
        ws["A1"] = "BODYTOKEN"

    views = ('<customSheetViews><customSheetView guid="{A3A9A64A-183C-48B5-B352-33619E42677F}" '
             'showPageBreaks="1"><pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" '
             f'header="0.3" footer="0.3"/>{view_extra}{view}</customSheetView></customSheetViews>')

    def sheet(xml: str) -> str:
        xml = re.sub(r"<headerFooter>.*?</headerFooter>|<headerFooter/>", "", xml, flags=re.S)
        xml = xml.replace("<pageMargins", views + "<pageMargins", 1)
        return xml.replace("</worksheet>", own + "</worksheet>")

    return _patch(_xlsx(build), dict({"xl/worksheets/sheet1.xml": sheet}, **(parts or {})))


def test_a_custom_views_print_header_is_never_the_sheets():
    """Review r3 (B2): a Custom View keeps its own copy of the print header
    and footer. ``.xls`` read whichever HEADER/FOOTER record came last, and
    Excel writes the view's after the sheet's, so the view's stale text
    replaced the sheet's or erased it; ``.xlsx`` took a header inside
    ``customSheetView`` as the sheet's when the sheet had none. Pinned on the
    reviewer's three real-Excel pairs rebuilt byte-equivalently (the view's
    records and element as Excel 16 wrote them), against what Excel itself
    reads back: header changed after the view (BRAVO, not ALPHA), set after
    it (CHARLIE, not the view's empty one), cleared after it (none, not
    DELTA).

    The class: everything else the readers take from a sheet that a view
    could also hold is taken from the sheet's own state only. ``.xls``:
    every record in the view's block is skipped -- a dialog flag, a text box,
    a background picture put there are not the sheet's -- and a block that
    never ends is named, not swallowed. ``.xlsx``: a hidden column, rows
    hidden by default, a drawing and a header picture inside the view are
    not the sheet's."""
    import struct

    body = [(0, 0, "BODYTOKEN", 0)]
    cases = [
        ("changed", "&CBRAVOHEAD", "&CBRAVOFOOT", "&CALPHAHEAD", "&CALPHAFOOT",
         "[sheet: Sheet9]\nBRAVOHEAD\nBODYTOKEN\nBRAVOFOOT"),
        ("late", "&CCHARLIEHEAD", "&CCHARLIEFOOT", "", "",
         "[sheet: Sheet9]\nCHARLIEHEAD\nBODYTOKEN\nCHARLIEFOOT"),
        ("cleared", "", "", "&CDELTAHEAD", "&CDELTAFOOT", "[sheet: Sheet9]\nBODYTOKEN"),
    ]
    for label, head, foot, view_head, view_foot, page in cases:
        xls = make_fixtures.xls_bytes([{
            "name": "Sheet9", "cells": body,
            "records": [_xls_hf(0x0014, head), _xls_hf(0x0015, foot),
                        _xls_view_block(view_head, view_foot)]}])
        _assert_doc(_extract(f"{label}.xls", xls), [page], [PAGE_NOTE_XLS])

        def hf(h, f):
            inner = "".join(f"<{t}>{html.escape(v)}</{t}>"
                            for t, v in (("oddHeader", h), ("oddFooter", f)) if v)
            return f"<headerFooter>{inner}</headerFooter>" if inner else ""

        _assert_doc(_extract(f"{label}.xlsx", _viewed_xlsx(hf(head, foot), hf(view_head, view_foot))),
                    [page], [PAGE_NOTE_XLSX])

    # .xls: what a view's block holds is never the sheet's, whatever it is.
    inside = (_xls_record(0x0081, struct.pack("<BB", 0xD1, 0x04))
              + _xls_shape(0x06, 40, "VIEWBOXWORD") + _xls_record(0x00E9, b"\0" * 8))
    xls = make_fixtures.xls_bytes([{
        "name": "Sheet9", "cells": body,
        "records": [_xls_hf(0x0014, "&CBRAVOHEAD"),
                    _xls_view_block("&C&GVIEWHEADWORD", "", inside)]}])
    _assert_doc(_extract("inside.xls", xls), ["[sheet: Sheet9]\nBRAVOHEAD\nBODYTOKEN"],
                [PAGE_NOTE_XLS])
    begin_only = _EXCEL_VIEW_BLOCK[:68]      # USERSVIEWBEGIN alone: 4 + 64 bytes
    xls = make_fixtures.xls_bytes([{
        "name": "Sheet9", "cells": body,
        "records": [begin_only, _xls_shape(0x06, 41, "AFTERBOXWORD"),
                    _xls_hf(0x0014, "&CLATEHEADWORD")]}])
    _assert_doc(_extract("open.xls", xls), ["[sheet: Sheet9]\nBODYTOKEN"],
                ["sheet 'Sheet9': a Custom View's records begin and never end, so every "
                 "record after them was skipped with them: its print header and footer and "
                 "dialog flag, and so a cell written twice, a comment missing its text, and a "
                 "text box, shape, chart, picture or control, may be lost with no note of its "
                 f"own ({UNREAD})", PAGE_NOTE_XLS])

    # .xlsx: the same, element by element.
    drawing_rels = {
        "xl/worksheets/_rels/sheet1.xml.rels": lambda t: (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rIdV" Type="http://schemas.openxmlformats.org/officeDocument/'
            '2006/relationships/drawing" Target="../drawings/drawing9.xml"/></Relationships>'),
        "xl/drawings/drawing9.xml": lambda t: _drawing_xml(_sp("VIEWBOXWORD", 2, box=True)),
    }
    view_extra = ('<sheetFormatPr defaultRowHeight="15" zeroHeight="1"/>'
                  '<cols><col min="1" max="1" width="0" hidden="1"/></cols>'
                  '<drawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
                  'relationships" r:id="rIdV"/>')
    raw = _viewed_xlsx("<headerFooter><oddHeader>&amp;CBRAVOHEAD</oddHeader></headerFooter>",
                       "<headerFooter><oddHeader>&amp;C&amp;GVIEWHEADWORD</oddHeader>"
                       "</headerFooter>", view_extra, drawing_rels)
    _assert_doc(_extract("inside.xlsx", raw), ["[sheet: Sheet9]\nBRAVOHEAD\nBODYTOKEN"],
                [PAGE_NOTE_XLSX])


_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _drawing_xml(*anchored: str) -> str:
    return (f'<xdr:wsDr xmlns:xdr="{_XDR}" xmlns:a="{_A}" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
            + "".join(anchored) + "</xdr:wsDr>")


def _anchor(inner: str) -> str:
    return ("<xdr:twoCellAnchor><xdr:from><xdr:col>2</xdr:col><xdr:colOff>0</xdr:colOff>"
            "<xdr:row>1</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>5"
            "</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>4</xdr:row><xdr:rowOff>0</xdr:rowOff>"
            f"</xdr:to>{inner}<xdr:clientData/></xdr:twoCellAnchor>")


def _sp(text: str, shape_id: int, box: bool = False, hidden: bool = False) -> str:
    """A DrawingML shape as Excel writes one: each ``\\n`` in ``text`` a new
    paragraph."""
    paragraphs = "".join(f'<a:p><a:r><a:rPr lang="en-US" sz="1100"/><a:t>{html.escape(p)}</a:t>'
                         "</a:r></a:p>" if p else '<a:p><a:endParaRPr lang="en-US"/></a:p>'
                         for p in text.split("\n"))
    hidden_attr = ' hidden="1"' if hidden else ""
    box_attr = ' txBox="1"' if box else ""
    return (f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{shape_id}" '
            f'name="Shape {shape_id}"{hidden_attr}/><xdr:cNvSpPr{box_attr}/></xdr:nvSpPr>'
            '<xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr><xdr:txBody>'
            f'<a:bodyPr/><a:lstStyle/>{paragraphs}</xdr:txBody></xdr:sp>')


def _frame(uri: str, shape_id: int) -> str:
    return (f'<xdr:graphicFrame macro=""><xdr:nvGraphicFramePr><xdr:cNvPr id="{shape_id}" '
            f'name="Frame {shape_id}"/><xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr>'
            '<xdr:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></xdr:xfrm><a:graphic>'
            f'<a:graphicData uri="{uri}"/></a:graphic></xdr:graphicFrame>')


def _choice(inner: str, fallback: str = "") -> str:
    return ('<mc:AlternateContent><mc:Choice xmlns:a14="http://schemas.microsoft.com/office/'
            f'drawing/2010/main" Requires="a14">{inner}</mc:Choice>'
            f"<mc:Fallback>{fallback}</mc:Fallback></mc:AlternateContent>")


def _board_xlsx(sheet_edit=None, parts_edit=None) -> bytes:
    """Excel's worksheet drawing constructs (as real Excel 16 wrote them in
    excel_fix3/excel_real/x04_drawings.xlsx): a two-paragraph text box, a
    shape with text, a group of two shapes with text, a shape with no text, a
    connector, WordArt, a picture, a chart, a SmartArt diagram, a form button
    and check box and an embedded object (each with Excel's hidden drawn copy
    in an mc:Choice, and in the legacy VML drawing), a comment's VML shape, a
    slicer's graphic frame whose mc:Fallback says it is one, and a picture in
    the print header."""
    def build(wb):
        ws = wb.active
        ws.title = "Board"
        ws["A1"], ws["A2"], ws["A3"] = "KITEROW", 5, 7

    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    drawing = _drawing_xml(
        _anchor(_sp("LOONBOX first line\nLOONSECOND line", 2, box=True)),
        _anchor(_sp("GREBESHAPE", 3)),
        _anchor('<xdr:grpSp><xdr:nvGrpSpPr><xdr:cNvPr id="6" name="Group 5"/><xdr:cNvGrpSpPr/>'
                "</xdr:nvGrpSpPr><xdr:grpSpPr/>" + _sp("COOTGROUPA", 4) + _sp("COOTGROUPB", 5)
                + "</xdr:grpSp>"),
        _anchor(_sp("", 7)),
        _anchor('<xdr:cxnSp macro=""><xdr:nvCxnSpPr><xdr:cNvPr id="8" name="Connector 7"/>'
                "<xdr:cNvCxnSpPr/></xdr:nvCxnSpPr><xdr:spPr/></xdr:cxnSp>"),
        _anchor(_sp("SCAUPART", 9)),
        _anchor('<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="11" name="Picture 10"/><xdr:cNvPicPr/>'
                '</xdr:nvPicPr><xdr:blipFill><a:blip r:embed="rId1"/></xdr:blipFill>'
                "<xdr:spPr/></xdr:pic>"),
        _anchor(_frame("http://schemas.openxmlformats.org/drawingml/2006/chart", 12)),
        _anchor(_frame("http://schemas.openxmlformats.org/drawingml/2006/diagram", 13)),
        _choice(_anchor(_sp("GADWALLBUTTON", 1025, hidden=True))),
        _choice(_anchor(_sp("WIGEONCHECK", 1026, hidden=True))),
        _choice(_anchor(_sp("", 1027, hidden=True))),
        _choice(_anchor(_frame("http://schemas.microsoft.com/office/drawing/2010/slicer", 14)),
                _anchor(_sp("FALLBACKSLICERWORD", 15))),
    )
    vml = ('<xml xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:'
           'office:office" xmlns:x="urn:schemas-microsoft-com:office:excel">'
           '<v:shape id="_x0000_s1025" type="#_x0000_t201"><v:textbox><div>GADWALLBUTTON</div>'
           '</v:textbox><x:ClientData ObjectType="Button"><x:Anchor>7, 56, 14, 0, 9, 18, 15, 8'
           '</x:Anchor></x:ClientData></v:shape>'
           '<v:shape id="_x0000_s1026" type="#_x0000_t201"><x:ClientData ObjectType="Checkbox">'
           '</x:ClientData></v:shape>'
           '<v:shape id="_x0000_s1027" type="#_x0000_t75"><v:imagedata o:relid="rId1"/>'
           '<x:ClientData ObjectType="Pict"><x:CF>Pict</x:CF></x:ClientData></v:shape>'
           '<v:shape id="_x0000_s1028" o:spid="_x0000_s1028" type="#_x0000_t202"><v:textbox>'
           '<div>NOTEVMLWORD</div></v:textbox><x:ClientData ObjectType="Note"><x:Row>1</x:Row>'
           '<x:Column>0</x:Column></x:ClientData></v:shape></xml>')
    tail = (f'<drawing xmlns:r="{rel}" r:id="rIdDr"/><legacyDrawing xmlns:r="{rel}" r:id="rIdVml"/>'
            f'<oleObjects xmlns:r="{rel}" xmlns:mc="http://schemas.openxmlformats.org/'
            'markup-compatibility/2006"><mc:AlternateContent><mc:Choice Requires="x14">'
            '<oleObject progId="Packager Shell Object" shapeId="1027" r:id="rIdOle"/></mc:Choice>'
            '<mc:Fallback><oleObject progId="Packager Shell Object" shapeId="1027" r:id="rIdOle"/>'
            '</mc:Fallback></mc:AlternateContent></oleObjects>'
            '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility'
            f'/2006"><mc:Choice Requires="x14"><controls xmlns:r="{rel}">'
            '<mc:AlternateContent><mc:Choice Requires="x14"><control shapeId="1025" r:id="rIdC1" '
            'name="Button 1"/></mc:Choice></mc:AlternateContent>'
            '<mc:AlternateContent><mc:Choice Requires="x14"><control shapeId="1026" r:id="rIdC2" '
            'name="Check Box 2"/></mc:Choice></mc:AlternateContent></controls></mc:Choice>'
            "</mc:AlternateContent>")
    rels = (f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rIdDr" Type="{rel}/drawing" Target="../drawings/drawing1.xml"/>'
            f'<Relationship Id="rIdVml" Type="{rel}/vmlDrawing" Target="../drawings/vmlDrawing1.vml"/>'
            f'<Relationship Id="rIdOle" Type="{rel}/oleObject" Target="../embeddings/oleObject1.bin"/>'
            f'<Relationship Id="rIdC1" Type="{rel}/ctrlProp" Target="../ctrlProps/ctrlProp1.xml"/>'
            f'<Relationship Id="rIdC2" Type="{rel}/ctrlProp" Target="../ctrlProps/ctrlProp2.xml"/>'
            "</Relationships>")

    def sheet(xml: str) -> str:
        xml = re.sub(r"<headerFooter>.*?</headerFooter>|<headerFooter/>", "", xml, flags=re.S)
        xml = xml.replace("</worksheet>", "<headerFooter><oddHeader>&amp;C&amp;G&amp;LSHOVELERHEAD"
                          "</oddHeader></headerFooter>" + tail + "</worksheet>")
        return sheet_edit(xml) if sheet_edit else xml

    parts = {"xl/worksheets/sheet1.xml": sheet,
             "xl/worksheets/_rels/sheet1.xml.rels": lambda t: rels,
             "xl/drawings/drawing1.xml": lambda t: drawing,
             "xl/drawings/vmlDrawing1.vml": lambda t: vml}
    parts.update(parts_edit or {})
    return _patch(_xlsx(build), parts)


def _board_xls(records=None) -> bytes:
    """:func:`_board_xlsx`'s constructs as real Excel 16 saved them to
    ``.xls`` (excel_fix3/excel_real/x04_drawings.xls): each an OBJ record of
    its type (a text box 0x06, a rectangle 0x02, an oval 0x03, a group 0x00,
    a connector 0x1E, a picture 0x08, a chart 0x05 with its own chart
    substream, a button 0x07, a check box 0x0B, an embedded object -- a
    picture with an ftPictFmla), the text in a TXO record. Excel saves a
    SmartArt diagram to ``.xls`` as a picture, and so does this."""
    import struct

    chart = (_xls_shape(0x05, 11)
             + _xls_record(0x0809, struct.pack("<HHHHII", 0x0600, 0x0020, 0, 0, 0, 0))
             + _xls_hf(0x0014, "&CCHARTHEADWORD")
             + _xls_record(0x0203, struct.pack("<HHHd", 0, 0, 0, 99.0))
             + _xls_record(0x000A, b""))
    board = [
        _xls_shape(0x06, 4, "LOONBOX first line\nLOONSECOND line"),
        _xls_shape(0x02, 5, "GREBESHAPE"),
        _xls_shape(0x00, 6), _xls_shape(0x03, 13, "COOTGROUPA"), _xls_shape(0x02, 14, "COOTGROUPB"),
        _xls_shape(0x02, 7), _xls_shape(0x1E, 8), _xls_shape(0x02, 9, "SCAUPART"),
        _xls_shape(0x08, 10, extra=_PICTURE), chart, _xls_shape(0x08, 12, extra=_PICTURE),
        _xls_shape(0x07, 1, "GADWALLBUTTON"), _xls_shape(0x0B, 2, "WIGEONCHECK"),
        _xls_shape(0x08, 3, extra=_EMBEDDED),
    ]
    return make_fixtures.xls_bytes([{
        "name": "Board", "header": "&C&G&LSHOVELERHEAD",
        "cells": [(0, 0, "KITEROW", 0), (1, 0, 5.0, 0), (2, 0, 7.0, 0)],
        "records": board if records is None else records}])


BOARD_PAGE = ("[sheet: Board]\nSHOVELERHEAD\nKITEROW\n5\n7\n"
              "[text box] LOONBOX first line ¶ LOONSECOND line\n[shape] GREBESHAPE\n"
              "[shape] COOTGROUPA\n[shape] COOTGROUPB\n[shape] SCAUPART")


def test_worksheet_drawings_are_read_or_named_in_both_formats(monkeypatch):
    """Review r3 (B4): a text box, a shape with text or a chart on an
    ordinary worksheet vanished with status FULL and no word, in both
    formats, while the same chart on its own tab was disclosed. A text box's
    or shape's text is READ -- one line per shape, ``[text box]`` or
    ``[shape]``, after the sheet's rows (and after a row-cap line) and before
    its comments, in the drawing's own order, a paragraph break as a cell's
    line break -- and what has no text to take is COUNTED in one FINAL note:
    charts, pictures, SmartArt, form or ActiveX controls, embedded objects,
    header pictures. A control's or object's drawn copy is not read as a
    shape; a comment's VML shape and a slicer's fallback text are not read
    at all; an embedded chart's own header and cells are not the sheet's.

    Pages are identical in both formats (the real-Excel pair is too,
    excel_fix3/x06); the notes differ only where the formats do: ``.xls``
    keeps no SmartArt (Excel saves the diagram as a picture) and has no
    slicer (counted in ``.xlsx`` as a drawing object of another kind). A
    drawing part missing from the file is FINAL, a drawing read that raised
    TRANSIENT."""
    counts = ("1 chart(s), {pictures}, {smartart}2 form or ActiveX control(s), 1 embedded "
              "or linked object(s){last}1 picture(s) in its print header or footer{other} "
              f"were not read ({UNREAD})")
    xlsx_note = "sheet 'Board': " + counts.format(
        pictures="1 picture(s)", smartart="1 SmartArt diagram(s), ", last=", ",
        other=" and 1 drawing object(s) of another kind")
    xls_note = "sheet 'Board': " + counts.format(pictures="2 picture(s)", smartart="",
                                                 last=" and ", other="")
    _assert_doc(_extract("board.xlsx", _board_xlsx()), [BOARD_PAGE], [xlsx_note, PAGE_NOTE_XLSX])
    _assert_doc(_extract("board.xls", _board_xls()), [BOARD_PAGE], [xls_note, PAGE_NOTE_XLS])
    assert ex.has_final_marker(xlsx_note) and not ex.has_transient_marker(xlsx_note)

    no_drawing_page = "[sheet: Board]\nSHOVELERHEAD\nKITEROW\n5\n7"
    missing = _board_xlsx(parts_edit={"xl/drawings/drawing1.xml": lambda t: None})
    note = ("sheet 'Board': its drawing part is missing from the file; a text box, shape, "
            f"chart, picture or control on it was not read ({UNREAD})")
    rest = ("sheet 'Board': 2 form or ActiveX control(s), 1 embedded or linked object(s) and 1 "
            f"picture(s) in its print header or footer were not read ({UNREAD})")
    _assert_doc(_extract("missing.xlsx", missing), [no_drawing_page], [note, rest, PAGE_NOTE_XLSX])

    real = ex._xlsx_drawing

    def raising(z, member, mirrored):
        raise MemoryError()

    monkeypatch.setattr(ex, "_xlsx_drawing", raising)
    note = ("sheet 'Board': its drawing could not be read (MemoryError), so a text box, "
            f"shape, chart, picture or control on it is not on its page or counted ({PART_UNREAD})")
    _assert_doc(_extract("raised.xlsx", _board_xlsx()), [no_drawing_page],
                [note, rest, PAGE_NOTE_XLSX])
    assert ex.has_transient_marker(note)
    monkeypatch.setattr(ex, "_xlsx_drawing", real)

    # The shapes follow the rows and a row-cap line, and precede the comments.
    def commented(wb):
        ws = wb.active
        ws.title = "Board"
        for r in range(1, 4):
            ws[f"A{r}"] = f"CAPROW{r - 1}"
        ws["B1"].comment = openpyxl.comments.Comment("NOTEWORD", "AUTHORWORD")

    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    raw = _patch(_xlsx(commented), {
        "xl/worksheets/sheet1.xml": lambda t: t.replace(
            "<legacyDrawing", f'<drawing xmlns:r="{rel}" r:id="rIdBox"/><legacyDrawing', 1),
        "xl/worksheets/_rels/sheet1.xml.rels": lambda t: t.replace(
            "</Relationships>", f'<Relationship Id="rIdBox" Type="{rel}/drawing" '
            'Target="../drawings/drawing7.xml"/></Relationships>'),
        "xl/drawings/drawing7.xml": lambda t: _drawing_xml(_anchor(_sp("BOXAFTERCAP", 2, box=True))),
    })
    monkeypatch.setattr(ex, "_XLSX_MAX_ROWS", 2)
    _assert_doc(_extract("capped.xlsx", raw),
                ["[sheet: Board]\nCAPROW0\nCAPROW1\n[... workbook truncated at 2 rows]\n"
                 "[text box] BOXAFTERCAP\n[comment on B1 by AUTHORWORD] NOTEWORD"],
                ["workbook truncated at 2 rows while reading 'Board'; the rest of that sheet "
                 f"was not read ({UNREAD})", PAGE_NOTE_XLSX])
    capped = make_fixtures.xls_bytes([{
        "name": "Board", "cells": [(r, 0, f"CAPROW{r}", 0) for r in range(3)],
        "notes": [(0, 1, "AUTHORWORD", "NOTEWORD")],
        "records": [_xls_shape(0x06, 7, "BOXAFTERCAP")]}])
    _assert_doc(_extract("capped.xls", capped),
                ["[sheet: Board]\nCAPROW0\nCAPROW1\n[... workbook truncated at 2 rows]\n"
                 "[text box] BOXAFTERCAP\n[comment on B1 by AUTHORWORD] NOTEWORD"],
                ["workbook truncated at 2 rows while reading 'Board'; the rest of that sheet "
                 f"was not read ({UNREAD})", PAGE_NOTE_XLS])


def test_a_raw_cell_failure_keeps_openpyxls_dates_times_and_durations(monkeypatch):
    """Review-fix round 3: dates, times and durations are now read from the
    cell's own ``<v>``. When that stream raises, openpyxl's converted value
    is shown under the rules it had -- a duration still ``30:00:00``, never
    Python's ``1 day, 6:00:00``; a time-only format's day count still the
    number, never a date -- and the transient note says what that costs."""
    def build(wb):
        ws = wb.active
        ws.title = "Clock"
        for c, (value, fmt) in enumerate(((1.25, "[h]:mm:ss"), (0.25, "h:mm"),
                                          (45489.25, "yyyy-mm-dd h:mm"), (1.5, "h:mm")), start=1):
            ws.cell(1, c, value).number_format = fmt

    real = ex._xlsx_iter_rows
    calls = {"n": 0}

    def raw_stream_fails(z, part):
        calls["n"] += 1
        if calls["n"] == 2:            # the extras stream first, then the raw cells
            raise MemoryError()
        return real(z, part)

    monkeypatch.setattr(ex, "_xlsx_iter_rows", raw_stream_fails)
    note = ("sheet 'Clock': its own cell XML could not be read (MemoryError), so a formula "
            "whose stored result is empty may be shown as its formula and counted as having no "
            "stored value, a date-formatted number that is no date may show as #VALUE!, a "
            "time-only cell may show the number openpyxl converted rather than the one stored, "
            "and a date, time or duration may be a millisecond off Excel's rounding "
            f"({PART_UNREAD})")
    _assert_doc(_extract("clock.xlsx", _xlsx(build)),
                ["[sheet: Clock]\n30:00:00\t06:00:00\t2024-07-16 06:00:00\t1.5"],
                [note, "1 cell(s) formatted as a time with no date held a value of a day or "
                 "more; the number is shown as stored", PAGE_NOTE_XLSX])


def test_a_hidden_sheet_is_said_read_only_when_its_rows_were():
    """Review r3 (C): the hidden-sheet note was settled before the rows were
    read, so a hidden sheet whose first row could not be read (a NaN value),
    or whose first row the row cap discarded, was said to be "read like any
    other sheet" beside the note saying it was not."""
    def build(wb):
        wb.active.title = "Shown"
        wb.active["A1"] = "SHOWNWORD"
        hid = wb.create_sheet("Hid")
        hid["A1"] = 1.5
        hid.sheet_state = "hidden"

    raw = _patch(_xlsx(build), {"xl/worksheets/sheet2.xml": lambda t: t.replace(
        "<v>1.5</v>", "<v>NaN</v>")})
    _assert_doc(_extract("nan.xlsx", raw), ["[sheet: Shown]\nSHOWNWORD", "[sheet: Hid]"],
                ["sheet 'Hid' is hidden in the workbook",
                 "sheet 'Hid': its rows from row 1 on could not be read (ValueError), so they "
                 f"are not on its page ({PART_UNREAD})", PAGE_NOTE_XLSX])


def test_a_hidden_sheet_the_cap_cut_before_its_first_row_is_not_said_read(monkeypatch):
    """Review r3 (C), the row-cap boundary: the first sheet holds exactly the
    cap, so the hidden second sheet is opened (its header shows) and its
    first row discarded -- none of its rows was read."""
    monkeypatch.setattr(ex, "_XLSX_MAX_ROWS", 2)

    def build(wb):
        wb.active.title = "Full"
        wb.active["A1"], wb.active["A2"] = "FULLONE", "FULLTWO"
        hid = wb.create_sheet("Hid")
        hid["A1"] = "HIDROW"
        hid.oddHeader.center.text = "TWITEHEAD"
        hid.sheet_state = "hidden"

    truncation = (f"workbook truncated at 2 rows while reading 'Hid'; the rest of that sheet "
                  f"was not read ({UNREAD})")
    page = ["[sheet: Full]\nFULLONE\nFULLTWO",
            "[sheet: Hid]\nTWITEHEAD\n[... workbook truncated at 2 rows]"]
    _assert_doc(_extract("cap.xlsx", _xlsx(build)), page,
                ["sheet 'Hid' is hidden in the workbook", truncation, PAGE_NOTE_XLSX])
    xls = make_fixtures.xls_bytes([
        {"name": "Full", "cells": [(0, 0, "FULLONE", 0), (1, 0, "FULLTWO", 0)]},
        {"name": "Hid", "visibility": 1, "header": "&CTWITEHEAD",
         "cells": [(0, 0, "HIDROW", 0)]}])
    _assert_doc(_extract("cap.xls", xls), page,
                ["sheet 'Hid' is hidden in the workbook", truncation, PAGE_NOTE_XLS])


def test_an_iso_date_cell_keeps_its_date_under_any_format():
    """Review r3 (C, a regression from round 2): a cell storing an ISO
    date-time (``t="d"``, which openpyxl writes with ``iso_dates``) under a
    number or time-only format read as a serial number, under a note saying
    the number was "shown as stored" -- the file stores the ISO text. It
    reads as stored, whatever the format, with no note."""
    def build(wb):
        wb.iso_dates = True
        ws = wb.active
        ws.title = "Iso"
        for c, fmt in enumerate(("General", "0.00", "h:mm", "yyyy-mm-dd"), start=1):
            ws.cell(1, c, datetime.datetime(2024, 7, 16, 8, 30)).number_format = fmt

    raw = _xlsx(build)
    sheet = zipfile.ZipFile(io.BytesIO(raw)).read("xl/worksheets/sheet1.xml").decode()
    assert sheet.count('t="d"') == 4, sheet
    _assert_doc(_extract("iso.xlsx", raw), ["[sheet: Iso]\n" + "\t".join(["2024-07-16 08:30:00"] * 4)],
                [PAGE_NOTE_XLSX])


def test_an_xls_cell_typed_a_date_with_no_format_found_says_so():
    """Review r3 (C): note 40 was listed as unreachable. An ``.xls`` cell
    whose format index is a CJK built-in date xlrd knows (57) but whose
    workbook declares no FORMAT for it reads as the number stored, under
    that note."""
    xls = make_fixtures.xls_bytes([{"name": "Dated", "cells": [(0, 0, 45489.0, 57)]}])
    _assert_doc(_extract("dated.xls", xls), ["[sheet: Dated]\n45489"],
                ["1 cell(s) typed as a date carry no number format this reader could find; "
                 "the number is shown as stored", PAGE_NOTE_XLS])


def test_the_persons_note_when_the_workbook_part_is_not_found(monkeypatch):
    """Review r3 (C, mutants U01/U03): when the workbook part itself cannot
    be found, threaded comments are credited to person ids under a TRANSIENT
    note -- never FINAL, never unsaid."""
    def not_found(z):
        raise MemoryError()

    monkeypatch.setattr(ex, "_xlsx_workbook_part", not_found)
    doc = _extract("threads.xlsx", _threads_workbook())
    persons = ("the workbook part was not found (MemoryError), so its persons part was not "
               "read; 1 threaded comment(s) are credited to person ids, not names "
               f"({PART_UNREAD})")
    assert persons in doc.notes, doc.notes
    assert ex.has_transient_marker(persons) and not ex.has_final_marker(persons)
    assert "[comment on A1 by {P1}] MAGPIE" in doc.pages[0].text, doc.pages[0].text

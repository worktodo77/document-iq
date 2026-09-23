"""The Word fidelity package: the Word reader's regression suite.

Every test here holds a part of the Word spec
(``docs/design/word_fidelity_spec.md``): every character in a Word
document's text-bearing parts reaches the page text, or is named in a note
carrying an evidence marker. Each test was watched failing against the
extractor that emitted paragraphs, then every table, then one note claiming
the file had no page boundaries, or against the defect named in its
docstring. Do not make a failure here pass by loosening an assertion.

One behaviour per test, per the spec. A position test uses ``str.find``/
``str.count`` and reports every position it computed; a built-package test
compares the exact lines or notes and reports what it got. Either way a
failure here is an assertion with a diagnostic message -- never a bare
``ValueError`` or ``IndexError`` from a sentinel that was never found.
"""

from __future__ import annotations

import io
import re
import zipfile

import pytest
from lxml import etree

from dociq.identify.bates import detect_candidates
from dociq.ingest import extract as ex
from dociq.ingest.pagemodel import normalize

from .conftest import FIXTURES
from .fixtures import document


def _pages(name: str, opt: ex.ExtractOptions | None = None):
    path = FIXTURES / name
    return ex.extract(path.name, path.read_bytes(), opt)


def _page_text(name: str) -> str:
    got = _pages(name)
    assert got.pages, f"{name} produced no pages at all: notes={got.notes!r}"
    return got.pages[0].text


def _positions(text: str, *sentinels: str) -> dict:
    return {s: text.find(s) for s in sentinels}


# ---------------------------------------------------------------------------
# 1. Body order
# ---------------------------------------------------------------------------


def test_body_order_follows_document_order_not_paragraphs_then_tables():
    text = _page_text("16_word_constructs.docx")
    order = ["AARDVARK", "BADGER", "CARIBOU", "DINGO", "ELAND", "FERRET",
             "GAZELLE", "HERON", "IBEX", "JACKAL", "LYNX", "MARMOT", "OCELOT",
             "QUAIL", "RAVEN", "TAPIR", "ALPACA"]
    pos = _positions(text, *order)
    found = [pos[s] for s in order]
    assert all(p != -1 for p in found) and found == sorted(found), (
        f"expected this sentinel order (document order, not "
        f"paragraphs-then-tables): {order}; positions found: {pos}")


# ---------------------------------------------------------------------------
# 2. Merged cells are not copied
# ---------------------------------------------------------------------------


def test_merged_cells_are_not_copied():
    text = _page_text("16_word_constructs.docx")
    n_badger, n_caribou = text.count("BADGER"), text.count("CARIBOU")
    assert n_badger == 1 and n_caribou == 1, (
        f"BADGER (gridSpan=2 cell) occurred {n_badger} time(s); CARIBOU "
        f"(vMerge restart cell) occurred {n_caribou} time(s); a spanned or "
        f"continuation grid slot must not repeat its neighbour's text")


# ---------------------------------------------------------------------------
# 3. Nested table read
# ---------------------------------------------------------------------------


def test_nested_table_is_read():
    text = _page_text("16_word_constructs.docx")
    i = text.find("FERRET")
    assert i != -1, (
        f"FERRET (the nested table's only cell) is not in the page text "
        f"(position {i}): {text!r}")


# ---------------------------------------------------------------------------
# 4. Content controls read
# ---------------------------------------------------------------------------


def test_content_controls_are_read():
    text = _page_text("16_word_constructs.docx")
    i_heron, i_ibex = text.find("HERON"), text.find("IBEX")
    assert i_heron != -1 and i_ibex != -1, (
        f"HERON (block content control) at {i_heron}; IBEX (inline content "
        f"control) at {i_ibex}; both must be read")


# ---------------------------------------------------------------------------
# 5. Tracked changes, DEFAULT view
# ---------------------------------------------------------------------------


def test_tracked_changes_keep_insertions_in_the_body_and_list_deletions_after_it():
    """D-56: the body reads with tracked changes accepted, and the deleted
    passage is listed after the notes and comments. Before the ruling this
    test held the omission and its count note; a listed deletion is no longer
    missing evidence, so no marked note may count it."""
    got = _pages("16_word_constructs.docx")
    text = got.pages[0].text
    lines = text.split("\n")
    i_jackal = text.find("JACKAL")
    listed = [i for i, line in enumerate(lines) if line == "[deleted by DocIQ fixtures] KOALA"]
    body_koala = [line for line in lines if "KOALA" in line and not line.startswith("[deleted by ")]
    i_vulture = next((i for i, line in enumerate(lines) if "VULTURE" in line), -1)
    marker_note = [n for n in got.notes if ex.has_evidence_marker(n) and "delet" in n.lower()]
    assert (i_jackal != -1 and listed and not body_koala and listed[0] > i_vulture != -1
            and not marker_note), (
        f"JACKAL (tracked insertion) at {i_jackal}; KOALA must be listed once as "
        f"'[deleted by DocIQ fixtures] KOALA' after the comment (VULTURE, line "
        f"{i_vulture}), at lines {listed}, and never in the body: {body_koala!r}; "
        f"no marked note may count a listed deletion: {marker_note!r}")


# ---------------------------------------------------------------------------
# 6. Fields
# ---------------------------------------------------------------------------


def test_fields_show_cached_and_complex_results_not_instruction_codes():
    text = _page_text("16_word_constructs.docx")
    i_lynx = text.find("LYNX")
    i_marmot = text.find("MARMOT")
    i_narwhal = text.find("NARWHAL")
    assert i_lynx != -1 and i_marmot != -1 and i_narwhal == -1, (
        f"LYNX (fldSimple cached result) at {i_lynx}; MARMOT (complex field "
        f"result run) at {i_marmot}; NARWHAL (the field's own instruction "
        f"code, must be excluded) at {i_narwhal}")


# ---------------------------------------------------------------------------
# 7. Hyperlink
# ---------------------------------------------------------------------------


def test_external_hyperlink_target_follows_its_display_text():
    text = _page_text("16_word_constructs.docx")
    i_ocelot, i_pelican = text.find("OCELOT"), text.find("PELICAN")
    assert i_ocelot != -1 and i_pelican != -1 and i_pelican > i_ocelot, (
        f"OCELOT (hyperlink display text) at {i_ocelot}; PELICAN (from the "
        f"hyperlink's external target, https://example.invalid/PELICAN) at "
        f"{i_pelican}; PELICAN must follow OCELOT")


# ---------------------------------------------------------------------------
# 8. Text box
# ---------------------------------------------------------------------------


def test_text_box_is_read_once_between_its_anchor_and_the_next_paragraph():
    text = _page_text("16_word_constructs.docx")
    n_raven = text.count("RAVEN")
    i_quail, i_raven, i_tapir = (text.find("QUAIL"), text.find("RAVEN"),
                                 text.find("TAPIR"))
    assert (n_raven == 1 and i_quail != -1 and i_raven != -1 and i_tapir != -1
            and i_quail < i_raven < i_tapir), (
        f"RAVEN occurred {n_raven} time(s) (it is stored twice in the XML, "
        f"once as mc:Choice and once as mc:Fallback, and must be read only "
        f"once); QUAIL at {i_quail}, RAVEN at {i_raven}, TAPIR at {i_tapir} "
        f"-- expected QUAIL < RAVEN < TAPIR")


# ---------------------------------------------------------------------------
# 9. Notes and comments
# ---------------------------------------------------------------------------


def test_notes_and_comments_are_disclosed_after_the_body_without_dates():
    got = _pages("16_word_constructs.docx")
    text = got.pages[0].text
    i_alpaca = text.find("ALPACA")
    positions = {
        "SALAMANDER": text.find("SALAMANDER"),
        "URCHIN": text.find("URCHIN"),
        "VULTURE": text.find("VULTURE"),
        "WALRUS Reviewer": text.find("WALRUS Reviewer"),
    }
    assert i_alpaca != -1 and all(p != -1 and p > i_alpaca for p in positions.values()), (
        f"ALPACA (last body paragraph) at {i_alpaca}; the footnote, endnote "
        f"and comment sentinels must all follow it: {positions}")
    assert "2019" not in text, (
        "the comment's own w:date (2019-01-02) must never reach the page "
        f"text; found '2019' in: {text!r}")


# ---------------------------------------------------------------------------
# 10. Headers lead
# ---------------------------------------------------------------------------


def test_headers_lead_the_page_text():
    text = _page_text("16_word_constructs.docx")
    first_lines = text.split("\n")[:3]
    head = "\n".join(first_lines)
    i_yak, i_xerus = head.find("YAK"), head.find("XERUS")
    assert i_yak != -1 and i_xerus != -1, (
        f"YAK (first-page header) and XERUS (default header) must both "
        f"be within the page text's first 3 lines; first 3 lines: "
        f"{first_lines!r}")


# ---------------------------------------------------------------------------
# 11. Footer trails, and is a Bates zone candidate
# ---------------------------------------------------------------------------


def test_footer_trails_the_page_text_and_is_a_bates_zone_candidate():
    got = _pages("16_word_constructs.docx")
    text = got.pages[0].text
    last_lines = text.split("\n")[-8:]
    tail = "\n".join(last_lines)
    i_zebu, i_stamp = tail.find("ZEBU"), tail.find("MNFV 000777")

    doc = document("16_word_constructs.docx", (got.pages[0],))
    candidates = detect_candidates((doc,))
    found_stamp = any(c.raw == "MNFV 000777" for c in candidates)

    assert i_zebu != -1 and i_stamp != -1 and found_stamp, (
        f"ZEBU (default footer, first paragraph) at tail-relative {i_zebu}; "
        f"MNFV 000777 (default footer, second paragraph) at tail-relative "
        f"{i_stamp}, within last 8 lines {last_lines!r}; Bates candidates "
        f"the existing detect_candidates() API found on this page: "
        f"{candidates!r}")


# ---------------------------------------------------------------------------
# 12. Unread parts disclosed
# ---------------------------------------------------------------------------


def test_altchunk_and_chart_are_disclosed_as_unread():
    got = _pages("16_word_constructs.docx")
    marked = [n for n in got.notes if ex.has_evidence_marker(n)]
    has_altchunk = any(re.search(r"\baltchunk\b", n, re.IGNORECASE) for n in marked)
    has_chart = any(re.search(r"\bchart\b", n, re.IGNORECASE) for n in marked)
    assert has_altchunk and has_chart, (
        f"expected marked (has_evidence_marker) notes naming an altChunk and "
        f"a chart, as whole words; marked notes: {marked!r}; all notes: "
        f"{got.notes!r}")


# ---------------------------------------------------------------------------
# 12b. A large body picture is disclosed as unread; a small one is not
# ---------------------------------------------------------------------------


def test_a_large_picture_is_disclosed_and_a_small_one_is_not():
    got = _pages("16_word_constructs.docx")
    picture_notes = [n for n in got.notes
                      if ex.has_evidence_marker(n) and "picture" in n.lower()]
    assert len(picture_notes) == 1, (
        f"expected exactly one marked note naming pictures -- the 6in x 8in "
        f"body picture covers ~51% of the page and must be disclosed; the "
        f"1in x 1in body picture covers ~1% and must not be: "
        f"{picture_notes!r}; all notes: {got.notes!r}")
    count = re.search(r"(\d+)\s+picture", picture_notes[0], re.IGNORECASE)
    assert count is not None and count.group(1) == "1", (
        f"expected the picture note's own count to be 1 (only the 6in x 8in "
        f"picture qualifies): {picture_notes[0]!r}")


# ---------------------------------------------------------------------------
# 13. The page note is true
# ---------------------------------------------------------------------------


def test_the_page_note_does_not_deny_the_rendered_page_break_it_carries():
    got = _pages("16_word_constructs.docx")
    denying = [n for n in got.notes if "no page boundaries" in n.lower()]
    assert not denying, (
        f"this fixture carries a rendered page break (GAZELLE's run holds "
        f"w:lastRenderedPageBreak); no note may claim the file carries none: "
        f"{denying!r}; all notes: {got.notes!r}")


# ---------------------------------------------------------------------------
# 14. DERIVED class test: every w:t reaches the page text, or is excluded by
#     name (mc:Fallback).
# ---------------------------------------------------------------------------

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_WT_TAGS = ("{%s}t" % _W_NS, "{%s}delText" % _W_NS)
_TEXTPATH_TAG = "{urn:schemas-microsoft-com:vml}textpath"
_TEXT_REL_TYPES = ("header", "footer", "footnotes", "endnotes", "comments")


def _wt_fragments(raw: bytes) -> list:
    """Every non-blank ``w:t`` and ``w:delText`` fragment, and every VML
    ``v:textpath`` string, in this DOCX's main document part and in every
    header, footer, footnotes, endnotes and comments part the main part's
    relationships name, skipping anything nested under ``mc:Fallback``.

    An independent, from-the-XML reading of "what this file's text-bearing
    parts actually hold", deliberately not sharing a code path with
    ``_extract_docx`` -- so it can catch what that extractor drops rather
    than agreeing with it by construction. The parts are found the way a
    package names them, through ``_rels/.rels`` and the main part's own
    relationships; this test once listed fixed file names, and so could not
    see a reader that did the same.
    """
    import posixpath

    rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    fragments: list = []
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        names = set(z.namelist())

        def targets(source: str, kinds) -> list:
            folder, _sep, base = source.rpartition("/")
            rels_name = f"{folder}/_rels/{base}.rels" if folder else f"_rels/{base}.rels"
            if rels_name not in names:
                return []
            out = []
            for rel in etree.fromstring(z.read(rels_name)).iter(rel_ns + "Relationship"):
                if rel.get("Type", "").rsplit("/", 1)[-1] in kinds and rel.get("TargetMode") != "External":
                    target = rel.get("Target")
                    out.append(posixpath.normpath(target.lstrip("/")) if target.startswith("/")
                               else posixpath.normpath(posixpath.join(folder, target)))
            return out

        main = targets("", ("officeDocument",))[0]
        for name in [main] + targets(main, _TEXT_REL_TYPES):
            root = etree.fromstring(z.read(name))
            found = [(t, t.text) for t in root.iter(*_WT_TAGS)]
            found += [(t, t.get("string")) for t in root.iter(_TEXTPATH_TAG)]
            for t, value in found:
                under_fallback = False
                ancestor = t.getparent()
                while ancestor is not None:
                    if etree.QName(ancestor).localname == "Fallback":
                        under_fallback = True
                        break
                    ancestor = ancestor.getparent()
                if under_fallback:
                    continue
                if value and value.strip():
                    fragments.append(value)
    return fragments


@pytest.mark.parametrize("name", ["16_word_constructs.docx", "05_letter.docx"])
def test_every_w_t_fragment_reaches_the_page_text_or_is_excluded_by_name(name):
    raw = (FIXTURES / name).read_bytes()
    fragments = _wt_fragments(raw)
    page_text = normalize(_page_text(name))
    missing = [f for f in fragments if normalize(f) not in page_text]
    assert not missing, (
        f"{name}: {len(missing)} of {len(fragments)} w:t, w:delText or "
        f"v:textpath fragment(s) from the main part and its related header, "
        f"footer, footnotes, endnotes and comments parts do not appear in the "
        f"normalized page text: {missing!r}")


# ---------------------------------------------------------------------------
# 15/16. The two unmarked OCR notes in _extract_pdf (Word spec part 10)
# ---------------------------------------------------------------------------


def test_ocr_disabled_note_carries_an_evidence_marker():
    got = _pages("02_scanned_instruction.pdf", ex.ExtractOptions(ocr_enabled=False))
    disabled_note = next((n for n in got.notes if "OCR disabled" in n), None)
    assert disabled_note is not None and ex.has_evidence_marker(disabled_note), (
        f"the 'OCR disabled' note must carry an evidence marker: "
        f"{disabled_note!r}; all notes: {got.notes!r}")


def test_ocr_unavailable_note_carries_an_evidence_marker(monkeypatch):
    monkeypatch.setattr(ex, "ocr_available", lambda: False)
    monkeypatch.setattr(
        ex, "ocr_models_present", lambda: (False, "stub: models not present"))
    got = _pages("02_scanned_instruction.pdf")
    unavailable_note = next((n for n in got.notes if "OCR is unavailable" in n), None)
    assert unavailable_note is not None and ex.has_evidence_marker(unavailable_note), (
        f"the 'OCR is unavailable' note must carry an evidence marker: "
        f"{unavailable_note!r}; all notes: {got.notes!r}")


# ---------------------------------------------------------------------------
# Word spec part 9, wherever the reader walks: an Office 2016+ chart, a
# construct outside the body, and a page that cannot be measured. Every input
# here is built in ``tmp_path``, not by changing fixture 16, so no existing
# position test above moves.
# ---------------------------------------------------------------------------


def _extract_bytes(name: str, raw: bytes):
    return ex.extract(name, raw)


def test_chartex_chart_is_disclosed_as_unread(tmp_path):
    """Word spec part 9: chart detection matched only a graphicData URI ending
    '/chart'. An Office 2016+ chartEx chart (waterfall, funnel, ...) uses the
    namespace ``.../office/drawing/2014/chartex`` verbatim instead and must
    be detected too."""
    import docx as _docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn

    d = _docx.Document()
    d.add_paragraph("BEFORE the chartEx chart.")
    sect_pr = d.element.body.find(qn("w:sectPr"))
    ns = (
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:cx="http://schemas.microsoft.com/office/drawing/2014/chartex"'
    )
    sect_pr.addprevious(parse_xml(
        '<w:p ' + ns + '><w:r><w:drawing>'
        '<wp:inline distT="0" distB="0" distL="0" distR="0">'
        '<wp:extent cx="1828800" cy="1828800"/>'
        '<wp:docPr id="1" name="ChartEx1"/>'
        '<a:graphic><a:graphicData '
        'uri="http://schemas.microsoft.com/office/drawing/2014/chartex">'
        '<cx:chart/>'
        '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>'))
    d.add_paragraph("AFTER the chartEx chart.")

    path = tmp_path / "chartex.docx"
    d.save(str(path))
    got = _extract_bytes(path.name, path.read_bytes())

    marked = [n for n in got.notes if ex.has_evidence_marker(n)]
    has_chart = any(re.search(r"\bchart\b", n, re.IGNORECASE) for n in marked)
    assert has_chart, (
        f"expected a marked note naming a chart for an Office 2016+ chartEx "
        f"graphic (graphicData uri is the chartex namespace, not a URI "
        f"ending '/chart'): marked={marked!r}; all notes={got.notes!r}")


def test_unread_constructs_are_disclosed_outside_the_body_too(tmp_path):
    """Word spec part 9: altChunk, chart, SmartArt and large-picture disclosure
    used to gate on ``in_body`` -- so one in a header, a footer, a footnote,
    an endnote, a comment or a text box was lost without a word. This proves
    three of those locations at once: a chart in a HEADER, an altChunk in a
    FOOTNOTE, and a chart PLUS a large picture inside a TEXT BOX (itself in
    the body).

    A presence-only check ("some marked chart/altChunk/picture note
    exists") cannot see a construct counted MORE than once.
    ``_note_drawing`` is called on the outer text-box ``<w:drawing>`` (whose
    own ``.iter()`` already walks down into anything nested inside it,
    including a chart or picture placed in the box) and is then called
    AGAIN when the walk separately reaches that nested ``<w:drawing>`` on
    its own -- so a single chart or picture inside a text box was measured
    to come out as 2 in a draft of the reader. Asserting the exact counts
    a CORRECT read produces (2 chart(s)
    -- header + text box, each real construct appearing once; 1 altChunk(s);
    1 picture(s)) pins the spec regardless of whether this particular
    double-count is what is fixed to satisfy it.
    """
    import docx as _docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn

    ns_all = (
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"'
    )
    w_only = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    r_only = ('xmlns:r="http://schemas.openxmlformats.org/officeDocument/'
             '2006/relationships"')

    d = _docx.Document()
    d.add_paragraph("BODY paragraph, nothing special.")
    sect_pr = d.element.body.find(qn("w:sectPr"))

    # A large picture (6in x 8in on an 8.5in x 11in page, ~51%) inside a
    # text box -- same size word_constructs_docx uses for its own body
    # picture disclosure test, so the >=25% threshold is comfortably cleared.
    picture_drawing = (
        '<w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
        '<wp:extent cx="5486400" cy="7315200"/>'
        '<wp:docPr id="9" name="BoxPic"/>'
        '<a:graphic><a:graphicData '
        'uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<pic:pic><pic:blipFill><a:blip r:embed="rIdBoxImg"/></pic:blipFill></pic:pic>'
        '</a:graphicData></a:graphic></wp:inline></w:drawing>'
    )
    # A chart alongside the picture, both inside the same text box -- the
    # second location the double count above needs.
    box_chart_drawing = (
        '<w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
        '<wp:extent cx="1828800" cy="1828800"/>'
        '<wp:docPr id="10" name="BoxChart"/>'
        '<a:graphic><a:graphicData '
        'uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
        '<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
        'r:id="rIdBoxChart"/>'
        '</a:graphicData></a:graphic></wp:inline></w:drawing>'
    )
    sect_pr.addprevious(parse_xml(
        '<w:p ' + ns_all + '>'
        '<w:r><w:t>Text box anchor.</w:t></w:r>'
        '<w:r><mc:AlternateContent><mc:Choice Requires="wps">'
        '<w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
        '<wp:extent cx="914400" cy="914400"/>'
        '<wp:docPr id="2" name="TextBox1"/>'
        '<a:graphic><a:graphicData '
        'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp><wps:cNvSpPr txBox="1"/>'
        '<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="914400"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr>'
        '<wps:txbx><w:txbxContent>'
        '<w:p><w:r><w:t>box text</w:t></w:r></w:p>'
        '<w:p><w:r>' + picture_drawing + '</w:r></w:p>'
        '<w:p><w:r>' + box_chart_drawing + '</w:r></w:p>'
        '</w:txbxContent></wps:txbx>'
        '<wps:bodyPr/></wps:wsp>'
        '</a:graphicData></a:graphic></wp:inline></w:drawing>'
        '</mc:Choice><mc:Fallback><w:pict/></mc:Fallback></mc:AlternateContent></w:r>'
        '</w:p>'))
    tapir_p = d.add_paragraph("Paragraph carrying a footnote reference.")
    tapir_p._p.append(parse_xml(
        '<w:r ' + w_only + '><w:footnoteReference w:id="1"/></w:r>'))
    # wire the header part into the body's own (final) sectPr
    sect_pr.insert(0, parse_xml(
        '<w:headerReference w:type="default" r:id="rIdHeader1" '
        + w_only + " " + r_only + "/>"))

    buf = io.BytesIO()
    d.save(buf)
    base = buf.getvalue()

    zin = zipfile.ZipFile(io.BytesIO(base))
    names = zin.namelist()
    ct = zin.read("[Content_Types].xml").decode("utf-8")
    rels = zin.read("word/_rels/document.xml.rels").decode("utf-8")
    pfx = "application/vnd.openxmlformats-officedocument.wordprocessingml."
    add_ct = (
        '<Override PartName="/word/footnotes.xml" ContentType="' + pfx + 'footnotes+xml"/>'
        '<Override PartName="/word/header1.xml" ContentType="' + pfx + 'header+xml"/>'
        '<Override PartName="/word/charts/chart1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.drawingml.chart+xml"/>'
        '<Override PartName="/word/charts/chart2.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.drawingml.chart+xml"/>'
        '<Default Extension="html" ContentType="text/html"/>'
    )
    assert "</Types>" in ct
    ct = ct.replace("</Types>", add_ct + "</Types>")
    rpfx = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    add_rel = (
        '<Relationship Id="rIdFootnotes" Type="' + rpfx + 'footnotes" Target="footnotes.xml"/>'
        '<Relationship Id="rIdHeader1" Type="' + rpfx + 'header" Target="header1.xml"/>'
        # The text-box chart lives in the BODY (word/document.xml), so its
        # own r:id is resolved through document.xml's OWN relationships,
        # unlike the header chart below (resolved through header1.xml.rels).
        '<Relationship Id="rIdBoxChart" Type="' + rpfx + 'chart" Target="charts/chart2.xml"/>'
    )
    assert "</Relationships>" in rels
    rels = rels.replace("</Relationships>", add_rel + "</Relationships>")

    footnotes_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:footnotes ' + w_only + '>'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0">'
        '<w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
        '<w:footnote w:id="1">'
        '<w:p><w:r><w:footnoteRef/></w:r>'
        '<w:r><w:t xml:space="preserve"> footnote text.</w:t></w:r></w:p>'
        '<w:altChunk r:id="rIdAltChunkFN" ' + r_only + '/>'
        '</w:footnote></w:footnotes>'
    )
    footnotes_rels_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdAltChunkFN" Type="' + rpfx
        + 'aFChunk" Target="altchunk_fn.html"/></Relationships>'
    )
    header_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:hdr ' + ns_all + '>'
        '<w:p><w:r><w:t>Header text.</w:t></w:r>'
        '<w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
        '<wp:extent cx="1828800" cy="1828800"/>'
        '<wp:docPr id="3" name="HeaderChart1"/>'
        '<a:graphic><a:graphicData '
        'uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
        '<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
        'r:id="rIdHeaderChart"/>'
        '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>'
        '</w:p></w:hdr>'
    )
    header_rels_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdHeaderChart" Type="' + rpfx
        + 'chart" Target="charts/chart1.xml"/></Relationships>'
    )
    chart_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">'
        '<c:chart/></c:chartSpace>'
    )

    path = tmp_path / "unread_outside_body.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            if n == "[Content_Types].xml":
                zout.writestr(n, ct)
            elif n == "word/_rels/document.xml.rels":
                zout.writestr(n, rels)
            else:
                zout.writestr(n, zin.read(n))
        zout.writestr("word/footnotes.xml", footnotes_xml)
        zout.writestr("word/_rels/footnotes.xml.rels", footnotes_rels_xml)
        zout.writestr("word/altchunk_fn.html",
                      b"<html><body>ALTCHUNK IN FOOTNOTE</body></html>")
        zout.writestr("word/header1.xml", header_xml)
        zout.writestr("word/_rels/header1.xml.rels", header_rels_xml)
        zout.writestr("word/charts/chart1.xml", chart_xml)
        zout.writestr("word/charts/chart2.xml", chart_xml)
    zin.close()

    got = _extract_bytes(path.name, path.read_bytes())
    marked = [n for n in got.notes if ex.has_evidence_marker(n)]
    altchunk_notes = [n for n in marked if re.search(r"\baltchunk\b", n, re.IGNORECASE)]
    chart_notes = [n for n in marked if re.search(r"\bchart\b", n, re.IGNORECASE)]
    picture_notes = [n for n in marked if "picture" in n.lower()]
    assert altchunk_notes and chart_notes and picture_notes, (
        f"expected marked notes for an altChunk (in a footnote), a chart "
        f"(in a header AND a text box) and a large picture (in a text box) "
        f"-- none of these live in the body: marked={marked!r}; all "
        f"notes={got.notes!r}")

    def _count(notes: list[str], word: str) -> int:
        m = re.search(r"(\d+)\s+" + word, notes[0], re.IGNORECASE) if notes else None
        assert m is not None, f"no count found in {notes!r}"
        return int(m.group(1))

    assert _count(chart_notes, "chart") == 2, (
        f"expected exactly 2 chart(s) -- one in the header, one in the text "
        f"box, each counted once: {chart_notes!r}")
    assert _count(altchunk_notes, "altChunk") == 1, (
        f"expected exactly 1 altChunk(s) (the footnote's): {altchunk_notes!r}")
    assert _count(picture_notes, "picture") == 1, (
        f"expected exactly 1 picture(s) (the text box's large picture): "
        f"{picture_notes!r}")


def test_unread_constructs_in_footer_endnote_and_comment(tmp_path):
    """Word spec part 9, the three locations the test above does
    not reach: a SmartArt drawing in a FOOTER, a chart in an ENDNOTE, and an
    altChunk in a COMMENT."""
    import docx as _docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn

    ns_all = (
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" '
        'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"'
    )
    w_only = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    r_only = ('xmlns:r="http://schemas.openxmlformats.org/officeDocument/'
             '2006/relationships"')

    d = _docx.Document()
    d.add_paragraph("BODY paragraph, nothing special.")
    sect_pr = d.element.body.find(qn("w:sectPr"))
    sect_pr.insert(0, parse_xml(
        '<w:footerReference w:type="default" r:id="rIdFooter1" '
        + w_only + " " + r_only + "/>"))

    buf = io.BytesIO()
    d.save(buf)
    base = buf.getvalue()

    zin = zipfile.ZipFile(io.BytesIO(base))
    names = zin.namelist()
    ct = zin.read("[Content_Types].xml").decode("utf-8")
    rels = zin.read("word/_rels/document.xml.rels").decode("utf-8")
    pfx = "application/vnd.openxmlformats-officedocument.wordprocessingml."
    add_ct = (
        '<Override PartName="/word/footer1.xml" ContentType="' + pfx + 'footer+xml"/>'
        '<Override PartName="/word/endnotes.xml" ContentType="' + pfx + 'endnotes+xml"/>'
        '<Override PartName="/word/comments.xml" ContentType="' + pfx + 'comments+xml"/>'
    )
    assert "</Types>" in ct
    ct = ct.replace("</Types>", add_ct + "</Types>")
    rpfx = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    # The endnotes and comments parts are related from the main part, as in
    # any package Word writes: a part nothing relates is not the document's
    # and is not read. This test first built them unrelated, which only a
    # reader that found parts by file name could pass.
    add_rel = ('<Relationship Id="rIdFooter1" Type="' + rpfx + 'footer" Target="footer1.xml"/>'
               '<Relationship Id="rIdEndnotes" Type="' + rpfx + 'endnotes" Target="endnotes.xml"/>'
               '<Relationship Id="rIdComments" Type="' + rpfx + 'comments" Target="comments.xml"/>')
    assert "</Relationships>" in rels
    rels = rels.replace("</Relationships>", add_rel + "</Relationships>")

    smartart_drawing = (
        '<w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
        '<wp:extent cx="1828800" cy="1828800"/>'
        '<wp:docPr id="30" name="FooterSmartArt"/>'
        '<a:graphic><a:graphicData '
        'uri="http://schemas.openxmlformats.org/drawingml/2006/diagram">'
        '<dgm:relIds r:dm="rIdDm" r:lo="rIdLo" r:qs="rIdQs" r:cs="rIdCs"/>'
        '</a:graphicData></a:graphic></wp:inline></w:drawing>'
    )
    chart_drawing = (
        '<w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
        '<wp:extent cx="1828800" cy="1828800"/>'
        '<wp:docPr id="31" name="EndnoteChart"/>'
        '<a:graphic><a:graphicData '
        'uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
        '<c:chart r:id="rIdEndChart"/>'
        '</a:graphicData></a:graphic></wp:inline></w:drawing>'
    )
    footer_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:ftr ' + ns_all + '><w:p><w:r>' + smartart_drawing
        + '</w:r></w:p></w:ftr>'
    )
    endnotes_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:endnotes ' + ns_all + '>'
        '<w:endnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:endnote>'
        '<w:endnote w:type="continuationSeparator" w:id="0">'
        '<w:p><w:r><w:continuationSeparator/></w:r></w:p></w:endnote>'
        '<w:endnote w:id="1"><w:p><w:r>' + chart_drawing
        + '</w:r></w:p></w:endnote></w:endnotes>'
    )
    # w:altChunk is a BLOCK-level construct -- a direct child of its
    # container (comment/footnote/endnote/body), a SIBLING of w:p, never
    # nested inside a w:r -- mirroring the existing footnote altChunk case
    # above (`_block_lines` only recognizes it as a direct child).
    comments_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:comments ' + ns_all + '>'
        '<w:comment w:id="1" w:author="Reviewer">'
        '<w:p><w:r><w:t>reviewer comment text.</w:t></w:r></w:p>'
        '<w:altChunk r:id="rIdCommentChunk"/>'
        '</w:comment>'
        '</w:comments>'
    )

    path = tmp_path / "unread_footer_endnote_comment.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            if n == "[Content_Types].xml":
                zout.writestr(n, ct)
            elif n == "word/_rels/document.xml.rels":
                zout.writestr(n, rels)
            else:
                zout.writestr(n, zin.read(n))
        zout.writestr("word/footer1.xml", footer_xml)
        zout.writestr("word/endnotes.xml", endnotes_xml)
        zout.writestr("word/comments.xml", comments_xml)
    zin.close()

    got = _extract_bytes(path.name, path.read_bytes())
    marked = [n for n in got.notes if ex.has_evidence_marker(n)]
    has_smartart = any(re.search(r"\bsmartart\b", n, re.IGNORECASE) for n in marked)
    has_chart = any(re.search(r"\bchart\b", n, re.IGNORECASE) for n in marked)
    has_altchunk = any(re.search(r"\baltchunk\b", n, re.IGNORECASE) for n in marked)
    assert has_smartart and has_chart and has_altchunk, (
        f"expected marked notes for a SmartArt drawing (in a footer), a "
        f"chart (in an endnote) and an altChunk (in a comment): "
        f"marked={marked!r}; all notes={got.notes!r}")


def test_chartex_in_real_alternatecontent_shape_with_picture_fallback_is_one_chart(tmp_path):
    """A chartEx graphic is never a bare ``graphicData`` in real Word output
    -- it is the ``mc:Choice`` of an ``mc:AlternateContent`` whose
    ``mc:Fallback`` is a plain picture (for a reader that does not
    understand chartEx). Nothing else pins this real shape: this guards
    against a fix that only matches
    a bare ``cx:chart`` graphicData and never looks inside
    ``mc:AlternateContent`` at all, which would silently double-count via
    the Fallback picture or miss the chart entirely."""
    import docx as _docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn

    ns_all = (
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
        'xmlns:cx1="http://schemas.microsoft.com/office/drawing/2015/9/8/chartex"'
    )
    fallback_picture = (
        '<w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
        '<wp:extent cx="5486400" cy="7315200"/>'
        '<wp:docPr id="40" name="ChartExFallbackPic"/>'
        '<a:graphic><a:graphicData '
        'uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<pic:pic><pic:blipFill><a:blip r:embed="rIdFallbackImg"/></pic:blipFill></pic:pic>'
        '</a:graphicData></a:graphic></wp:inline></w:drawing>'
    )
    d = _docx.Document()
    sect_pr = d.element.body.find(qn("w:sectPr"))
    sect_pr.addprevious(parse_xml(
        '<w:p ' + ns_all + '><w:r><mc:AlternateContent>'
        '<mc:Choice Requires="cx1">'
        '<w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
        '<wp:extent cx="1828800" cy="1828800"/>'
        '<wp:docPr id="41" name="ChartEx1"/>'
        '<a:graphic><a:graphicData '
        'uri="http://schemas.microsoft.com/office/drawing/2014/chartex">'
        '<cx:chart xmlns:cx="http://schemas.microsoft.com/office/drawing/2014/chartex" '
        'r:id="rIdChartEx"/>'
        '</a:graphicData></a:graphic></wp:inline></w:drawing>'
        '</mc:Choice>'
        '<mc:Fallback>' + fallback_picture + '</mc:Fallback>'
        '</mc:AlternateContent></w:r></w:p>'))

    path = tmp_path / "chartex_real_shape.docx"
    d.save(str(path))
    got = _extract_bytes(path.name, path.read_bytes())

    marked = [n for n in got.notes if ex.has_evidence_marker(n)]
    chart_notes = [n for n in marked if re.search(r"\bchart\b", n, re.IGNORECASE)]
    picture_notes = [n for n in marked if "picture" in n.lower()]
    assert len(chart_notes) == 1, (
        f"expected exactly 1 marked chart note (the chartEx, once): "
        f"{chart_notes!r}; all notes={got.notes!r}")
    count = re.search(r"(\d+)\s+chart", chart_notes[0], re.IGNORECASE)
    assert count is not None and count.group(1) == "1", (
        f"expected the chart note's own count to be 1: {chart_notes[0]!r}")
    assert not picture_notes, (
        f"the Fallback picture must not ALSO be disclosed as an unread "
        f"picture -- a reader that understands chartEx takes mc:Choice, "
        f"never mc:Fallback: {picture_notes!r}")


@pytest.mark.parametrize("mutate_sect,case_id", [
    (lambda sect_pr, qn_: sect_pr.remove(sect_pr.find(qn_("w:pgSz"))), "removed"),
    (lambda sect_pr, qn_: sect_pr.find(qn_("w:pgSz")).set(qn_("w:w"), "0"), "w_zero"),
    (lambda sect_pr, qn_: sect_pr.find(qn_("w:pgSz")).attrib.pop(qn_("w:h")), "h_missing"),
], ids=["pgsz_removed", "pgsz_w_zero", "pgsz_h_missing"])
def test_unusable_pgsz_variants_all_disclose_unmeasurable_pictures(
        tmp_path, mutate_sect, case_id):
    """The test below covers only a REMOVED w:pgSz. An
    unusable one (w:w="0", or a missing w:h) must be treated the same way,
    not silently pass the "has a pgSz element" check and then divide by
    zero or KeyError -- or, worse, silently skip disclosure."""
    import docx as _docx
    from docx.oxml.ns import qn
    from docx.shared import Inches
    from lxml import etree

    d = _docx.Document()
    d.add_paragraph("BEFORE the unmeasurable picture.")
    pixel_buf = io.BytesIO()
    from PIL import Image
    Image.new("RGB", (1, 1), (0, 0, 0)).save(pixel_buf, format="PNG")
    d.add_picture(io.BytesIO(pixel_buf.getvalue()), width=Inches(6), height=Inches(8))
    d.add_paragraph("AFTER the unmeasurable picture.")

    buf = io.BytesIO()
    d.save(buf)
    base = buf.getvalue()

    zin = zipfile.ZipFile(io.BytesIO(base))
    names = zin.namelist()
    root = etree.fromstring(zin.read("word/document.xml"))
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    body = next(c for c in root if c.tag == w + "body")
    sect_pr = next(c for c in reversed(list(body)) if c.tag == w + "sectPr")
    mutate_sect(sect_pr, qn)
    new_doc_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                                 standalone=True)

    path = tmp_path / f"unusable_pgsz_{case_id}.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, new_doc_xml if n == "word/document.xml" else zin.read(n))
    zin.close()

    got = _extract_bytes(path.name, path.read_bytes())
    marked = [n for n in got.notes if ex.has_evidence_marker(n)]
    unmeasurable = [n for n in marked if "could not be measured" in n.lower()]
    assert unmeasurable, (
        f"expected a marked note saying pictures could not be measured "
        f"against the page ({case_id}): marked={marked!r}; all "
        f"notes={got.notes!r}")


def test_no_usable_pgsz_and_no_picture_discloses_nothing_unmeasurable(tmp_path):
    """The negative case the parametrized test above cannot cover: no usable
    ``w:pgSz`` AND no picture at all must not, by itself, produce an
    "unmeasurable" note -- there is nothing to measure. Guards against a fix
    that notes the page's own unmeasurability regardless of whether any
    picture exists."""
    import docx as _docx
    from docx.oxml.ns import qn
    from lxml import etree

    d = _docx.Document()
    d.add_paragraph("No pictures anywhere in this document.")

    buf = io.BytesIO()
    d.save(buf)
    base = buf.getvalue()

    zin = zipfile.ZipFile(io.BytesIO(base))
    names = zin.namelist()
    root = etree.fromstring(zin.read("word/document.xml"))
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    body = next(c for c in root if c.tag == w + "body")
    sect_pr = next(c for c in reversed(list(body)) if c.tag == w + "sectPr")
    sect_pr.remove(sect_pr.find(qn("w:pgSz")))
    new_doc_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                                 standalone=True)

    path = tmp_path / "no_pgsz_no_picture.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, new_doc_xml if n == "word/document.xml" else zin.read(n))
    zin.close()

    got = _extract_bytes(path.name, path.read_bytes())
    marked = [n for n in got.notes if ex.has_evidence_marker(n)]
    unmeasurable = [n for n in marked if "could not be measured" in n.lower()]
    assert not unmeasurable, (
        f"a document with no picture at all owes no 'could not be "
        f"measured' note, usable page size or not: {unmeasurable!r}; all "
        f"notes={got.notes!r}")


def test_a_final_section_with_no_usable_pgsz_discloses_unmeasurable_pictures(tmp_path):
    """Word spec part 9: when the final section has no usable ``w:pgSz``, large-
    picture disclosure used to be skipped SILENTLY. A marked note must say
    pictures could not be measured against the page, mirroring the PDF
    path's own wording when image geometry cannot be measured."""
    import docx as _docx
    from docx.shared import Inches
    from lxml import etree
    from PIL import Image

    d = _docx.Document()
    d.add_paragraph("BEFORE the unmeasurable picture.")
    pixel_buf = io.BytesIO()
    Image.new("RGB", (1, 1), (0, 0, 0)).save(pixel_buf, format="PNG")
    d.add_picture(io.BytesIO(pixel_buf.getvalue()), width=Inches(6), height=Inches(8))
    d.add_paragraph("AFTER the unmeasurable picture.")

    buf = io.BytesIO()
    d.save(buf)
    base = buf.getvalue()

    zin = zipfile.ZipFile(io.BytesIO(base))
    names = zin.namelist()
    root = etree.fromstring(zin.read("word/document.xml"))
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    body = next(c for c in root if c.tag == w + "body")
    sect_pr = next(c for c in reversed(list(body)) if c.tag == w + "sectPr")
    pg_sz = next((c for c in sect_pr if c.tag == w + "pgSz"), None)
    assert pg_sz is not None, "python-docx's own sectPr must carry a w:pgSz to remove"
    sect_pr.remove(pg_sz)
    new_doc_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                                 standalone=True)

    path = tmp_path / "no_pgsz.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, new_doc_xml if n == "word/document.xml" else zin.read(n))
    zin.close()

    got = _extract_bytes(path.name, path.read_bytes())
    marked = [n for n in got.notes if ex.has_evidence_marker(n)]
    unmeasurable = [n for n in marked if "could not be measured" in n.lower()]
    assert unmeasurable, (
        f"expected a marked note saying pictures could not be measured "
        f"against the page (no usable w:pgSz on the final section): "
        f"marked={marked!r}; all notes={got.notes!r}")


# ---------------------------------------------------------------------------
# Built packages, exact text: separators, wrappers, and the constructs the
# reader used to drop without a note
# ---------------------------------------------------------------------------

_RAW_NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
    'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
    'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
    'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:w16se="http://schemas.microsoft.com/office/word/2015/wordml/symex" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:o="urn:schemas-microsoft-com:office:office" '
    'xmlns:inv="urn:invented:markup-this-reader-does-not-know" '
    'mc:Ignorable="w14"'
)
_RAW_RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
_RAW_WML = "application/vnd.openxmlformats-officedocument.wordprocessingml."
_LETTER = '<w:pgSz w:w="12240" w:h="15840"/>'


def _raw_docx(body: str, *, main: str = "word/document.xml", sect: str = "",
              pg: str = _LETTER, parts: dict | None = None,
              rels: list | None = None) -> bytes:
    """A Word package built by hand: ``body`` inside ``w:body``, a final
    ``w:sectPr`` holding ``sect`` and ``pg``, and ``parts`` ({name: xml}).
    ``rels`` are the main part's relationships, ``(id, type, target)``,
    ``type`` being the last segment of the relationship type URI; a target
    beginning ``http`` or ``file:`` is external."""
    main_dir, main_base = main.rsplit("/", 1)
    kinds = {"header": "header+xml", "footer": "footer+xml",
             "footnotes": "footnotes+xml", "endnotes": "endnotes+xml",
             "settings": "settings+xml", "comments": "comments+xml"}
    overrides = [f'<Override PartName="/{main}" ContentType="{_RAW_WML}document.main+xml"/>']
    for name in (parts or {}):
        base = name.rsplit("/", 1)[-1]
        kind = next((v for k, v in kinds.items() if base.startswith(k)), None)
        if kind:
            overrides.append(f'<Override PartName="/{name}" ContentType="{_RAW_WML}{kind}"/>')
    ct = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" '
          'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/>'
          + "".join(overrides) + '</Types>')
    rel_xml = "".join(
        f'<Relationship Id="{i}" Type="{_RAW_RT}{t}" Target="{g}"'
        + (' TargetMode="External"' if g.startswith(("http", "file:")) else "") + "/>"
        for i, t, g in (rels or []))
    doc = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           f'<w:document {_RAW_NS}><w:body>{body}<w:sectPr>{sect}{pg}</w:sectPr>'
           '</w:body></w:document>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", ct)
        zf.writestr("_rels/.rels",
                    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                    f'<Relationship Id="rId1" Type="{_RAW_RT}officeDocument" Target="{main}"/>'
                    '</Relationships>')
        zf.writestr(main, doc)
        zf.writestr(f"{main_dir}/_rels/{main_base}.rels",
                    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                    + rel_xml + '</Relationships>')
        for name, xml in (parts or {}).items():
            zf.writestr(name, xml)
    return buf.getvalue()


def _p(*runs: str, ppr: str = "") -> str:
    return "<w:p>" + ppr + "".join(runs) + "</w:p>"


def _t(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def _part(tag: str, inner: str) -> str:
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:{tag} {_RAW_NS}>{inner}</w:{tag}>')


def _read(raw: bytes):
    got = ex.extract("built.docx", raw)
    assert got.pages, (got.status, got.error, got.notes)
    return got.pages[0].text, got.notes


_REV = 'w:author="Invented Reviewer" w:date="2018-05-06T00:00:00Z"'


def test_run_separators_table_joins_and_note_labels_render_exactly():
    """Word spec parts 2, 4, 6 and 8 as exact strings. The class test above
    normalizes and matches substrings, so it cannot see a separator added or
    lost: tab-stop DEFINITIONS in ``w:pPr`` came out as tab characters the
    document does not hold, a positional tab (``w:ptab``, Word's own
    three-column header layout) was dropped and joined the words beside it,
    and a deleted paragraph MARK counted as deleted text."""
    cell = lambda text, pr="": f"<w:tc>{pr}{_p(_t(text))}</w:tc>"  # noqa: E731
    body = (
        _p('<w:r><w:t>A</w:t><w:tab/><w:t>B</w:t></w:r>',
           ppr='<w:pPr><w:tabs><w:tab w:val="left" w:pos="720"/>'
               '<w:tab w:val="right" w:leader="dot" w:pos="9350"/></w:tabs></w:pPr>')
        + _p('<w:r><w:t>C</w:t><w:br/><w:t>D</w:t><w:cr/><w:t>E</w:t></w:r>')
        + _p('<w:r><w:t>F</w:t><w:noBreakHyphen/><w:t>G</w:t></w:r>')
        + _p(_t("LEFT"), '<w:r><w:ptab w:relativeTo="margin" w:alignment="right" '
                         'w:leader="none"/></w:r>', _t("RIGHT"))
        + _p(f'<w:moveFrom w:id="1" {_REV}>' + _t("GONE") + '</w:moveFrom>',
             f'<w:moveTo w:id="2" {_REV}>' + _t("MOVED") + '</w:moveTo>')
        + _p(_t("MARK"), ppr=f'<w:pPr><w:rPr><w:del w:id="3" {_REV}/></w:rPr></w:pPr>')
        + "<w:tbl><w:tr>" + cell("R1C1")
        + cell("V", '<w:tcPr><w:vMerge w:val="restart"/></w:tcPr>') + "</w:tr>"
        + "<w:tr>" + cell("R2C1") + cell("V", "<w:tcPr><w:vMerge/></w:tcPr>")
        + "</w:tr></w:tbl>"
        + _p(_t("REF"), '<w:r><w:footnoteReference w:id="1"/></w:r>'))
    footnotes = _part("footnotes",
                      '<w:footnote w:type="separator" w:id="-1">'
                      '<w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
                      '<w:footnote w:id="1">' + _p(_t("NOTE")) + '</w:footnote>')
    text, notes = _read(_raw_docx(body, parts={"word/footnotes.xml": footnotes},
                                  rels=[("rIdFn", "footnotes", "footnotes.xml")]))
    # D-56: the moved-from text names no move range, so no destination is
    # known for it and it is listed as deleted; the deleted paragraph mark
    # holds no character and lists nothing.
    assert text.split("\n") == [
        "A\tB", "C", "D", "E", "F-G", "LEFT\tRIGHT", "MOVED", "MARK",
        "R1C1\tV", "R2C1", "REF", "[footnote 1] NOTE",
        "[deleted by Invented Reviewer] GONE"], text.split("\n")
    assert not any(ex.M_WORD_TRACKED_DELETION in n for n in notes), notes


def test_rows_and_cells_inside_content_controls_and_custom_xml_are_read():
    """Word spec part 1 recurses into ``w:sdt`` and ``w:customXml``; the
    table reader took only DIRECT ``w:tr`` and ``w:tc`` children, so a row
    or a cell wrapped in either (a repeating-section content control) lost
    its text with no note."""
    cell = lambda text: f"<w:tc>{_p(_t(text))}</w:tc>"  # noqa: E731
    body = (
        "<w:tbl><w:tr>" + cell("PLAIN1") + cell("PLAIN2") + "</w:tr>"
        "<w:sdt><w:sdtPr/><w:sdtContent><w:tr>" + cell("ROWSDT") + cell("X")
        + "</w:tr></w:sdtContent></w:sdt>"
        "<w:tr>" + cell("CELLPLAIN") + "<w:sdt><w:sdtContent>" + cell("CELLSDT")
        + "</w:sdtContent></w:sdt></w:tr>"
        '<w:customXml w:element="row"><w:tr>' + cell("CUSTOMROW")
        + '<w:customXml w:element="cell">' + cell("CUSTOMCELL") + "</w:customXml>"
        "</w:tr></w:customXml></w:tbl>")
    text, _notes = _read(_raw_docx(body))
    assert text.split("\n") == ["PLAIN1\tPLAIN2", "ROWSDT\tX",
                                "CELLPLAIN\tCELLSDT", "CUSTOMROW\tCUSTOMCELL"], text


def test_a_symbol_character_is_rendered_or_disclosed():
    """``w:sym`` holds its character as a hex code. A real Unicode code point
    is text; a Wingdings code names a picture no standard table maps, so it
    keeps its place as U+FFFD (it used to vanish and join the words beside
    it) and is counted in a marked note."""
    body = _p(_t("TICK"), '<w:r><w:sym w:font="Segoe UI Symbol" w:char="2713"/></w:r>',
              _t("BOX"), '<w:r><w:sym w:font="Wingdings" w:char="F0FC"/></w:r>')
    text, notes = _read(_raw_docx(body))
    assert text == "TICK\u2713BOX\ufffd", repr(text)
    marked = [n for n in notes if ex.has_evidence_marker(n) and "symbol" in n.lower()]
    assert len(marked) == 1 and re.search(r"\b1 symbol", marked[0]), notes


def test_ruby_nested_field_codes_block_alternate_content_and_subdocuments():
    """Four constructs the reader got wrong without a note: a phonetic guide
    (``w:ruby``) ran into its base text; the cached result of a field nested
    inside another field's CODE was emitted as text; ``mc:AlternateContent``
    directly under the body was skipped whole; a ``w:subDoc`` link (a master
    document's sub-document, stored elsewhere) was ignored."""
    body = (
        _p('<w:r><w:ruby><w:rubyPr/><w:rt>' + _t("GUIDE") + '</w:rt><w:rubyBase>'
           + _t("BASE") + '</w:rubyBase></w:ruby></w:r>')
        + _p('<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
             '<w:r><w:instrText xml:space="preserve"> IF </w:instrText></w:r>'
             '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
             '<w:r><w:instrText xml:space="preserve"> MERGEFIELD Flag </w:instrText></w:r>'
             '<w:r><w:fldChar w:fldCharType="separate"/></w:r>' + _t("INNERCACHED")
             + '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
             '<w:r><w:instrText xml:space="preserve"> = "1" "YES" "NO" </w:instrText></w:r>'
             '<w:r><w:fldChar w:fldCharType="separate"/></w:r>' + _t("OUTERRESULT")
             + '<w:r><w:fldChar w:fldCharType="end"/></w:r>')
        + '<mc:AlternateContent><mc:Choice Requires="w14">' + _p(_t("BLOCKCHOICE"))
        + '</mc:Choice><mc:Fallback>' + _p(_t("BLOCKFALLBACK"))
        + '</mc:Fallback></mc:AlternateContent>'
        + _p('<w:subDoc r:id="rIdSub"/>')
        + _p(_t("END")))
    text, notes = _read(_raw_docx(
        body, rels=[("rIdSub", "subDocument", "file:///C:/Invented/sub.docx")]))
    assert text.split("\n") == ["BASE(GUIDE)", "OUTERRESULT", "BLOCKCHOICE", "", "END"], (
        text.split("\n"))
    marked = [n for n in notes if ex.has_evidence_marker(n) and "sub-document" in n.lower()]
    assert len(marked) == 1, notes


def test_a_chart_and_a_large_picture_inside_a_group_shape_are_disclosed():
    """A group shape (``wpg:wgp``) is one graphic whose own URI is neither a
    chart nor a picture; the chart and the picture it groups were never
    looked at. The reader already walked into groups for text-box text."""
    chart = ('<wpg:graphicFrame><wpg:cNvPr id="2" name="c"/><wpg:cNvFrPr/><wpg:xfrm/>'
             '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
             '<c:chart r:id="rIdC"/></a:graphicData></a:graphic></wpg:graphicFrame>')
    picture = ('<pic:pic><pic:nvPicPr><pic:cNvPr id="3" name="p"/><pic:cNvPicPr/></pic:nvPicPr>'
               '<pic:blipFill><a:blip r:embed="rIdImg"/></pic:blipFill><pic:spPr><a:xfrm>'
               '<a:off x="0" y="0"/><a:ext cx="5486400" cy="7315200"/></a:xfrm></pic:spPr>'
               '</pic:pic>')
    body = _p('<w:r><w:drawing><wp:inline><wp:extent cx="5486400" cy="7315200"/>'
              '<wp:docPr id="1" name="g"/><a:graphic><a:graphicData '
              'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup">'
              '<wpg:wgp><wpg:cNvGrpSpPr/><wpg:grpSpPr/>' + chart + picture
              + '</wpg:wgp></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')
    _text, notes = _read(_raw_docx(body))
    marked = [n for n in notes if ex.has_evidence_marker(n)]
    assert any(re.search(r"\b1 chart", n) for n in marked), notes
    assert any(re.search(r"\b1 picture", n) and "25%" in n for n in marked), notes


@pytest.mark.parametrize("case", ["negative_page_width", "picture_without_extent"])
def test_an_unmeasurable_picture_is_disclosed_as_unmeasured(case):
    """A negative page width made the page area negative, and a picture with
    no ``wp:extent`` returned early: either way a picture that may cover the
    page was neither counted nor called unmeasurable."""
    extent = "" if case == "picture_without_extent" else '<wp:extent cx="5486400" cy="7315200"/>'
    body = _p('<w:r><w:drawing><wp:inline>' + extent + '<wp:docPr id="1" name="p"/>'
              '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
              '<pic:pic><pic:blipFill><a:blip r:embed="rIdImg"/></pic:blipFill></pic:pic>'
              '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')
    pg = '<w:pgSz w:w="-12240" w:h="15840"/>' if case == "negative_page_width" else _LETTER
    _text, notes = _read(_raw_docx(body, pg=pg))
    assert any(ex.has_evidence_marker(n) and "could not be measured" in n for n in notes), notes


def test_a_hyperlink_target_is_kept_when_the_display_text_merely_contains_it():
    """Word spec part 2: the target follows the display text "when the two
    differ". A substring test dropped a target the display text contains but
    is not, so ``.../rev2`` shown over a link to ``.../rev`` lost the real
    referent."""
    body = (_p('<w:hyperlink r:id="rIdA">' + _t("http://example.invalid/docs/rev2")
               + '</w:hyperlink>')
            + _p('<w:hyperlink r:id="rIdB">' + _t("http://example.invalid/same")
                 + '</w:hyperlink>'))
    text, _notes = _read(_raw_docx(body, rels=[
        ("rIdA", "hyperlink", "http://example.invalid/docs/rev"),
        ("rIdB", "hyperlink", "http://example.invalid/same")]))
    assert text.split("\n") == [
        "http://example.invalid/docs/rev2 <http://example.invalid/docs/rev>",
        "http://example.invalid/same"], text


def test_a_header_part_no_section_displays_is_disclosed():
    """Word spec part 5 reads a first-page part only under ``w:titlePg`` and an
    even-page part only under ``w:evenAndOddHeaders``. A part with text that
    the flags leave undisplayed (Word keeps it when "Different first page" is
    unticked) was dropped with no note, and ``w:val="0"`` on either flag,
    which turns it OFF, was read as on."""
    hdr = lambda text: _part("hdr", _p(_t(text)))  # noqa: E731
    rels = [("rIdF", "header", "header1.xml"), ("rIdD", "header", "header2.xml"),
            ("rIdE", "header", "header3.xml"), ("rIdS", "settings", "settings.xml")]
    parts = {"word/header1.xml": hdr("FIRSTHIDDEN"), "word/header2.xml": hdr("DEFAULTSHOWN"),
             "word/header3.xml": hdr("EVENHIDDEN"),
             "word/settings.xml": _part("settings", '<w:evenAndOddHeaders w:val="0"/>')}
    sect = ('<w:headerReference w:type="first" r:id="rIdF"/>'
            '<w:headerReference w:type="default" r:id="rIdD"/>'
            '<w:headerReference w:type="even" r:id="rIdE"/>')
    for title_pg in ("", '<w:titlePg w:val="0"/>'):
        text, notes = _read(_raw_docx(_p(_t("BODY")), sect=sect + title_pg,
                                      pg=_LETTER, parts=parts, rels=rels))
        assert text.split("\n") == ["DEFAULTSHOWN", "BODY"], (title_pg, text)
        marked = [n for n in notes if ex.has_evidence_marker(n) and "header" in n.lower()]
        assert len(marked) == 1 and re.search(r"\b2 header/footer part", marked[0]), (
            title_pg, notes)


def test_a_package_absolute_header_target_is_read():
    """``Target="/word/header1.xml"`` is a valid OPC part name; joined onto
    ``word/`` it became ``/word/header1.xml``, matched no part, and the header
    (and any stamp in it) vanished with no note."""
    header = _part("hdr", _p(_t("ABSHEADER MNFV 000042")))
    text, _notes = _read(_raw_docx(
        _p(_t("BODYTEXT")), sect='<w:headerReference w:type="default" r:id="rIdH"/>',
        parts={"word/header1.xml": header}, rels=[("rIdH", "header", "/word/header1.xml")]))
    assert text.split("\n") == ["ABSHEADER MNFV 000042", "BODYTEXT"], text


def test_the_main_document_part_is_found_through_the_package_relationships():
    """A package names its main part in ``_rels/.rels``; ``word/document2.xml``
    is as valid as ``word/document.xml``. The reader opened the literal name,
    so a file python-docx had read became FAILED."""
    header = _part("hdr", _p(_t("HEADER-TWO")))
    footnotes = _part("footnotes", '<w:footnote w:id="1">' + _p(_t("NOTE-TWO")) + '</w:footnote>')
    text, _notes = _read(_raw_docx(
        _p(_t("MAINPARTTWO")), main="word/document2.xml",
        sect='<w:headerReference w:type="default" r:id="rIdH"/>',
        parts={"word/header1.xml": header, "word/footnotes.xml": footnotes},
        rels=[("rIdH", "header", "header1.xml"), ("rIdFn", "footnotes", "footnotes.xml")]))
    assert text.split("\n") == ["HEADER-TWO", "MAINPARTTWO", "[footnote 1] NOTE-TWO"], text


def test_fixture16_header_note_and_footer_boundaries_are_exact():
    """Word spec parts 5 and 6, pinned line for line. "Within the first 3
    lines" and "within the last 8" held under a header label line, a footer
    label line, notes moved after the footer and the first-page and default
    headers swapped, because fixture 16 has two header lines, three note
    lines and two footer lines."""
    lines = _page_text("16_word_constructs.docx").split("\n")
    assert lines[:3] == ["YAK first-page header line.", "XERUS default header line.",
                         "AARDVARK opens this synthetic construct fixture."], lines[:3]
    # D-56 puts the deletion list between the comments and the footer.
    assert lines[-6:] == [
        "[footnote 1]  SALAMANDER is the footnote text.",
        "[endnote 1]  URCHIN is the endnote text.",
        "[comment by WALRUS Reviewer] VULTURE flags this passage for review.",
        "[deleted by DocIQ fixtures] KOALA",
        "ZEBU default footer line one.", "MNFV 000777"], lines[-6:]


@pytest.mark.parametrize("even_and_odd", [False, True])
def test_a_shared_header_part_is_read_once_and_even_pages_follow_settings(even_and_odd):
    """Two sections referencing one default header part read it once (Word
    spec part 5: "each part emitted once"), and the even-page part is read
    only when settings ask for even and odd headers. Fixture 16 has one
    section and no even-page part, so neither rule was pinned."""
    hdr = lambda text: _part("hdr", _p(_t(text)))  # noqa: E731
    refs = '<w:headerReference w:type="default" r:id="rIdH"/>'
    body = (_p(_t("BODY-ONE"), ppr=f"<w:pPr><w:sectPr>{refs}{_LETTER}</w:sectPr></w:pPr>")
            + _p(_t("BODY-TWO")))
    parts = {"word/header1.xml": hdr("SHARED-HEADER"), "word/header2.xml": hdr("EVEN-HEADER")}
    rels = [("rIdH", "header", "header1.xml"), ("rIdE", "header", "header2.xml")]
    if even_and_odd:
        parts["word/settings.xml"] = _part("settings", "<w:evenAndOddHeaders/>")
        rels.append(("rIdS", "settings", "settings.xml"))
    text, _notes = _read(_raw_docx(
        body, sect=refs + '<w:headerReference w:type="even" r:id="rIdE"/>',
        parts=parts, rels=rels))
    expected = (["SHARED-HEADER", "EVEN-HEADER"] if even_and_odd else ["SHARED-HEADER"])
    assert text.split("\n") == expected + ["BODY-ONE", "BODY-TWO"], text


# ---------------------------------------------------------------------------
# D-56: tracked deletions are listed after the body. The Word spec addendum
# (2026-09-14) states the order, the grouping and each shape pinned here.
# ---------------------------------------------------------------------------

_WHEN = 'w:date="2018-05-06T00:00:00Z"'


def _del(author: str, *runs: str) -> str:
    return f'<w:del w:id="1" w:author="{author}" {_WHEN}>' + "".join(runs) + "</w:del>"


def _dt(text: str) -> str:
    return f'<w:r><w:delText xml:space="preserve">{text}</w:delText></w:r>'


def _text_box(blocks: str) -> str:
    return ('<w:r><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:inline>'
            '<wp:extent cx="914400" cy="914400"/><wp:docPr id="7" name="tb"/><a:graphic>'
            '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            '<wps:wsp><wps:txbx><w:txbxContent>' + blocks + '</w:txbxContent></wps:txbx>'
            '</wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice>'
            '<mc:Fallback><w:pict/></mc:Fallback></mc:AlternateContent></w:r>')


def test_tracked_deletions_are_listed_after_the_notes_and_before_the_footer():
    """D-56. The body reads with changes accepted. After the footnotes,
    endnotes and comments, and before the footer, each deleted passage is one
    line, ``[deleted by <author>] <text>``, in page-text order: the header
    parts' deletions, the body's, the notes' and comments', then the footer
    parts'. Consecutive deleted runs by one author in one paragraph are one
    passage (formatting splits a deletion into several w:del runs); a kept
    character, another author, or the paragraph's end starts the next one.
    A table cell's deletions are listed in the body's order, and a text
    box's after its anchoring paragraph's own."""
    cell = (f"<w:tbl><w:tr><w:tc>{_p(_t('CELL'), _del('Cy', _dt('GONECELL')))}</w:tc>"
            "</w:tr></w:tbl>")
    bold = '<w:r><w:rPr><w:b/></w:rPr><w:delText>two</w:delText></w:r>'
    body = (
        _p(_t("ALPHA "), _del("Ann", _dt("one ")), '<w:bookmarkStart w:id="0" w:name="x"/>',
           _del("Ann", bold), _t(" BETA"), _del("Bob", _dt("three")))
        + _p(_del("Ann", _dt("x")), _del("Bob", _dt("y")), _del("Ann", _dt("z")))
        + cell
        + _p(_t("ANCHOR"), _text_box(_p(_t("BOXTEXT"), _del("Dee", _dt("GONEBOX")))),
             _t(" TAIL"), _del("Ann", _dt("AFTERBOX")))
        + _p(_t("REF"), '<w:r><w:footnoteReference w:id="1"/></w:r>'))
    parts = {
        "word/header1.xml": _part("hdr", _p(_t("HDR"), _del("Ann", _dt("OLDHDR")))),
        "word/footer1.xml": _part("ftr", _p(_t("FTR"), _del("Gus", _dt("OLDFTR")))),
        "word/footnotes.xml": _part("footnotes", '<w:footnote w:id="1">'
                                    + _p(_t("NOTE"), _del("Eve", _dt("GONENOTE")))
                                    + "</w:footnote>"),
        "word/comments.xml": _part("comments", '<w:comment w:id="0" w:author="Rev">'
                                   + _p(_t("CMT"), _del("Fay", _dt("GONECMT")))
                                   + "</w:comment>"),
    }
    rels = [("rIdH", "header", "header1.xml"), ("rIdF", "footer", "footer1.xml"),
            ("rIdFn", "footnotes", "footnotes.xml"), ("rIdC", "comments", "comments.xml")]
    sect = ('<w:headerReference w:type="default" r:id="rIdH"/>'
            '<w:footerReference w:type="default" r:id="rIdF"/>')
    text, notes = _read(_raw_docx(body, sect=sect, parts=parts, rels=rels))
    assert text.split("\n") == [
        "HDR", "ALPHA  BETA", "", "CELL", "ANCHOR TAIL", "BOXTEXT", "REF",
        "[footnote 1] NOTE", "[comment by Rev] CMT",
        "[deleted by Ann] OLDHDR",
        "[deleted by Ann] one two", "[deleted by Bob] three",
        "[deleted by Ann] x", "[deleted by Bob] y", "[deleted by Ann] z",
        "[deleted by Cy] GONECELL", "[deleted by Ann] AFTERBOX", "[deleted by Dee] GONEBOX",
        "[deleted by Eve] GONENOTE", "[deleted by Fay] GONECMT",
        "[deleted by Gus] OLDFTR", "FTR"], text.split("\n")
    assert not any(ex.has_evidence_marker(n) for n in notes), notes


def test_moved_text_is_listed_only_when_its_destination_is_not_in_the_document():
    """D-56 left open where moved-from text goes. Moved-from text whose move
    names a ``w:moveToRangeStart`` the reader read is not listed: its words
    stand in the body where they were moved to, and listing them again would
    put every date and number in them into the output twice; a plain note
    counts such moves. Moved-from text with no destination (no move range,
    or a name no moveTo carries) reads, with changes accepted, as deleted,
    and is listed."""
    move_from = lambda author, text: (  # noqa: E731
        f'<w:moveFrom w:id="2" w:author="{author}" {_WHEN}>' + _t(text) + "</w:moveFrom>")
    body = (
        f'<w:moveFromRangeStart w:id="11" w:name="move1" w:author="Ann" {_WHEN}/>'
        + _p(_t("P1 "), move_from("Ann", "MOVEDTEXT"))
        + '<w:moveFromRangeEnd w:id="11"/>'
        + _p(_t("MIDDLE"))
        + f'<w:moveToRangeStart w:id="13" w:name="move1" w:author="Ann" {_WHEN}/>'
        + _p(f'<w:moveTo w:id="14" w:author="Ann" {_WHEN}>' + _t("MOVEDTEXT") + "</w:moveTo>")
        + '<w:moveToRangeEnd w:id="13"/>'
        + _p(_t("P2 "), move_from("Bob", "ORPHANMOVE"))
        + _p(_t("P3 "), f'<w:moveFromRangeStart w:id="16" w:name="move9" w:author="Cy" {_WHEN}/>',
             move_from("Cy", "NODEST"), '<w:moveFromRangeEnd w:id="16"/>'))
    text, notes = _read(_raw_docx(body))
    assert text.split("\n") == ["P1", "MIDDLE", "MOVEDTEXT", "P2", "P3",
                                "[deleted by Bob] ORPHANMOVE", "[deleted by Cy] NODEST"], (
        text.split("\n"))
    moved = [n for n in notes if "moved" in n]
    assert moved == ["1 passage(s) moved under tracked changes are shown only where "
                     "they were moved to"], notes
    assert not any(ex.has_evidence_marker(n) for n in notes), notes


def test_nested_revisions_field_results_paragraph_marks_and_rows():
    """D-56, the remaining shapes. A deletion inside an insertion, and an
    insertion inside a deletion, are deleted text by the deletion's author. A
    deleted field result is listed and its code is not. A deleted paragraph
    mark holds no character: nothing is listed and the two paragraphs keep
    their own lines. A deleted table row leaves the body, and its text is
    listed once whether only the row, or also its runs, carry the deletion.
    A line break inside a deleted passage is a space, so the label stands on
    the one line the passage occupies."""
    field = lambda *runs: "".join(runs)  # noqa: E731
    body = (
        _p(_t("A"), f'<w:ins w:id="3" w:author="Ann" {_WHEN}>' + _del("Bob", _dt("INSDEL"))
           + "</w:ins>", _t("B"))
        + _p(_t("C"), _del("Cy", f'<w:ins w:id="4" w:author="Dee" {_WHEN}>' + _dt("DELINS")
                           + "</w:ins>"), _t("D"))
        + _p(_t("F1"), _del("Ann", field(
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>',
            '<w:r><w:delInstrText xml:space="preserve"> DATE </w:delInstrText></w:r>',
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>', _dt("12 March 2019"),
            '<w:r><w:fldChar w:fldCharType="end"/></w:r>')))
        + _p('<w:r><w:fldChar w:fldCharType="begin"/></w:r>',
             '<w:r><w:instrText xml:space="preserve"> REF Clause </w:instrText></w:r>',
             '<w:r><w:fldChar w:fldCharType="separate"/></w:r>', _del("Bob", _dt("OLDRESULT")),
             _t("NEWRESULT"), '<w:r><w:fldChar w:fldCharType="end"/></w:r>')
        + _p(_t("MERGE1"), ppr=f'<w:pPr><w:rPr><w:del w:id="5" w:author="Cy" {_WHEN}/></w:rPr></w:pPr>')
        + _p(_t("MERGE2"))
        + "<w:tbl>"
        + f'<w:tr><w:trPr><w:del w:id="6" w:author="Dee" {_WHEN}/></w:trPr><w:tc>'
        + _p(_t("ROWGONE")) + "</w:tc></w:tr>"
        + f'<w:tr><w:trPr><w:del w:id="7" w:author="Eve" {_WHEN}/></w:trPr><w:tc>'
        + _p(_del("Eve", _dt("WORDROW"))) + "</w:tc></w:tr>"
        + "<w:tr><w:tc>" + _p(_t("ROWKEPT")) + "</w:tc>"
        + f'<w:tc><w:tcPr><w:cellDel w:id="8" w:author="Gus" {_WHEN}/></w:tcPr>'
        + _p(_t("CELLGONE")) + "</w:tc></w:tr></w:tbl>"
        + _p(_t("BR"), _del("Fay", '<w:r><w:delText>L1</w:delText><w:br/>'
                                    '<w:delText>L2</w:delText></w:r>')))
    text, notes = _read(_raw_docx(body))
    assert text.split("\n") == [
        "AB", "CD", "F1", "NEWRESULT", "MERGE1", "MERGE2", "ROWKEPT", "BR",
        "[deleted by Bob] INSDEL", "[deleted by Cy] DELINS", "[deleted by Ann] 12 March 2019",
        "[deleted by Bob] OLDRESULT", "[deleted by Dee] ROWGONE", "[deleted by Eve] WORDROW",
        "[deleted by Gus] CELLGONE", "[deleted by Fay] L1 L2"], text.split("\n")
    assert not any(ex.has_evidence_marker(n) for n in notes), notes


def test_a_date_or_a_stamp_in_deleted_text_is_never_the_documents_own():
    """D-56's reason for listing deletions after the body: a date in deleted
    text never becomes the first date while the body has one, and a
    stamp-shaped number in it is never read from a Bates zone. The listing
    sits above a one-line footer, inside the tail zone by position, so the
    zone itself must skip it: otherwise the deleted number refuses the
    footer's real stamp as ambiguous, or becomes the locator of a page that
    has none."""
    from dociq.identify.bates import (BatesDecision, DecisionStatus, apply_bates,
                                      detect_candidates, propose_format)
    from dociq.ingest.walker import _dated

    body = (_p(_t("Letter of 3 March 2019, invented."))
            + _p(_t("Superseded text:"), _del("Ann", _dt("issued 12 January 2018 as MNFV 000999"))))
    got = ex.extract("stamped.docx", _raw_docx(
        body, sect='<w:footerReference w:type="default" r:id="rIdF"/>',
        parts={"word/footer1.xml": _part("ftr", _p(_t("MNFV 000777")))},
        rels=[("rIdF", "footer", "footer1.xml")]))
    lines = got.pages[0].text.split("\n")
    assert lines == ["Letter of 3 March 2019, invented.", "Superseded text:",
                     "[deleted by Ann] issued 12 January 2018 as MNFV 000999", "MNFV 000777"], lines
    dates, _notes = _dated(got.pages)
    assert dates[:1] == ("2019-03-03",), dates
    doc = document("stamped.docx", got.pages)
    assert [c.raw for c in detect_candidates((doc,))] == ["MNFV 000777"]
    decision = BatesDecision(DecisionStatus.CONFIRMED, propose_format((doc,), min_pages=1).format)
    assert apply_bates((doc,), decision)[0].pages[0].bates == "MNFV 000777"

    bare = ex.extract("bare.docx", _raw_docx(_p(_t("Short note.")) + _p(_del("Ann", _dt("MNFV 000999")))))
    assert bare.pages[0].text.split("\n") == ["Short note.", "", "[deleted by Ann] MNFV 000999"], (
        bare.pages[0].text)
    bare_doc = document("bare.docx", bare.pages)
    assert detect_candidates((bare_doc,)) == ()
    assert apply_bates((bare_doc,), decision)[0].pages[0].bates is None


_DELETED_GRAPHIC_NOTE = (
    "{n} deleted drawing(s), picture(s) or object(s); the text of a text box or watermark "
    "in one is listed with the deletions, the text of a chart or SmartArt drawing is not "
    "read, and a document embedded in one is recovered, where it can be, as a child "
    "document whose note names the deletion")


def test_a_deleted_chart_is_counted_as_a_deletion_not_as_an_unread_chart():
    """A deleted picture, chart or object holds nothing to list. It is counted
    under the deletion marker, once per outermost graphic, and never as an
    unread chart the document shows. The note said any text inside them was
    listed with the deletions, which was false for a chart's own text, for
    SmartArt and for an embedded document; it says what happens to each."""
    chart = ('<w:r><w:drawing><wp:inline><wp:extent cx="1828800" cy="1828800"/>'
             '<wp:docPr id="1" name="c"/><a:graphic><a:graphicData '
             'uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
             '<c:chart r:id="rIdC"/></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')
    text, notes = _read(_raw_docx(_p(_t("KEPT"), _del("Ann", chart))))
    assert text == "KEPT", text
    marked = [n for n in notes if ex.has_evidence_marker(n)]
    assert marked == [f"{ex.M_WORD_TRACKED_DELETION}: " + _DELETED_GRAPHIC_NOTE.format(n=1)], notes

    # Once per OUTERMOST graphic, however it nests (Word review round 4, C): a
    # text box's paragraphs are read by a call of their own, and an object,
    # picture or chart inside a deleted text box used to be counted again.
    obj = '<w:r><w:object><o:OLEObject Type="Embed" ProgID="Package" r:id="rIdO"/></w:object></w:r>'
    vml_box = ('<w:r><w:pict><v:shape><v:textbox><w:txbxContent>{}</w:txbxContent>'
               '</v:textbox></v:shape></w:pict></w:r>')
    for label, body, n in [
        ("text box holding an object", _p(_t("KEPT"), _del("Ann", _text_box(_p(obj)))), 1),
        ("text box holding a chart", _p(_t("KEPT"), _del("Ann", _text_box(_p(chart)))), 1),
        ("VML text box holding an object",
         _p(_t("KEPT"), _del("Ann", vml_box.format(_p(obj)))), 1),
        ("text box in a text box holding an object",
         _p(_t("KEPT"), _del("Ann", _text_box(_p(_text_box(_p(obj)))))), 1),
        ("shown text box holding a deleted object",
         _p(_t("KEPT"), _text_box(_p(_del("Ann", obj)))), 1),
        ("a deleted text box and a deleted object beside it",
         _p(_t("KEPT"), _del("Ann", _text_box(_p(_t("x"))) + obj)), 2),
    ]:
        got = ex.extract("built.docx", _raw_docx(body, rels=[("rIdO", "oleObject",
                                                               "embeddings/x.xml")],
                                                 parts={"word/embeddings/x.xml": "INVENTED"}))
        marked = [x for x in got.notes if x.startswith(ex.M_WORD_TRACKED_DELETION)]
        assert marked == [f"{ex.M_WORD_TRACKED_DELETION}: "
                          + _DELETED_GRAPHIC_NOTE.format(n=n)], (label, got.notes)


# ---------------------------------------------------------------------------
# Parts are found through the main part's relationships, never by file name
# ---------------------------------------------------------------------------


def test_notes_comments_and_settings_are_found_through_the_main_parts_relationships():
    """Footnotes, endnotes, comments and settings were opened by their default
    file names. A valid package that names them otherwise lost the notes and
    comments with no marker and called a displayed even-page header hidden;
    stale parts left under the default names were read as the document's
    own. Unrelated parts are not the document's: not read, and a plain note
    counts the ones holding text."""
    hdr = lambda text: _part("hdr", _p(_t(text)))  # noqa: E731
    parts = {
        "word/header1.xml": hdr("HDR-DEFAULT"), "word/header2.xml": hdr("HDR-EVEN"),
        "word/footnotes2.xml": _part("footnotes", '<w:footnote w:id="2">'
                                     + _p(_t("FOOT-PANGOLIN")) + "</w:footnote>"),
        "word/endnotes7.xml": _part("endnotes", '<w:endnote w:id="3">'
                                    + _p(_t("END-QUOKKA")) + "</w:endnote>"),
        "word/comments1.xml": _part("comments", '<w:comment w:id="0" w:author="Made-Up Person">'
                                    + _p(_t("CMT-AXOLOTL")) + "</w:comment>"),
        "word/settings3.xml": _part("settings", "<w:evenAndOddHeaders/>"),
        "word/footnotes.xml": _part("footnotes", '<w:footnote w:id="9">'
                                    + _p(_t("DECOY-STALE-FOOT")) + "</w:footnote>"),
        "word/comments.xml": _part("comments", '<w:comment w:id="9" w:author="Stale">'
                                   + _p(_t("DECOY-STALE-CMT")) + "</w:comment>"),
        "word/settings.xml": _part("settings", '<w:evenAndOddHeaders w:val="0"/>'),
    }
    rels = [("rIdH", "header", "header1.xml"), ("rIdE", "header", "header2.xml"),
            ("rIdFn", "footnotes", "footnotes2.xml"), ("rIdEn", "endnotes", "endnotes7.xml"),
            ("rIdC", "comments", "comments1.xml"), ("rIdS", "settings", "settings3.xml")]
    sect = ('<w:headerReference w:type="default" r:id="rIdH"/>'
            '<w:headerReference w:type="even" r:id="rIdE"/>')
    text, notes = _read(_raw_docx(_p(_t("BODY")), sect=sect, parts=parts, rels=rels))
    assert text.split("\n") == [
        "HDR-DEFAULT", "HDR-EVEN", "BODY", "[footnote 2] FOOT-PANGOLIN",
        "[endnote 3] END-QUOKKA", "[comment by Made-Up Person] CMT-AXOLOTL"], text.split("\n")
    assert not any(ex.has_evidence_marker(n) for n in notes), notes
    assert ("2 Word part(s) holding text that no relationship in the package reaches "
            "were not read; Word does not show them") in notes, notes


def test_parts_outside_word_resolve_against_the_folder_of_the_part_naming_them():
    """Relationship targets resolve against the folder of the part that names
    them (``doc/``, not ``word/``), and a footnotes part is whatever the
    footnotes relationship names, here ``notes.xml``."""
    text, _notes = _read(_raw_docx(
        _p(_t("BODY")), main="doc/main.xml",
        sect='<w:headerReference w:type="default" r:id="rIdH"/>',
        parts={"doc/header1.xml": _part("hdr", _p(_t("HEADER-IN-DOC-FOLDER"))),
               "doc/notes.xml": _part("footnotes", '<w:footnote w:id="1">'
                                      + _p(_t("NOTE-IN-DOC-FOLDER")) + "</w:footnote>")},
        rels=[("rIdH", "header", "header1.xml"), ("rIdFn", "footnotes", "notes.xml")]))
    assert text.split("\n") == ["HEADER-IN-DOC-FOLDER", "BODY",
                                "[footnote 1] NOTE-IN-DOC-FOLDER"], text.split("\n")


# ---------------------------------------------------------------------------
# Symbols: the Symbol font mapped, anything unmapped kept in place and named
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("font,char,expected", [
    ("Symbol", "F0B1", "\u00b1"), ("Symbol", "F0B3", "\u2265"), ("Symbol", "F0B0", "\u00b0"),
    ("Symbol", "00B1", "\u00b1"), ("Symbol", "0061", "\u03b1"), ("symbol", "F0E6", "\u239b"),
    ("Symbol", "F0D2", "\u00ae"), ("Segoe UI Symbol", "2713", "\u2713"),
])
def test_a_symbol_font_character_is_rendered_in_place(font, char, expected):
    """A Symbol-font ``w:sym`` was deleted in place, so "10 ± 2 mm" read
    "102 mm", a figure the document never states, under a note that called
    the character untranslatable."""
    text, notes = _read(_raw_docx(_p(_t("Tolerance 10"),
                                     f'<w:r><w:sym w:font="{font}" w:char="{char}"/></w:r>',
                                     _t("2 mm"))))
    assert text == f"Tolerance 10{expected}2 mm", repr(text)
    assert not any("symbol" in n.lower() for n in notes), notes


def test_the_symbol_font_table_matches_the_adobe_symbol_encoding():
    """The table is checked against reportlab's ``symbol`` codec, the Adobe
    Symbol encoding: every code it maps to an ordinary character maps to the
    same one here, every code it leaves undefined is the placeholder, and each
    private-use value it uses is replaced by the character it draws."""
    import codecs

    import reportlab.pdfbase.rl_codecs as rl

    rl.RL_Codecs.register()
    codecs.lookup("symbol")
    drawn = {0xF6D9: "\u00a9", 0xF6DA: "\u00ae", 0xF6DB: "\u2122", 0xF8E6: "\u23d0",
             0xF8E7: "\u23af", 0xF8E8: "\u00ae", 0xF8E9: "\u00a9", 0xF8EA: "\u2122"}
    drawn.update({pua: chr(u) for pua, u in zip(
        range(0xF8EB, 0xF8FF),
        [0x239B, 0x239C, 0x239D, 0x23A1, 0x23A2, 0x23A3, 0x23A7, 0x23A8, 0x23A9, 0x23AA,
         0x23AE, 0x239E, 0x239F, 0x23A0, 0x23A4, 0x23A5, 0x23A6, 0x23AB, 0x23AC, 0x23AD])})
    wrong = []
    for code in list(range(0x20, 0x7F)) + list(range(0xA0, 0x100)):
        try:
            theirs = bytes([code]).decode("symbol")
        except UnicodeDecodeError:
            theirs = ex.SYMBOL_PLACEHOLDER
        if 0xE000 <= ord(theirs) <= 0xF8FF:
            theirs = drawn.get(ord(theirs), ex.SYMBOL_PLACEHOLDER)
        ours = (ex._SYMBOL_FONT_TABLE[0][code - 0x20] if code < 0x7F
                else ex._SYMBOL_FONT_TABLE[1][code - 0xA0])
        if ours != theirs:
            wrong.append((hex(code), ours, theirs))
    assert len(ex._SYMBOL_FONT_TABLE[0]) == 95 and len(ex._SYMBOL_FONT_TABLE[1]) == 96
    assert not wrong, wrong


def test_an_unmapped_symbol_keeps_its_place_and_is_disclosed():
    """A code no standard table maps (Wingdings, whose ``00FC`` is a tick and
    not u-umlaut; the Symbol font's radical extender; a surrogate) stands as
    U+FFFD where it was, never nothing, and a marked note names the count and
    the fonts. A symbol inside a field's code is code, not text."""
    sym = lambda font, char: f'<w:r><w:sym w:font="{font}" w:char="{char}"/></w:r>'  # noqa: E731
    body = (_p(_t("ITEM"), sym("Wingdings", "F0FC"), _t("DONE"), sym("Wingdings", "00FC"),
               _t("X"), sym("Symbol", "F060"), _t("Y"), sym("Invented Font", "D800"), _t("Z"))
            + _p('<w:r><w:fldChar w:fldCharType="begin"/></w:r>',
                 '<w:r><w:instrText xml:space="preserve"> QUOTE </w:instrText></w:r>',
                 sym("Symbol", "F0B1"), '<w:r><w:fldChar w:fldCharType="separate"/></w:r>',
                 _t("SHOWN"), '<w:r><w:fldChar w:fldCharType="end"/></w:r>'))
    text, notes = _read(_raw_docx(body))
    assert text.split("\n") == ["ITEM\ufffdDONE\ufffdX\ufffdY\ufffdZ", "SHOWN"], repr(text)
    marked = [n for n in notes if ex.has_evidence_marker(n) and "symbol" in n.lower()]
    assert marked == [f"{ex.M_WORD_UNREAD}: 4 symbol character(s) with no standard text "
                      "mapping (font(s): 'Invented Font', 'Symbol', 'Wingdings'), each shown as "
                      "U+FFFD"], notes


def test_a_symex_character_is_read_once_and_an_unknown_choice_falls_back():
    """``w16se:symEx`` (Word 2016+) holds its character like ``w:sym``, inside
    an ``mc:Choice`` whose Fallback repeats it; it read as nothing, with no
    note. A Choice requiring markup this reader does not understand is not
    taken: its Fallback is what the markup-compatibility rules give such a
    reader."""
    body = (_p(_t("Status "), '<mc:AlternateContent><mc:Choice Requires="w16se"><w:r>'
               '<w16se:symEx w16se:font="Segoe UI Emoji" w16se:char="2705"/></w:r></mc:Choice>'
               '<mc:Fallback><w:r><w:t>\u2705</w:t></w:r></mc:Fallback></mc:AlternateContent>',
               _t(" accepted"))
            + _p('<mc:AlternateContent><mc:Choice Requires="inv"><w:r><w:t>CHOICE-UNKNOWN</w:t>'
                 '</w:r></mc:Choice><mc:Fallback><w:r><w:t>FALLBACK-READ</w:t></w:r>'
                 '</mc:Fallback></mc:AlternateContent>'))
    text, notes = _read(_raw_docx(body))
    assert text.split("\n") == ["Status \u2705 accepted", "FALLBACK-READ"], text.split("\n")
    assert not any(ex.has_evidence_marker(n) for n in notes), notes


# ---------------------------------------------------------------------------
# Text kept outside w:t: watermarks and legacy form fields; rows in alternate
# content; header parts no section displays
# ---------------------------------------------------------------------------

_WATERMARK = ('<w:r><w:pict><v:shapetype id="_x0000_t136" coordsize="21600,21600"/>'
              '<v:shape id="PowerPlusWaterMarkObject1" type="#_x0000_t136" '
              'style="position:absolute;rotation:315">'
              '<v:textpath style="font-family:Calibri" string="{text}"/></v:shape></w:pict></w:r>')


def test_a_watermark_is_read_as_its_own_line_after_its_anchor():
    """A VML watermark keeps its text only in ``v:textpath/@string``; a DRAFT or
    SUPERSEDED stamp vanished with no note. It is read as a line of its own
    directly after the paragraph that anchors it, as a text box is."""
    text, notes = _read(_raw_docx(
        _p(_t("BODYTEXT")), sect='<w:headerReference w:type="default" r:id="rIdH"/>',
        parts={"word/header1.xml": _part("hdr", _p(_t("HDRTEXT"), _WATERMARK.format(
            text="DRAFT-NOT-FOR-CONSTRUCTION")) + _p(_t("HDRLINE2")))},
        rels=[("rIdH", "header", "header1.xml")]))
    assert text.split("\n") == ["HDRTEXT", "DRAFT-NOT-FOR-CONSTRUCTION", "HDRLINE2",
                                "BODYTEXT"], text.split("\n")
    assert not any(ex.has_evidence_marker(n) for n in notes), notes


@pytest.mark.parametrize("hidden_part", ["watermark_only", "deletion_only", "unreadable"])
def test_a_header_part_no_section_displays_counts_whatever_content_it_holds(hidden_part):
    """Whether an undisplayed header part holds content is decided by the
    reader itself, so it can never be narrower than what the reader reads: it
    counted ``w:t`` only, and a hidden part holding only a watermark or only
    deleted text was not disclosed. A part that will not parse counts too."""
    content = {"watermark_only": _part("hdr", _p(_WATERMARK.format(text="HIDDENWM"))),
               "deletion_only": _part("hdr", _p(_del("Ann", _dt("HIDDENDEL")))),
               "unreadable": "<w:hdr this is not xml"}[hidden_part]
    text, notes = _read(_raw_docx(
        _p(_t("BODYTEXT")),
        sect=('<w:headerReference w:type="first" r:id="rIdF"/>'
              '<w:headerReference w:type="default" r:id="rIdH"/>'),
        parts={"word/header1.xml": _part("hdr", _p(_t("DEFAULTHDR"))),
               "word/coverpage.xml": content},
        rels=[("rIdH", "header", "header1.xml"), ("rIdF", "header", "coverpage.xml")]))
    assert text.split("\n") == ["DEFAULTHDR", "BODYTEXT"], text.split("\n")
    marked = [n for n in notes if ex.has_evidence_marker(n) and "header/footer part" in n]
    assert len(marked) == 1 and re.search(r"\b1 header/footer part", marked[0]), notes


@pytest.mark.parametrize("first_footer_type", ["footer", "invented"])
def test_header_and_footer_parts_no_section_references_are_disclosed(first_footer_type):
    """A header part the main part relates but no section references, and a
    first-page footer without ``w:titlePg``, are each a part no section
    displays; both were dropped with no note (the footer because only header
    references were counted). A footer reference is counted even when its
    relationship is not typed as a footer."""
    text, notes = _read(_raw_docx(
        _p(_t("BODYTEXT")),
        sect=('<w:footerReference w:type="default" r:id="rIdD"/>'
              '<w:footerReference w:type="first" r:id="rIdFF"/>'),
        parts={"word/header1.xml": _part("hdr", _p(_t("RELATEDHEADER"))),
               "word/footer1.xml": _part("ftr", _p(_t("SHOWNFOOTER"))),
               "word/footer2.xml": _part("ftr", _p(_t("FIRSTFOOTER")))},
        rels=[("rIdH", "header", "header1.xml"), ("rIdD", "footer", "footer1.xml"),
              ("rIdFF", first_footer_type, "footer2.xml")]))
    assert text.split("\n") == ["BODYTEXT", "SHOWNFOOTER"], text.split("\n")
    marked = [n for n in notes if ex.has_evidence_marker(n) and "header/footer part" in n]
    assert len(marked) == 1 and re.search(r"\b2 header/footer part", marked[0]), notes


def test_legacy_form_field_values_and_rows_inside_alternate_content_are_read():
    """A legacy check box or drop-down shows its value from ``w:ffData``, not
    from a result run, and read as nothing ("Approved:  by engineer"). A
    table row wrapped in ``mc:AlternateContent`` was dropped, as rows in
    content controls once were."""
    def checkbox(default: str, checked: str = "") -> str:
        return ('<w:r><w:fldChar w:fldCharType="begin"><w:ffData><w:name w:val="Check1"/>'
                f'<w:enabled/><w:checkBox><w:sizeAuto/><w:default w:val="{default}"/>{checked}'
                '</w:checkBox></w:ffData></w:fldChar></w:r>'
                '<w:r><w:instrText xml:space="preserve"> FORMCHECKBOX </w:instrText></w:r>'
                '<w:r><w:fldChar w:fldCharType="end"/></w:r>')
    dropdown = ('<w:r><w:fldChar w:fldCharType="begin"><w:ffData><w:name w:val="Drop1"/>'
                '<w:ddList><w:result w:val="1"/><w:listEntry w:val="PENDING"/>'
                '<w:listEntry w:val="APPROVED"/></w:ddList></w:ffData></w:fldChar></w:r>'
                '<w:r><w:instrText xml:space="preserve"> FORMDROPDOWN </w:instrText></w:r>'
                '<w:r><w:fldChar w:fldCharType="end"/></w:r>')
    cell = lambda text: f"<w:tc>{_p(_t(text))}</w:tc>"  # noqa: E731
    body = (_p(_t("Approved: "), checkbox("1"), _t(" by engineer"))
            + _p(_t("Rejected: "), checkbox("1", '<w:checked w:val="0"/>'), _t(" end"))
            + _p(_t("Status: "), dropdown)
            + "<w:tbl><w:tr>" + cell("ROW1") + "</w:tr>"
            + '<mc:AlternateContent><mc:Choice Requires="w14"><w:tr>' + cell("ROWCHOICE")
            + "</w:tr></mc:Choice><mc:Fallback><w:tr>" + cell("ROWFALLBACK")
            + "</w:tr></mc:Fallback></mc:AlternateContent></w:tbl>")
    text, _notes = _read(_raw_docx(body))
    assert text.split("\n") == ["Approved: \u2612 by engineer", "Rejected: \u2610 end",
                                "Status: APPROVED", "ROW1", "ROWCHOICE"], text.split("\n")


def test_a_field_code_running_on_past_a_paragraph_mark_stays_code():
    """A field's code can continue into the next paragraph. The field state
    ended at each paragraph, so a nested field's cached result inside the
    outer code came out as text; a code that never ends is disclosed."""
    body = (_p('<w:r><w:fldChar w:fldCharType="begin"/></w:r>',
               '<w:r><w:instrText xml:space="preserve"> IF 1 = 1 "</w:instrText></w:r>')
            + _p('<w:r><w:fldChar w:fldCharType="begin"/></w:r>',
                 '<w:r><w:instrText xml:space="preserve"> DOCPROPERTY Title </w:instrText></w:r>',
                 '<w:r><w:fldChar w:fldCharType="separate"/></w:r>', _t("INNERCACHED"),
                 '<w:r><w:fldChar w:fldCharType="end"/></w:r>',
                 '<w:r><w:instrText xml:space="preserve">" "" </w:instrText></w:r>',
                 '<w:r><w:fldChar w:fldCharType="separate"/></w:r>', _t("OUTERRESULT"),
                 '<w:r><w:fldChar w:fldCharType="end"/></w:r>')
            + _p(_t("AFTER")))
    text, notes = _read(_raw_docx(body))
    assert text.split("\n") == ["", "OUTERRESULT", "AFTER"] or text.split("\n") == [
        "OUTERRESULT", "AFTER"], text.split("\n")
    assert not any(ex.has_evidence_marker(n) for n in notes), notes
    open_code = _p('<w:r><w:fldChar w:fldCharType="begin"/></w:r>',
                   '<w:r><w:instrText xml:space="preserve"> TOC </w:instrText></w:r>') + _p(_t("LOST"))
    _text, notes = _read(_raw_docx(_p(_t("BEFORE")) + open_code))
    assert [n for n in notes if ex.has_evidence_marker(n)] == [
        f"{ex.M_WORD_UNREAD}: 1 field code(s) never ended; the text after each, to the end "
        "of its story (the body, a table cell, a text box, a header or footer, a note or a "
        "comment), was read as field code and left out"], notes


# ---------------------------------------------------------------------------
# Pictures measured at the size they are drawn; the page area and threshold
# ---------------------------------------------------------------------------


def _group_drawing(children: str, *, ext=(5486400, 7315200), ch_ext=None) -> str:
    xfrm = ""
    if ch_ext is not None:
        xfrm = (f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{ext[0]}" cy="{ext[1]}"/>'
                f'<a:chOff x="0" y="0"/><a:chExt cx="{ch_ext[0]}" cy="{ch_ext[1]}"/></a:xfrm>')
    return (f'<w:r><w:drawing><wp:inline><wp:extent cx="{ext[0]}" cy="{ext[1]}"/>'
            '<wp:docPr id="1" name="g"/><a:graphic><a:graphicData '
            'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup">'
            f'<wpg:wgp><wpg:cNvGrpSpPr/><wpg:grpSpPr>{xfrm}</wpg:grpSpPr>' + children
            + '</wpg:wgp></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')


def _group_picture(cx: int, cy: int) -> str:
    return ('<pic:pic><pic:nvPicPr><pic:cNvPr id="3" name="p"/><pic:cNvPicPr/></pic:nvPicPr>'
            '<pic:blipFill><a:blip r:embed="rIdImg"/></pic:blipFill><pic:spPr><a:xfrm>'
            f'<a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm></pic:spPr></pic:pic>')


@pytest.mark.parametrize("child_scale,picture_scale,nested,large", [
    (1.0, 1.0, False, True), (0.1, 0.1, False, True), (10.0, 1.0, False, False),
    (0.5, 0.5, True, True)])
def test_a_picture_inside_a_resized_group_is_measured_at_its_drawn_size(
        child_scale, picture_scale, nested, large):
    """A group draws its child coordinate space (``chExt``) at its extent
    (6x8in here). A picture filling a child space ten times smaller than the
    group is drawn at 6x8in, and went undisclosed; a 6x8in-in-child-units
    picture in a child space ten times larger is drawn at 0.6x0.8in, and was
    reported. A nested group compounds its own ``ext / chExt``."""
    ext = (5486400, 7315200)
    ch = (int(ext[0] * child_scale), int(ext[1] * child_scale))
    pic = (int(ext[0] * picture_scale), int(ext[1] * picture_scale))
    if nested:
        inner = ('<wpg:grpSp><wpg:grpSpPr><a:xfrm><a:off x="0" y="0"/>'
                 f'<a:ext cx="{pic[0]}" cy="{pic[1]}"/><a:chOff x="0" y="0"/>'
                 f'<a:chExt cx="{pic[0] // 4}" cy="{pic[1] // 4}"/></a:xfrm></wpg:grpSpPr>'
                 + _group_picture(pic[0] // 4, pic[1] // 4) + "</wpg:grpSp>")
    else:
        inner = _group_picture(*pic)
    _text, notes = _read(_raw_docx(_p(_group_drawing(inner, ext=ext, ch_ext=ch))))
    marked = [n for n in notes if ex.has_evidence_marker(n) and "picture" in n]
    assert bool(marked) is large, (child_scale, picture_scale, nested, notes)


def test_the_page_area_is_the_final_sections_and_a_quarter_page_picture_is_disclosed():
    """The large-picture share is measured against the FINAL section's page,
    not the first section's, and a picture covering exactly 25% is disclosed
    (the threshold is at least 25%)."""
    quarter = ('<w:r><w:drawing><wp:inline><wp:extent cx="3886200" cy="5029200"/>'
               '<wp:docPr id="1" name="p"/><a:graphic><a:graphicData '
               'uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic>'
               '<pic:blipFill><a:blip r:embed="rIdImg"/></pic:blipFill></pic:pic>'
               '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')
    huge_first = '<w:pPr><w:sectPr><w:pgSz w:w="144000" w:h="144000"/></w:sectPr></w:pPr>'
    _text, notes = _read(_raw_docx(_p(_t("SECTION ONE"), ppr=huge_first) + _p(quarter)))
    marked = [n for n in notes if ex.has_evidence_marker(n) and "picture" in n]
    assert len(marked) == 1 and re.search(r"\b1 picture", marked[0]), notes


@pytest.mark.parametrize("value", ["false", "off", "0"])
def test_an_on_off_flag_written_false_or_off_turns_the_header_off(value):
    """``w:titlePg`` and ``w:evenAndOddHeaders`` are off when their ``w:val``
    is ``0``, ``false`` or ``off``."""
    hdr = lambda text: _part("hdr", _p(_t(text)))  # noqa: E731
    text, _notes = _read(_raw_docx(
        _p(_t("BODY")),
        sect=('<w:headerReference w:type="first" r:id="rIdF"/>'
              '<w:headerReference w:type="default" r:id="rIdD"/>'
              '<w:headerReference w:type="even" r:id="rIdE"/>'
              f'<w:titlePg w:val="{value}"/>'),
        parts={"word/header1.xml": hdr("FIRST"), "word/header2.xml": hdr("DEFAULT"),
               "word/header3.xml": hdr("EVEN"),
               "word/settings.xml": _part("settings", f'<w:evenAndOddHeaders w:val="{value}"/>')},
        rels=[("rIdF", "header", "header1.xml"), ("rIdD", "header", "header2.xml"),
              ("rIdE", "header", "header3.xml"), ("rIdS", "settings", "settings.xml")]))
    assert text.split("\n") == ["DEFAULT", "BODY"], text.split("\n")


def test_a_picture_with_a_negative_extent_is_unmeasured():
    body = _p('<w:r><w:drawing><wp:inline><wp:extent cx="-5486400" cy="7315200"/>'
              '<wp:docPr id="1" name="p"/><a:graphic><a:graphicData '
              'uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic>'
              '<pic:blipFill><a:blip r:embed="rIdImg"/></pic:blipFill></pic:pic>'
              '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')
    _text, notes = _read(_raw_docx(body))
    assert any(ex.has_evidence_marker(n) and "could not be measured" in n for n in notes), notes


def test_fixture16_text_box_stands_on_its_own_line_after_its_anchor():
    """Word spec part 3, line for line: the box's text is not inlined into the
    anchoring paragraph's line."""
    lines = _page_text("16_word_constructs.docx").split("\n")
    raven = [i for i, line in enumerate(lines) if "RAVEN" in line]
    quail = [i for i, line in enumerate(lines) if "QUAIL" in line]
    assert len(raven) == 1 and len(quail) == 1 and raven[0] == quail[0] + 1, (
        [lines[i] for i in quail + raven])
    assert "RAVEN" not in lines[quail[0]] and "QUAIL" not in lines[raven[0]], lines


def test_the_true_page_note_is_on_every_word_record_and_carries_no_marker():
    """Word spec part 7: the page note describes the page model on the record and
    on its page, and is not an evidence marker."""
    got = _pages("16_word_constructs.docx")
    assert got.notes[0] == ex.WORD_LAYOUT_NOTE and got.pages[0].notes[0] == ex.WORD_LAYOUT_NOTE, (
        got.notes)
    assert not ex.has_evidence_marker(ex.WORD_LAYOUT_NOTE)


# ---------------------------------------------------------------------------
# Review-fix round 3: characters in ordinary runs drawn in a symbol font
# (addendum A2.4)
# ---------------------------------------------------------------------------

_W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _font_run(text: str, rpr: str = "", tag: str = "t") -> str:
    return f'<w:r><w:rPr>{rpr}</w:rPr><w:{tag} xml:space="preserve">{text}</w:{tag}></w:r>'


def _fonts(**slots) -> str:
    return "<w:rFonts " + " ".join(f'w:{k}="{v}"' for k, v in slots.items()) + "/>"


def _styles(*styles: str, defaults: str = "") -> str:
    head = (f"<w:docDefaults><w:rPrDefault><w:rPr>{defaults}</w:rPr></w:rPrDefault>"
            "</w:docDefaults>" if defaults else "")
    return _part("styles", head + "".join(styles))


def _theme(minor_latin: str) -> str:
    fonts = ('<a:latin typeface="{}"/><a:ea typeface=""/><a:cs typeface=""/>')
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="t">'
            '<a:themeElements><a:fontScheme name="f"><a:majorFont>' + fonts.format("Arial")
            + '</a:majorFont><a:minorFont>' + fonts.format(minor_latin)
            + '</a:minorFont></a:fontScheme></a:themeElements></a:theme>')


_DINGS = _fonts(ascii="Wingdings", hAnsi="Wingdings")
_SYM = _fonts(ascii="Symbol", hAnsi="Symbol")
_STYLE_RELS = [("rIdS", "styles", "styles.xml"), ("rIdT", "theme", "theme/theme1.xml")]
_SYMBOL_RUN_CASES = {
    "symbol_direct_micro": (_p(_t("Coating 250 "), _font_run("m", _SYM), _t("m thick")),
                            {}, "Coating 250 µm thick", ""),
    "symbol_private_use_form": (_p(_t("Pressure "), _font_run("&#xF0B3;", _SYM), _t(" 40 bar")),
                                {}, "Pressure ≥ 40 bar", ""),
    "wingdings_latin1_arrow": (_p(_t("Next "), _font_run("&#xE0;", _DINGS), _t(" step")),
                               {}, "Next � step", "1:'Wingdings'"),
    "wingdings_ascii_slot_only": (_p(_t("Thanks "), _font_run("J", _fonts(ascii="Wingdings"))),
                                  {}, "Thanks �", "1:'Wingdings'"),
    "wingdings_private_use_and_spaces": (
        _p(_t("Passed "), _font_run("&#xF0FC; &#xFC;", _DINGS)), {}, "Passed � �",
        "2:'Wingdings'"),
    "zapf_dingbats_tick": (_p(_t("Checked "), _font_run("3", _fonts(ascii="ZapfDingbats"))),
                           {}, "Checked ✓", ""),
    "character_style": (_p(_t("Mark "), _font_run("&#xFC;", '<w:rStyle w:val="Dings"/>')),
                        {"word/styles.xml": _styles(
                            '<w:style w:type="character" w:styleId="Dings"><w:rPr>'
                            + _DINGS + '</w:rPr></w:style>')},
                        "Mark �", "1:'Wingdings'"),
    "paragraph_style_based_on": (
        _p(_t("a"), ppr='<w:pPr><w:pStyle w:val="Greek"/></w:pPr>'),
        {"word/styles.xml": _styles(
            '<w:style w:type="paragraph" w:styleId="Base"><w:rPr>' + _SYM + '</w:rPr></w:style>'
            '<w:style w:type="paragraph" w:styleId="Greek"><w:basedOn w:val="Base"/></w:style>')},
        "α", ""),
    "document_defaults": (_p(_t("J")), {"word/styles.xml": _styles(defaults=_DINGS)},
                          "�", "1:'Wingdings'"),
    "theme_font": (_p(_font_run("J", '<w:rFonts w:asciiTheme="minorHAnsi" w:ascii="Arial"/>')),
                   {"word/theme/theme1.xml": _theme("Wingdings")}, "�", "1:'Wingdings'"),
    "complex_script_slot_needs_rtl": (
        _p(_font_run("J", _fonts(ascii="Arial", cs="Wingdings")),
           _font_run("J", _fonts(ascii="Arial", cs="Wingdings") + "<w:rtl/>")),
        {}, "J�", "1:'Wingdings'"),
    "deleted_text_in_a_symbol_font": (
        _p(_t("Kept"), _del("Ann", _font_run("&#xE8;", _DINGS, tag="delText"))), {},
        "Kept\n[deleted by Ann] �", "1:'Wingdings'"),
    "non_breaking_hyphen_and_drop_down_in_symbol_fonts": (
        _p(_t("A"), f"<w:r><w:rPr>{_DINGS}</w:rPr><w:noBreakHyphen/></w:r>", _t("B "),
           f'<w:r><w:rPr>{_SYM}</w:rPr><w:fldChar w:fldCharType="begin"><w:ffData><w:ddList>'
           '<w:result w:val="0"/><w:listEntry w:val="a"/></w:ddList></w:ffData></w:fldChar></w:r>'
           '<w:r><w:instrText xml:space="preserve"> FORMDROPDOWN </w:instrText></w:r>'
           '<w:r><w:fldChar w:fldCharType="end"/></w:r>'),
        {}, "A�B α", "1:'Wingdings'"),
}


@pytest.mark.parametrize("case", sorted(_SYMBOL_RUN_CASES))
def test_ordinary_text_in_a_symbol_font_run_reads_as_the_font_draws_it(case):
    """Review round 3: a ``w:t`` whose run font is Symbol, Zapf Dingbats,
    Wingdings or another symbol font, set directly (any of ``w:ascii``,
    ``w:hAnsi``, ``w:eastAsia``, ``w:cs``, by the character's slot), by a
    character or paragraph style, by the document defaults or through the
    theme, was copied through as the Latin letter sharing its code: a
    Symbol-font ``m`` (a micro sign) read ``mm``, a thousand times the
    figure, and 125 Wingdings arrows in the corpus read as ``a-grave``. Each
    character now maps as a ``w:sym`` does, placeholders counted in the same
    marked note; a space stays a space."""
    body, parts, expected, disclosed = _SYMBOL_RUN_CASES[case]
    rels = [rel for rel in _STYLE_RELS if f"word/{rel[2]}" in parts]
    text, notes = _read(_raw_docx(body, parts=parts, rels=rels))
    assert text == expected, (case, text)
    marked = [n for n in notes if ex.has_evidence_marker(n) and "symbol" in n]
    count, _sep, fonts = disclosed.partition(":")
    assert marked == ([f"{ex.M_WORD_UNREAD}: {count} symbol character(s) with no standard text "
                       f"mapping (font(s): {fonts}), each shown as U+FFFD"] if disclosed else []), (
        case, notes)


def _font_run(font: str, text: str) -> str:
    return (f'<w:r><w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}"/></w:rPr>'
            f'<w:t xml:space="preserve">{text}</w:t></w:r>')


@pytest.mark.parametrize("font,micro,umlaut,disclosed", [
    ("Symbol", "\u00b5", "\u23ab", None),
    ("SymbolMT", "\u00b5", "\u23ab", None),
    ("Symbol MT", "\u00b5", "\u23ab", None),
    ("symbol-regular", "\u00b5", "\u23ab", None),
    ("ZapfDingbatsITC", "\u274d", "\u27bc", None),
    ("ITC Zapf Dingbats", "\u274d", "\u27bc", None),
    ("Wingdings-Regular", "\ufffd", "\ufffd", "Wingdings-Regular"),
    ("Wingdings2", "\ufffd", "\ufffd", "Wingdings2"),
    ("MTExtra", "\ufffd", "\ufffd", "MTExtra"),
    ("Invented Dingbats", "\ufffd", "\ufffd", "Invented Dingbats"),
    ("Invented Serif", "m", "\u00fc", None),
    ("ArialMT", "m", "\u00fc", None),
])
def test_a_symbol_font_is_known_by_its_aliases_and_by_the_files_own_font_table(
        font, micro, umlaut, disclosed):
    """Word review round 4, C: symbol fonts were known by a closed list of
    names, so the PostScript and foundry names a converted file carries
    (``SymbolMT``, ``ZapfDingbatsITC``, ``Wingdings-Regular``) and a font the
    file's own font table declares symbol-charset (``w:charset w:val="02"``)
    read as Latin letters with no note. A name is now folded to its family
    (case, spaces, hyphens, and ``MT``/``ITC``/``Regular`` affixes ignored),
    and a declared symbol-charset font is read as Wingdings is. A text font
    declared with another charset, or an ``MT`` text font, is untouched."""
    table = _part("fonts", "".join(
        f'<w:font w:name="{name}"><w:charset w:val="{cs}"/><w:family w:val="auto"/></w:font>'
        for name, cs in (("Invented Dingbats", "02"), ("Invented Serif", "00"),
                         ("Symbol", "02"), ("Wingdings2", "02"))))
    body = _p(_font_run("Arial", "dose 5 "), _font_run(font, "m"), _font_run("Arial", "g"),
              _font_run(font, " \u00fc"))
    text, notes = _read(_raw_docx(body, parts={"word/fontTable.xml": table},
                                  rels=[("rIdFT", "fontTable", "fontTable.xml")]))
    assert text == f"dose 5 {micro}g {umlaut}", (font, text)
    marked = [n for n in notes if ex.has_evidence_marker(n) and "symbol" in n]
    assert marked == ([f"{ex.M_WORD_UNREAD}: 2 symbol character(s) with no standard text "
                       f"mapping (font(s): '{disclosed}'), each shown as U+FFFD"]
                      if disclosed else []), (font, notes)


def test_without_a_font_table_an_undeclared_font_reads_as_its_characters():
    """The charset rule reads only what the file declares: the same run in
    ``Invented Dingbats`` with no font table part is its characters."""
    body = _p(_font_run("Invented Dingbats", "m"))
    text, notes = _read(_raw_docx(body))
    assert text == "m" and not [n for n in notes if "symbol" in n], (text, notes)


def test_the_zapf_dingbats_table_matches_reportlabs_codec():
    """Zapf Dingbats was disclosed as having no standard mapping, though
    reportlab, a declared dependency, ships one: the table here is held equal
    to its ``zapfdingbats`` codec, U+FFFD where the codec defines nothing."""
    import codecs

    import reportlab.pdfbase.rl_codecs as rl

    rl.RL_Codecs.register()
    codecs.lookup("zapfdingbats")
    wrong = []
    for code in range(0x20, 0x100):
        try:
            theirs = bytes([code]).decode("zapfdingbats")
        except UnicodeDecodeError:
            theirs = ex.SYMBOL_PLACEHOLDER
        if ex._ZAPF_DINGBATS_TABLE[code - 0x20] != theirs:
            wrong.append((hex(code), ex._ZAPF_DINGBATS_TABLE[code - 0x20], theirs))
    assert len(ex._ZAPF_DINGBATS_TABLE) == 224 and not wrong, wrong


def test_a_zapf_dingbats_symbol_maps_and_monotype_sorts_stays_disclosed():
    """``w:sym`` in Zapf Dingbats maps through its table in either code form;
    Monotype Sorts, which no dependency's table names, stays a counted
    placeholder, as does a control code, which would otherwise break the
    line (a ``w:sym`` of ``000A`` split one)."""
    sym = lambda font, char: f'<w:r><w:sym w:font="{font}" w:char="{char}"/></w:r>'  # noqa: E731
    text, notes = _read(_raw_docx(_p(
        _t("A"), sym("ZapfDingbats", "F033"), _t("B"), sym("ITC Zapf Dingbats", "0034"),
        _t("C"), sym("Monotype Sorts", "0033"), _t("D"), sym("Arial", "000A"), _t("E"))))
    assert text == "A✓B✔C�D�E", repr(text)
    assert [n for n in notes if ex.has_evidence_marker(n)] == [
        f"{ex.M_WORD_UNREAD}: 2 symbol character(s) with no standard text mapping (font(s): "
        "'Arial', 'Monotype Sorts'), each shown as U+FFFD"], notes


def test_a_watermark_set_in_a_symbol_font_reads_as_the_font_draws_it():
    """A VML watermark's text is drawn in the font its style names; in
    Wingdings its letters are pictures."""
    mark = _WATERMARK.format(text="J").replace("font-family:Calibri", "font-family:&quot;Wingdings&quot;")
    text, notes = _read(_raw_docx(
        _p(_t("BODYTEXT")), sect='<w:headerReference w:type="default" r:id="rIdH"/>',
        parts={"word/header1.xml": _part("hdr", _p(_t("HDR"), mark))},
        rels=[("rIdH", "header", "header1.xml")]))
    assert text.split("\n") == ["HDR", "�", "BODYTEXT"], text.split("\n")
    assert [n for n in notes if ex.has_evidence_marker(n)] == [
        f"{ex.M_WORD_UNREAD}: 1 symbol character(s) with no standard text mapping (font(s): "
        "'Wingdings'), each shown as U+FFFD"], notes


# ---------------------------------------------------------------------------
# Review-fix round 3: values from the file never start a page line
# ---------------------------------------------------------------------------


def test_a_value_the_file_supplies_to_a_label_or_target_never_starts_a_page_line():
    """Deletion authors had their line breaks read as spaces; a comment's
    author, a note's id and a hyperlink's target did not, so a file could put
    a stamp-shaped line of its own choosing on the page, inside the Bates
    tail zone. Every value DocIQ writes into a label or beside page text is
    now one line."""
    from dociq.identify.bates import detect_candidates

    body = (_p(_t("Body."), '<w:r><w:footnoteReference w:id="1"/></w:r>')
            + _p('<w:hyperlink r:id="rIdL">' + _t("the register") + "</w:hyperlink>")
            + _p(f'<w:del w:id="1" w:author="Ann&#10;MNFV 000999" {_WHEN}>' + _dt("x") + "</w:del>"))
    parts = {"word/footnotes.xml": _part("footnotes", '<w:footnote w:id="1&#10;MNFV 000997">'
                                         + _p(_t("NOTE")) + "</w:footnote>"),
             "word/comments.xml": _part("comments", '<w:comment w:id="0" '
                                        'w:author="Rev&#13;&#10;MNFV 000998">'
                                        + _p(_t("CMT")) + "</w:comment>")}
    rels = [("rIdFn", "footnotes", "footnotes.xml"), ("rIdC", "comments", "comments.xml"),
            ("rIdL", "hyperlink", "http://example.com/&#10;MNFV 000996&#10;")]
    got = ex.extract("labels.docx", _raw_docx(body, parts=parts, rels=rels))
    lines = got.pages[0].text.split("\n")
    assert lines == ["Body.", "the register <http://example.com/ MNFV 000996>", "",
                     "[footnote 1 MNFV 000997] NOTE", "[comment by Rev MNFV 000998] CMT",
                     "[deleted by Ann MNFV 000999] x"], lines
    assert detect_candidates((document("labels.docx", got.pages),)) == ()


def test_an_email_header_value_never_starts_a_page_line(monkeypatch):
    """The same class outside Word: an email reader writes ``Subject: `` and
    the header's value after it, and an encoded word (RFC 2047) can decode to
    a line break. The value then put a line of the sender's choosing on the
    page, inside the Bates head zone. A header reads as one line, as a mail
    client shows it; so does a ``.msg`` header."""
    import extract_msg

    from dociq.identify.bates import detect_candidates

    raw = (b"From: sender@example.com\r\n"
           b"Subject: =?utf-8?q?Invented_report=0AMNFV_000777?=\r\n"
           b"Date: Sun, 03 Mar 2019 09:00:00 +0000\r\n\r\nBody text.\r\n")
    got = ex.extract("mail.eml", raw)
    lines = got.pages[0].text.split("\n")
    assert lines[:3] == ["From: sender@example.com", "Subject: Invented report MNFV 000777",
                         "Date: Sun, 03 Mar 2019 09:00:00 +0000 (2019-03-03)"], lines
    assert detect_candidates((document("mail.eml", got.pages),)) == ()

    class _Message:
        sender = "sender@example.com"
        to = "a@example.com\r\nMNFV 000778"
        cc = None
        subject = "Invented\nMNFV 000779"
        date = "Sun, 03 Mar 2019 09:00:00 +0000"
        body = "Body text."

        def __init__(self, _path):
            pass

    monkeypatch.setattr(extract_msg, "Message", _Message)
    got = ex.extract("mail.msg", b"not read: the reader is replaced")
    assert got.pages[0].text.split("\n") == [
        "From: sender@example.com", "To: a@example.com MNFV 000778",
        "Subject: Invented MNFV 000779", "Date: Sun, 03 Mar 2019 09:00:00 +0000", "",
        "Body text."], got.pages[0].text


# ---------------------------------------------------------------------------
# Review-fix round 3: D-56 passages, moves and field codes, stated exactly
# ---------------------------------------------------------------------------


def test_deleted_and_moved_from_text_side_by_side_by_one_author_is_one_passage():
    """Addendum A4: consecutive removed characters by one author in one
    paragraph are one passage. Deleted text followed at once by moved-from
    text with no destination, by the same author, was listed as two. Another
    author, or a kept character, still starts the next passage."""
    move_from = lambda author, text: (  # noqa: E731
        f'<w:moveFrom w:id="2" w:author="{author}" {_WHEN}>' + _t(text) + "</w:moveFrom>")
    body = (_p(_t("P1 "), _del("Ann", _dt("ALPHA ")), move_from("Ann", "BETA"),
               _del("Ann", _dt(" GAMMA")))
            + _p(_t("P2 "), _del("Ann", _dt("DELTA")), move_from("Bob", "EPSILON"))
            + _p(_t("P3 "), move_from("Cy", "ZETA"), _t(" kept "), _del("Cy", _dt("ETA"))))
    text, notes = _read(_raw_docx(body))
    assert text.split("\n") == [
        "P1", "P2", "P3  kept",
        "[deleted by Ann] ALPHA BETA GAMMA", "[deleted by Ann] DELTA", "[deleted by Bob] EPSILON",
        "[deleted by Cy] ZETA", "[deleted by Cy] ETA"], text.split("\n")
    assert not any(ex.has_evidence_marker(n) for n in notes), notes


def test_a_deleted_hyperlink_or_ruby_guide_in_several_runs_is_read_as_one_passage():
    """What a deleted hyperlink shows, and a deleted ruby guide, are read
    across every run of them since they began: the target follows the display
    text only when the two differ, and the whole guide stands in brackets.
    Read from the last run alone, a link whose display text is its target in
    two runs repeated the target, and a two-run guide bracketed only its end."""
    link = ('<w:hyperlink r:id="rIdL">' + _dt("http://example.com") + _dt("/invented")
            + "</w:hyperlink>")
    ruby = ('<w:r><w:ruby><w:rubyPr/><w:rt>' + _dt("GUIDE") + _dt("TWO") + "</w:rt><w:rubyBase>"
            + _dt("BASE") + "</w:rubyBase></w:ruby></w:r>")
    text, _notes = _read(_raw_docx(
        _p(_t("KEPT "), _del("Ann", link)) + _p(_t("ALSO "), _del("Ann", ruby)),
        rels=[("rIdL", "hyperlink", "http://example.com/invented")]))
    assert text.split("\n") == ["KEPT", "ALSO", "[deleted by Ann] http://example.com/invented",
                                "[deleted by Ann] BASE(GUIDETWO)"], text.split("\n")


def test_the_moved_note_counts_passages_and_ranges_inside_paragraphs_pair_moves():
    """The note said "1 passage(s)" for a move of three paragraphs: it counted
    move names. It counts passages. Move ranges Word writes inside the
    paragraphs (its usual shape) pair a move with its destination as well as
    ranges between blocks do; ignored, the moved sentence was listed a second
    time as deleted."""
    mf = lambda text: f'<w:moveFrom w:id="2" w:author="Ann" {_WHEN}>' + _t(text) + "</w:moveFrom>"  # noqa: E731
    mt = lambda text: f'<w:moveTo w:id="4" w:author="Ann" {_WHEN}>' + _t(text) + "</w:moveTo>"  # noqa: E731
    body = (f'<w:moveFromRangeStart w:id="1" w:name="move1" w:author="Ann" {_WHEN}/>'
            + _p(mf("ONE")) + _p(mf("TWO")) + _p(mf("THREE"))
            + '<w:moveFromRangeEnd w:id="1"/>' + _p(_t("MIDDLE"))
            + f'<w:moveToRangeStart w:id="3" w:name="move1" w:author="Ann" {_WHEN}/>'
            + _p(mt("ONE")) + _p(mt("TWO")) + _p(mt("THREE")) + '<w:moveToRangeEnd w:id="3"/>'
            + _p(_t("P1 "), f'<w:moveFromRangeStart w:id="5" w:name="move2" w:author="Ann" {_WHEN}/>',
                 mf("MOVED 3 March 2019"), '<w:moveFromRangeEnd w:id="5"/>')
            + _p(_t("P2 "), f'<w:moveToRangeStart w:id="6" w:name="move2" w:author="Ann" {_WHEN}/>',
                 mt("MOVED 3 March 2019"), '<w:moveToRangeEnd w:id="6"/>'))
    text, notes = _read(_raw_docx(body))
    # The three paragraphs moved away read as blank lines, which normalization
    # strips from the start of the page.
    assert text.split("\n") == ["MIDDLE", "ONE", "TWO", "THREE", "P1",
                                "P2 MOVED 3 March 2019"], text.split("\n")
    assert [n for n in notes if "moved" in n] == [
        "4 passage(s) moved under tracked changes are shown only where they were moved to"], notes


def test_an_open_field_code_is_counted_only_where_it_swallowed_text_in_its_own_story():
    """The note said the text after a field that never ended was left out
    "in its part". A field's state is per story: a field left open in a table
    cell loses nothing after the table, and one left open at the end of the
    body with nothing after it loses nothing; neither is counted. One left
    open in the body before a table swallows the body's later paragraphs but
    not the cell's text."""
    open_field = ('<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                  '<w:r><w:instrText xml:space="preserve"> TOC </w:instrText></w:r>')
    cell = lambda text: f"<w:tbl><w:tr><w:tc>{text}</w:tc></w:tr></w:tbl>"  # noqa: E731
    cases = {
        "open_in_a_cell": (_p(_t("FIRST")) + cell(_p(open_field)) + _p(_t("AFTER")),
                           ["FIRST", "", "AFTER"], 0),
        "open_at_the_end": (_p(_t("BEFORE")) + _p(open_field), ["BEFORE"], 0),
        "open_before_a_table": (_p(_t("FIRST"), open_field) + cell(_p(_t("CELLTEXT")))
                                + _p(_t("LOST")) + _p(_t("LOST TOO")),
                                ["FIRST", "CELLTEXT"], 1),
    }
    for case, (body, expected, swallowed) in cases.items():
        text, notes = _read(_raw_docx(body))
        assert text.split("\n") == expected, (case, text.split("\n"))
        marked = [n for n in notes if ex.has_evidence_marker(n)]
        assert marked == ([f"{ex.M_WORD_UNREAD}: 1 field code(s) never ended; the text after "
                           "each, to the end of its story (the body, a table cell, a text box, a "
                           "header or footer, a note or a comment), was read as field code and "
                           "left out"] if swallowed else []), (case, notes)


def test_run_content_standing_among_the_blocks_is_read():
    """The schema lets a ``w:del``, ``w:ins`` or math paragraph holding runs
    stand directly among a story's blocks. Everything there but paragraphs,
    tables and content controls was dropped with no note: a block-level
    deletion was not listed and a block-level insertion's text was lost."""
    body = (_p(_t("BEFORE"))
            + f'<w:del w:id="1" w:author="Ann" {_WHEN}>' + _dt("BLOCKDEL") + "</w:del>"
            + f'<w:ins w:id="2" w:author="Bob" {_WHEN}>' + _t("BLOCKINS") + "</w:ins>"
            + '<w:bookmarkStart w:id="0" w:name="b"/><w:bookmarkEnd w:id="0"/>'
            + _p(_t("AFTER")))
    text, notes = _read(_raw_docx(body))
    assert text.split("\n") == ["BEFORE", "BLOCKINS", "AFTER", "[deleted by Ann] BLOCKDEL"], (
        text.split("\n"))
    assert not any(ex.has_evidence_marker(n) for n in notes), notes


def test_an_unrelated_word_part_is_counted_by_its_root_whatever_its_content_type():
    """A part no relationship reaches is counted in the plain note when it
    holds Word text, found by its root element as well as by content type: a
    stale footnotes part typed only by the ``xml`` default was not counted. A
    thumbnail or media part is not a text part and is never counted."""
    parts = {"word/stale_notes.xml": _part("footnotes", '<w:footnote w:id="4">'
                                           + _p(_t("STALE-FOOTNOTE")) + "</w:footnote>"),
             "word/media/image9.png": b"\x89PNG\r\n\x1a\n" + b"\x00" * 64,
             "docProps/thumbnail.jpeg": b"\xff\xd8\xff\xe0" + b"\x00" * 64}
    text, notes = _read(_raw_docx(_p(_t("BODY")), parts=parts))
    assert text == "BODY", text
    assert [n for n in notes if "no relationship" in n] == [
        "1 Word part(s) holding text that no relationship in the package reaches were not "
        "read; Word does not show them"], notes


# ---------------------------------------------------------------------------
# Review-fix round 3: three behaviours no test held (review finding 4)
# ---------------------------------------------------------------------------


def test_eight_or_more_deletion_passages_above_a_footer_stamp_leave_it_in_the_zone():
    """A Word page's Bates zone is the one it would have without its deletion
    list. With a single deletion line, shorter than the eight-line tail zone,
    no test could tell a tail bound counted over every line from one counted
    over the zone's own lines; with twelve listed passages above a two-line
    footer, the first loses the footer's stamp."""
    from dociq.identify.bates import BatesDecision, DecisionStatus, apply_bates, propose_format

    body = (_p(_t("Letter of 3 March 2019, invented.")) + _p(_t("Second line."))
            + _p(_t("Third line.")) + _p(_t("Fourth line."))
            + "".join(_p(_t(f"W{i} "), _del("Ann", _dt(f"old wording {i}"))) for i in range(12)))
    got = ex.extract("many.docx", _raw_docx(
        body, sect='<w:footerReference w:type="default" r:id="rIdF"/>',
        parts={"word/footer1.xml": _part("ftr", _p(_t("Page footer text")) + _p(_t("MNFV 000777")))},
        rels=[("rIdF", "footer", "footer1.xml")]))
    lines = got.pages[0].text.split("\n")
    assert len(lines) == 30 and lines[-3:] == ["[deleted by Ann] old wording 11",
                                               "Page footer text", "MNFV 000777"], lines
    doc = document("many.docx", got.pages)
    assert [(c.raw, c.line_index) for c in detect_candidates((doc,))] == [("MNFV 000777", 29)]
    decision = BatesDecision(DecisionStatus.CONFIRMED, propose_format((doc,), min_pages=1).format)
    assert apply_bates((doc,), decision)[0].pages[0].bates == "MNFV 000777"


def test_a_table_inside_a_deleted_text_box_row_or_cell_is_listed():
    """A table nested in deleted content is deleted content: its text is
    listed, by the enclosing deletion's author. A table inside a deleted text
    box, or nested in a deleted row or cell, is reached through the deletion
    it sits in; losing that inheritance dropped its text from both the body
    and the list, with no note for the row and the cell."""
    box = ('<w:r><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:inline>'
           '<wp:extent cx="914400" cy="914400"/><wp:docPr id="7" name="tb"/><a:graphic>'
           '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
           '<wps:wsp><wps:txbx><w:txbxContent>' + _p(_t("BOXPARA"))
           + "<w:tbl><w:tr><w:tc>" + _p(_t("BOXCELL 12 May 2017")) + "</w:tc></w:tr></w:tbl>"
           + '</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:inline>'
           '</w:drawing></mc:Choice><mc:Fallback><w:pict/></mc:Fallback></mc:AlternateContent></w:r>')
    nested = lambda text: "<w:tbl><w:tr><w:tc>" + _p(_t(text)) + "</w:tc></w:tr></w:tbl>" + _p("")  # noqa: E731
    body = (_p(_t("KEPT"), _del("Ann", box))
            + "<w:tbl>"
            + f'<w:tr><w:trPr><w:del w:id="6" w:author="Cal" {_WHEN}/></w:trPr><w:tc>'
            + nested("NESTINROW 9 Sept 2015") + "</w:tc></w:tr>"
            + "<w:tr><w:tc>" + _p(_t("ROWKEPT")) + "</w:tc>"
            + f'<w:tc><w:tcPr><w:cellDel w:id="8" w:author="Dan" {_WHEN}/></w:tcPr>'
            + nested("NESTINCELL 2 Feb 2014") + "</w:tc></w:tr></w:tbl>")
    text, notes = _read(_raw_docx(body))
    assert text.split("\n") == [
        "KEPT", "ROWKEPT", "[deleted by Ann] BOXPARA", "[deleted by Ann] BOXCELL 12 May 2017",
        "[deleted by Cal] NESTINROW 9 Sept 2015", "[deleted by Dan] NESTINCELL 2 Feb 2014"], (
        text.split("\n"))
    assert [n for n in notes if ex.has_evidence_marker(n)] == [
        f"{ex.M_WORD_TRACKED_DELETION}: " + _DELETED_GRAPHIC_NOTE.format(n=1)], notes


def test_a_hidden_first_page_header_holding_only_a_chart_is_disclosed_once():
    """A first-page header no section displays, holding nothing but a chart,
    is a part with content: the one marked note naming it is its only record.
    Reading it to decide must not leave its chart counted on the document, as
    if the page showed one."""
    chart = ('<w:r><w:drawing><wp:inline><wp:extent cx="1828800" cy="1828800"/>'
             '<wp:docPr id="2" name="c"/><a:graphic><a:graphicData '
             'uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
             '<c:chart r:id="rIdC"/></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')
    text, notes = _read(_raw_docx(
        _p(_t("BODYTEXT")),
        sect=('<w:headerReference w:type="first" r:id="rIdF"/>'
              '<w:headerReference w:type="default" r:id="rIdH"/>'),
        parts={"word/header1.xml": _part("hdr", _p(_t("DEFAULTHDR"))),
               "word/header2.xml": _part("hdr", _p(chart))},
        rels=[("rIdH", "header", "header1.xml"), ("rIdF", "header", "header2.xml")]))
    assert text.split("\n") == ["DEFAULTHDR", "BODYTEXT"], text.split("\n")
    assert [n for n in notes if ex.has_evidence_marker(n)] == [
        f"{ex.M_WORD_UNREAD}: 1 header/footer part(s) that no section displays (referenced "
        "under a type the section does not show, such as a first-page part without w:titlePg "
        "or an even-page part without w:evenAndOddHeaders, or referenced by no section)"], notes
